#!/usr/bin/env python3
"""ProfitPulse tests: build + formula wiring + LibreOffice recalc vs ground truth."""
import csv
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent.parent
DATA = HERE / "data"
XLSX = HERE / "profitpulse-ratio-dashboard.xlsx"

sys.path.insert(0, str(HERE))
import build_workbook as B  # noqa: E402

YEARS = B.YEARS
ERROR_TOKENS = ("#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#NULL!", "#NUM!", "#N/A")


# ------------------------------------------------------------- fixtures ----
def read_csv(name):
    with open(DATA / name, newline="") as f:
        return list(csv.reader(f))


@pytest.fixture(scope="module")
def wb():
    assert XLSX.exists(), "workbook not built — run build_workbook.py first"
    return load_workbook(XLSX)


@pytest.fixture(scope="module")
def recalc_wb(tmp_path_factory):
    """Headless LibreOffice recalc in ISOLATED dirs (src != outdir)."""
    tmp = tmp_path_factory.mktemp("lo")
    srcdir, outdir = tmp / "src", tmp / "out"
    srcdir.mkdir()
    outdir.mkdir()
    src = srcdir / XLSX.name
    shutil.copy(XLSX, src)
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, f"soffice failed: {r.stderr[-2000:]}"
    out = outdir / (XLSX.stem + ".xlsx")
    assert out.exists(), "LO produced no output"
    return load_workbook(out, data_only=True)


# ------------------------------------------------------------- ground truth
def ground_truth():
    pnl = {r[0]: [float(x) for x in r[1:]] for r in read_csv("pnl.csv")[1:]}
    bs = {r[0]: [float(x) for x in r[1:]] for r in read_csv("bs.csv")[1:]}
    g = lambda k, i: pnl[k][i]          # i: 0..4 -> FY21..FY25
    h = lambda k, i: bs[k][i + 1]       # bs col index; FY21..FY25 -> 1..5
    hp = lambda k, i: bs[k][i]          # previous year
    out = {}
    for i, y in enumerate(YEARS):
        rev, cogs = g("Revenue", i), g("COGS", i)
        tca = sum(h(k, i) for k in ("Cash & bank (plug)", "Receivables", "Inventory", "Other current assets"))
        tcl = sum(h(k, i) for k in ("Payables", "Short-term debt", "Other current liabilities"))
        ta = tca + sum(h(k, i) for k in ("PPE (net)", "Other non-current assets"))
        tl = tcl + sum(h(k, i) for k in ("Long-term debt", "Other non-current liabilities"))
        te = ta - tl

        def bs_total(i_):
            tca_ = sum(bs[k][i_] for k in ("Cash & bank (plug)", "Receivables", "Inventory", "Other current assets"))
            tnca_ = sum(bs[k][i_] for k in ("PPE (net)", "Other non-current assets"))
            tcl_ = sum(bs[k][i_] for k in ("Payables", "Short-term debt", "Other current liabilities"))
            tncl_ = sum(bs[k][i_] for k in ("Long-term debt", "Other non-current liabilities"))
            te_ = bs["Share capital"][i_] + bs["Reserves & surplus"][i_]
            return tca_ + tnca_, tcl_ + tncl_ + te_, te_
        ta_c, tle_c, te_c = bs_total(i + 1)
        ta_p, tle_p, te_p = bs_total(i)
        avgA, avgE = (ta_c + ta_p) / 2, (te_c + te_p) / 2
        avgRecv = (bs["Receivables"][i + 1] + bs["Receivables"][i]) / 2
        avgInv = (bs["Inventory"][i + 1] + bs["Inventory"][i]) / 2
        avgPay = (bs["Payables"][i + 1] + bs["Payables"][i]) / 2
        gross = g("Gross profit", i) if "Gross profit" in pnl else rev - cogs
        out[y] = {
            "Current ratio": tca / tcl,
            "Quick ratio": (h("Cash & bank (plug)", i) + h("Receivables", i)) / tcl,
            "Working capital": tca - tcl,
            "Gross margin": (rev - cogs) / rev,
            "Net margin": g("Net income", i) / rev,
            "ROE": g("Net income", i) / avgE,
            "ROA": g("Net income", i) / avgA,
            "ROCE": g("EBIT", i) / (te_c + sum(bs[k][i + 1] for k in ("Long-term debt", "Other non-current liabilities"))),
            "D/E": (ta_c - te_c) / te_c,
            "Interest coverage": g("EBIT", i) / g("Interest", i),
            "Receivable days": avgRecv / rev * 365,
            "Inventory days": avgInv / cogs * 365,
            "Payable days": avgPay / cogs * 365,
            "Asset turnover": rev / avgA,
            "Revenue": rev, "Net income": g("Net income", i),
            "Assets": ta_c, "Equity": te_c,
        }
        out[y]["CCC"] = out[y]["Receivable days"] + out[y]["Inventory days"] - out[y]["Payable days"]
    return out


# ------------------------------------------------------------------ tests --
def test_data_files_exist():
    for n in ("pnl.csv", "bs.csv", "meta.csv"):
        assert (DATA / n).exists(), n


def test_workbook_builds():
    B.main()
    assert XLSX.exists() and XLSX.stat().st_size > 10_000


def test_sheets_present(wb):
    for s in ("Guide", "Raw_PNL", "Raw_BS", "Ratios", "DuPont",
              "Dashboard", "AI_Report", "Checks"):
        assert s in wb.sheetnames, s


def test_no_invalid_range_literals(wb):
    """Excel range literals must repeat the column letter on both ends."""
    bad = re.compile(r"\$[A-Z]+\$\d+:\$\d+")
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and bad.search(c.value):
                    hits.append(f"{ws.title}!{c.coordinate}: {c.value[:80]}")
    assert not hits, f"invalid range literals: {hits[:5]}"


def test_ratios_are_formulas_not_values(wb):
    ws = wb["Ratios"]
    for row in range(2, 2 + len(B.RATIO_DEFS)):
        for col in range(2, 7):
            v = ws.cell(row=row, column=col).value
            assert isinstance(v, str) and (v.startswith("=") or v == "n/a"), \
                f"Ratios!{B.C(col)}{row} is a hard value: {v!r}"


def test_formula_wiring_spot_checks(wb):
    r, ws = B.RATIO_R, wb["Ratios"]
    cell = ws.cell(row=r["Current ratio"], column=2).value  # FY21
    assert "Raw_BS!C8" in cell and "Raw_BS!C14" in cell, cell
    cell = ws.cell(row=r["Receivable days"], column=6).value  # FY25
    assert "365" in cell and "Raw_PNL!F2" in cell, cell
    assert wb["Dashboard"].cell(row=2, column=2).value == "=Raw_PNL!F2"
    assert wb["DuPont"].cell(row=7, column=6).value == "=DuPont!F5-DuPont!F6"


def test_no_numeric_blank_guards(wb):
    """Numeric formula columns must guard to 0, never ''."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and '""' in v:
                    assert False, f"{ws.title}!{c.coordinate} uses blank-guard \"\": {v[:90]}"


def test_lo_recalc_matches_ground_truth(recalc_wb):
    gt = ground_truth()
    ws = recalc_wb["Ratios"]
    r = B.RATIO_R
    mapping = {
        "Current ratio": r["Current ratio"], "Quick ratio": r["Quick ratio"],
        "Working capital": r["Working capital"], "Gross margin": r["Gross margin"],
        "Net margin": r["Net margin"], "ROE": r["Return on equity (ROE)"],
        "ROA": r["Return on assets (ROA)"], "ROCE": r["Return on capital employed (ROCE)"],
        "D/E": r["Debt / Equity"], "Interest coverage": r["Interest coverage (x)"],
        "Receivable days": r["Receivable days"], "Inventory days": r["Inventory days"],
        "Payable days": r["Payable days"], "Asset turnover": r["Asset turnover (x)"],
    }
    for y in YEARS:
        j = YEARS.index(y) + 2
        for name, row in mapping.items():
            got = ws.cell(row=row, column=j).value
            assert isinstance(got, (int, float)), f"{y} {name}: {got!r}"
            assert got == pytest.approx(gt[y][name], rel=1e-6, abs=1e-3), \
                f"{y} {name}: LO={got} vs truth={gt[y][name]}"
        ccc = ws.cell(row=r["Cash conversion cycle (days)"], column=j).value
        assert ccc == pytest.approx(gt[y]["CCC"], rel=1e-6, abs=1e-2), f"{y} CCC"


def test_no_formula_errors_after_recalc(recalc_wb):
    bad = []
    for ws in recalc_wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERROR_TOKENS:
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not bad, f"formula errors: {bad[:10]}"


def test_checks_all_pass_after_recalc(recalc_wb):
    ws = recalc_wb["Checks"]
    for row in range(3, 11):
        for col in range(2, 7):
            v = ws.cell(row=row, column=col).value
            assert v is True or v == pytest.approx(0, abs=1e-6), \
                f"Checks row {row} col {col}: {v!r}"


def test_dashboard_has_three_charts(wb):
    assert len(wb["Dashboard"]._charts) == 3
