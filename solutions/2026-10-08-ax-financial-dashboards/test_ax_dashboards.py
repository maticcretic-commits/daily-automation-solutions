"""Tests for the AX 2012 financial dashboards generator."""
import csv
import os
import subprocess
import unittest

import openpyxl

SRC = "sample_ax_export.csv"
XLSX = "AX_2012_Financial_Dashboards.xlsx"
RECALC_DIR = "/tmp/df-20261008/recalc"

EXPECTED_SHEETS = ["Data", "Income Statement", "Balance Sheet",
                   "KPI Dashboard", "ChartData", "Charts"]
CATEGORIES = {"Revenue", "Expense", "Asset", "Asset-Contra", "Liability", "Equity"}
SIGN = {"Revenue": -1, "Expense": 1, "Asset": 1, "Asset-Contra": 1,
        "Liability": -1, "Equity": -1}


def load_csv(path=SRC):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def signed(r):
    return SIGN[r["AccountCategory"]] * (float(r["Debit"]) - float(r["Credit"]))


def annual_total(rows, predicate):
    return round(sum(signed(r) for r in rows if predicate(r)), 2)


def recalc_values():
    """Roundtrip through LibreOffice so formulas get computed values."""
    os.makedirs(RECALC_DIR, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                    "--outdir", RECALC_DIR, XLSX],
                   check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(os.path.join(RECALC_DIR, XLSX), data_only=True)
    return wb


class TestSampleData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = load_csv()

    def test_expected_columns(self):
        self.assertEqual(
            list(self.rows[0].keys()),
            ["FiscalYear", "Period", "MainAccount", "AccountName",
             "AccountCategory", "Debit", "Credit"])

    def test_row_count(self):
        # 26 accounts x 12 periods
        self.assertEqual(len(self.rows), 312)

    def test_each_account_in_each_period_once(self):
        seen = set()
        for r in self.rows:
            key = (r["MainAccount"], int(r["Period"]))
            self.assertNotIn(key, seen)
            seen.add(key)
        self.assertEqual(len(seen), 312)

    def test_periods_are_1_to_12(self):
        self.assertEqual(sorted({int(r["Period"]) for r in self.rows}),
                         list(range(1, 13)))

    def test_categories_known(self):
        self.assertTrue({r["AccountCategory"] for r in self.rows} <= CATEGORIES)

    def test_trial_balance_proves_out_every_period(self):
        for p in range(1, 13):
            prows = [r for r in self.rows if int(r["Period"]) == p]
            dr = sum(float(r["Debit"]) for r in prows)
            cr = sum(float(r["Credit"]) for r in prows)
            self.assertAlmostEqual(dr, cr, delta=0.05,
                                   msg=f"period {p}: Dr {dr} != Cr {cr}")

    def test_revenue_is_positive_when_signed(self):
        rev = annual_total(self.rows, lambda r: r["AccountCategory"] == "Revenue")
        self.assertGreater(rev, 0)

    def test_assets_equal_liabilities_plus_equity_in_data(self):
        # Balance sheet as REPORTED: retained earnings = opening RE + profit.
        for p in range(1, 13):
            prows = [r for r in self.rows if int(r["Period"]) == p]
            a = sum(signed(r) for r in prows
                    if r["AccountCategory"] in ("Asset", "Asset-Contra"))
            l = sum(signed(r) for r in prows
                    if r["AccountCategory"] == "Liability")
            e = sum(signed(r) for r in prows
                    if r["AccountCategory"] == "Equity")
            rev = sum(signed(r) for r in prows
                      if r["AccountCategory"] == "Revenue")
            exp = sum(signed(r) for r in prows
                      if r["AccountCategory"] == "Expense")
            self.assertAlmostEqual(a, l + e + (rev - exp), delta=0.05,
                                   msg=f"period {p} balance sheet does not balance")

    def test_deterministic_regeneration(self):
        import make_sample_data
        before = load_csv()
        make_sample_data.main()
        after = load_csv()
        self.assertEqual(before, after)


class TestWorkbookStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(XLSX)

    def test_all_sheets_present(self):
        self.assertEqual(self.wb.sheetnames, EXPECTED_SHEETS)

    def test_data_sheet_has_all_rows(self):
        ws = self.wb["Data"]
        self.assertEqual(ws.max_row, 313)  # header + 312 rows

    def test_data_sheet_signed_column_spot_check(self):
        ws = self.wb["Data"]
        for row in ws.iter_rows(min_row=2, max_row=313, values_only=True):
            _, _, _, _, cat, dr, cr, sgn = row
            self.assertAlmostEqual(sgn, SIGN[cat] * (dr - cr), delta=0.01)

    def test_income_statement_period_headers(self):
        ws = self.wb["Income Statement"]
        self.assertEqual([ws.cell(row=1, column=c).value for c in range(3, 15)],
                         list(range(1, 13)))
        self.assertEqual(ws.cell(row=1, column=15).value, "Total")

    def test_account_cells_use_sumifs_on_data(self):
        ws = self.wb["Income Statement"]
        f = ws["C3"].value
        self.assertIn("SUMIFS", f)
        self.assertIn("Data!$H:$H", f)
        self.assertIn("Data!$B:$B", f)
        self.assertIn("Data!$C:$C", f)

    def test_total_revenue_row_sums_revenue_lines(self):
        ws = self.wb["Income Statement"]
        # revenue accounts in rows 3-5, total in row 6
        self.assertEqual(ws["B6"].value, "Total Revenue")
        self.assertEqual(ws["C6"].value, "=SUM(C3:C5)")
        self.assertEqual(ws["O6"].value, "=SUM(C6:N6)")

    def test_net_income_row_is_revenue_minus_expenses(self):
        ws = self.wb["Income Statement"]
        self.assertEqual(ws["B19"].value, "Total Expenses")
        self.assertEqual(ws["B20"].value, "Net Income")
        self.assertEqual(ws["C20"].value, "=C6-C19")
        self.assertEqual(ws["O20"].value, "=O6-O19")

    def test_balance_sheet_check_row_exists(self):
        ws = self.wb["Balance Sheet"]
        self.assertEqual(ws["B24"].value, "Check (Assets - L&E) = 0")
        self.assertEqual(ws["C24"].value, "=C9-C23")

    def test_balance_sheet_total_assets_row(self):
        ws = self.wb["Balance Sheet"]
        self.assertEqual(ws["B9"].value, "Total Assets")
        self.assertEqual(ws["C9"].value, "=SUM(C3:C8)")

    def test_retained_earnings_rolls_in_current_profit(self):
        ws = self.wb["Balance Sheet"]
        self.assertEqual(ws["B21"].value,
                         "Retained Earnings (incl. current profit)")
        f = ws["C21"].value
        self.assertIn("SUMIFS", f)
        self.assertIn("'Income Statement'!C20", f)

    def test_no_broken_formulas_anywhere(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        self.assertNotIn("#REF!", v, f"{ws.title}!{cell.coordinate}")
                        self.assertNotIn("#VALUE!", v, f"{ws.title}!{cell.coordinate}")

    def test_kpi_sheet_has_ten_items(self):
        ws = self.wb["KPI Dashboard"]
        labels = [ws.cell(row=r, column=1).value for r in range(3, 13)]
        self.assertEqual(len([l for l in labels if l]), 10)

    def test_kpi_margin_formula(self):
        ws = self.wb["KPI Dashboard"]
        self.assertEqual(ws["A6"].value, "Net Margin")
        self.assertEqual(ws["B6"].value, "=B5/B3")
        self.assertEqual(ws["B6"].number_format, "0.0%")

    def test_kpi_ratio_formulas(self):
        ws = self.wb["KPI Dashboard"]
        self.assertEqual(ws["B9"].value, "=B7/B8")   # current ratio
        self.assertEqual(ws["B12"].value, "=B10/B11")  # debt-to-equity

    def test_chartdata_references_income_statement(self):
        ws = self.wb["ChartData"]
        self.assertEqual(ws["B2"].value, "='Income Statement'!C6")
        self.assertEqual(ws["D13"].value, "='Income Statement'!N20")

    def test_three_charts_present(self):
        ws = self.wb["Charts"]
        self.assertEqual(len(ws._charts), 3)
        kinds = sorted(type(c).__name__ for c in ws._charts)
        self.assertEqual(kinds, ["BarChart", "LineChart", "PieChart"])

    def test_bar_chart_has_two_series(self):
        ws = self.wb["Charts"]
        bar = next(c for c in ws._charts if type(c).__name__ == "BarChart")
        self.assertEqual(len(bar.series), 2)


class TestRecalculatedValues(unittest.TestCase):
    """LibreOffice recalc roundtrip: formulas must compute the right numbers."""

    @classmethod
    def setUpClass(cls):
        cls.rows = load_csv()
        cls.wb = recalc_values()

    def test_total_revenue_value(self):
        expected = annual_total(self.rows,
                                lambda r: r["AccountCategory"] == "Revenue")
        got = self.wb["Income Statement"]["O6"].value
        self.assertAlmostEqual(got, expected, delta=1.0)

    def test_total_expenses_value(self):
        expected = annual_total(self.rows,
                                lambda r: r["AccountCategory"] == "Expense")
        got = self.wb["Income Statement"]["O19"].value
        self.assertAlmostEqual(got, expected, delta=1.0)

    def test_net_income_value(self):
        rev = annual_total(self.rows, lambda r: r["AccountCategory"] == "Revenue")
        exp = annual_total(self.rows, lambda r: r["AccountCategory"] == "Expense")
        got = self.wb["Income Statement"]["O20"].value
        self.assertAlmostEqual(got, rev - exp, delta=1.0)

    def test_balance_sheet_balances(self):
        got = self.wb["Balance Sheet"]["O24"].value
        self.assertAlmostEqual(got or 0, 0, delta=1.0)

    def test_net_margin_value(self):
        rev = annual_total(self.rows, lambda r: r["AccountCategory"] == "Revenue")
        exp = annual_total(self.rows, lambda r: r["AccountCategory"] == "Expense")
        got = self.wb["KPI Dashboard"]["B6"].value
        self.assertAlmostEqual(got, (rev - exp) / rev, delta=0.001)

    def test_current_ratio_value(self):
        ca = annual_total(self.rows, lambda r: r["MainAccount"] in
                          ("101000", "102000", "103000", "104000"))
        cl = annual_total(self.rows, lambda r: r["MainAccount"] in
                          ("201000", "202000", "203000", "205000"))
        got = self.wb["KPI Dashboard"]["B9"].value
        self.assertAlmostEqual(got, ca / cl, delta=0.01)

    def test_monthly_revenue_series(self):
        ws = self.wb["ChartData"]
        rev_p1 = sum(signed(r) for r in self.rows
                     if int(r["Period"]) == 1 and r["AccountCategory"] == "Revenue")
        self.assertAlmostEqual(ws["B2"].value, rev_p1, delta=1.0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
