#!/usr/bin/env python3
"""Tests for InvoiceLedger.

Covers CSV integrity, workbook structure/formulas, and a full LibreOffice
headless recalculation compared against an independent Python ground truth
computed straight from the CSVs.
"""
import csv
import os
import shutil
import subprocess
import unittest
from datetime import date

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
XLSX = os.path.join(BASE, "invoice-ledger-reports.xlsx")
REPORT_DATE = date(2026, 9, 30)
MONTHS = [(2025, 10), (2025, 11), (2025, 12), (2026, 1), (2026, 2), (2026, 3),
          (2026, 4), (2026, 5), (2026, 6), (2026, 7), (2026, 8), (2026, 9)]


def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.DictReader(f))


def ground_truth():
    inv = read_csv("invoices.csv")
    pay = read_csv("payments.csv")
    g = {"monthly_rev": {(y, m): 0.0 for y, m in MONTHS},
         "monthly_net": {(y, m): 0.0 for y, m in MONTHS},
         "buckets": {"0-30": 0.0, "31-60": 0.0, "61-90": 0.0, "90+": 0.0},
         "outstanding": 0.0, "receipts": 0.0, "paid_inv": 0.0,
         "paid_exp": 0.0, "gst_net": 0.0, "rev": 0.0, "cogs": 0.0,
         "exp": 0.0, "coll_days": []}
    for i in inv:
        d = date.fromisoformat(i["date"])
        amt, gst = float(i["amount"]), float(i["gst_amt"])
        key = (d.year, d.month)
        if i["type"] == "SALES":
            g["rev"] += amt
            g["monthly_rev"][key] += amt
            g["gst_net"] += gst
            if i["paid_date"]:
                g["coll_days"].append(
                    (date.fromisoformat(i["paid_date"]) - d).days)
        elif i["type"] == "PURCHASE":
            g["cogs"] += amt
            g["gst_net"] -= gst
        else:
            g["exp"] += amt
            g["gst_net"] -= gst
        bal = float(i["total"]) - float(i["paid_amount"])
        if bal > 0.005:
            g["outstanding"] += bal
            days = (REPORT_DATE - d).days
            b = ("0-30" if days <= 30 else "31-60" if days <= 60
                 else "61-90" if days <= 90 else "90+")
            g["buckets"][b] += bal
    for p in pay:
        a = float(p["amount"])
        if p["type"] == "SALES":
            g["receipts"] += a
        elif p["type"] == "PURCHASE":
            g["paid_inv"] += a
        else:
            g["paid_exp"] += a
    # monthly net profit = rev - cogs - exp per month
    for i in inv:
        d = date.fromisoformat(i["date"])
        key = (d.year, d.month)
        a = float(i["amount"])
        if i["type"] == "SALES":
            g["monthly_net"][key] += a
        else:
            g["monthly_net"][key] -= a
    g["net"] = g["rev"] - g["cogs"] - g["exp"]
    g["gross"] = g["rev"] - g["cogs"]
    g["net_cash"] = g["receipts"] - g["paid_inv"] - g["paid_exp"]
    g["avg_coll"] = (sum(g["coll_days"]) / len(g["coll_days"])
                     if g["coll_days"] else 0.0)
    return g


def recalc():
    """LibreOffice headless recalculation of a COPY in a separate outdir."""
    src = "/tmp/il_src"
    out = "/tmp/il_out"
    shutil.rmtree(src, ignore_errors=True)
    shutil.rmtree(out, ignore_errors=True)
    os.makedirs(src)
    os.makedirs(out)
    shutil.copy(XLSX, os.path.join(src, "workbook.xlsx"))
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out,
         os.path.join(src, "workbook.xlsx")],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:500]}"
    return load_workbook(os.path.join(out, "workbook.xlsx"), data_only=True)


class TestData(unittest.TestCase):
    def test_csvs_exist_and_counts(self):
        inv = read_csv("invoices.csv")
        pay = read_csv("payments.csv")
        self.assertGreater(len(inv), 500)
        self.assertGreater(len(pay), 400)
        sales = [i for i in inv if i["type"] == "SALES"]
        self.assertGreater(len(sales), 300)

    def test_invoice_totals_tie(self):
        for i in read_csv("invoices.csv"):
            self.assertAlmostEqual(float(i["total"]),
                                   float(i["amount"]) + float(i["gst_amt"]),
                                   places=2, msg=i["inv_id"])

    def test_paid_never_exceeds_total(self):
        for i in read_csv("invoices.csv"):
            self.assertLessEqual(float(i["paid_amount"]),
                                 float(i["total"]) + 0.005, msg=i["inv_id"])

    def test_dates_in_range(self):
        for i in read_csv("invoices.csv"):
            d = date.fromisoformat(i["date"])
            self.assertTrue(date(2025, 10, 1) <= d <= REPORT_DATE)


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(XLSX)

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames,
                         ["Guide", "Raw_Invoices", "Raw_Payments", "P&L",
                          "CashFlow", "Receivables", "Dashboard",
                          "AI_Report", "Checks"])

    def test_formula_cells_present(self):
        pl = self.wb["P&L"]
        self.assertTrue(str(pl["B3"].value).startswith("=SUMIFS"))
        self.assertEqual(pl["B6"].value, "=SUM(B3:B5)")
        self.assertEqual(self.wb["Checks"]["B3"].value[:5], "=ABS(")
        self.assertEqual(self.wb["Dashboard"]["B3"].value, "='P&L'!N6")
        self.assertIn("Guide!$B$2", str(self.wb["Receivables"]["G4"].value))

    def test_charts_exist(self):
        charts = self.wb["Dashboard"]._charts
        self.assertEqual(len(charts), 3)


class TestRecalc(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = recalc()
        cls.g = ground_truth()

    def v(self, sheet, cell):
        return self.wb[sheet][cell].value

    def test_pl_fy(self):
        self.assertAlmostEqual(self.v("P&L", "N6"), self.g["rev"], delta=0.05)
        self.assertAlmostEqual(self.v("P&L", "N7"), self.g["cogs"], delta=0.05)
        self.assertAlmostEqual(self.v("P&L", "N8"), self.g["gross"], delta=0.05)
        self.assertAlmostEqual(self.v("P&L", "N16"), self.g["exp"], delta=0.05)
        self.assertAlmostEqual(self.v("P&L", "N17"), self.g["net"], delta=0.05)
        self.assertAlmostEqual(self.v("P&L", "N22"), self.g["gst_net"],
                               delta=0.05)

    def test_pl_months(self):
        letters = "BCDEFGHIJKLM"
        for L, ym in zip(letters, MONTHS):
            self.assertAlmostEqual(self.v("P&L", f"{L}6"),
                                   self.g["monthly_rev"][ym], delta=0.05,
                                   msg=f"rev {ym}")
        for L, ym in zip(letters, MONTHS):
            self.assertAlmostEqual(self.v("P&L", f"{L}17"),
                                   self.g["monthly_net"][ym], delta=0.05,
                                   msg=f"net {ym}")

    def test_cashflow(self):
        self.assertAlmostEqual(self.v("CashFlow", "N3"), self.g["receipts"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("CashFlow", "N4"), self.g["paid_inv"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("CashFlow", "N5"), self.g["paid_exp"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("CashFlow", "N6"), self.g["net_cash"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("CashFlow", "N7"), self.g["net_cash"],
                               delta=0.05)

    def test_receivables(self):
        self.assertAlmostEqual(self.v("Receivables", "K7"),
                               self.g["outstanding"], delta=0.05)
        for cell, b in zip(("K3", "K4", "K5", "K6"),
                           ("0-30", "31-60", "61-90", "90+")):
            self.assertAlmostEqual(self.v("Receivables", cell),
                                   self.g["buckets"][b], delta=0.05,
                                   msg=f"bucket {b}")
        self.assertAlmostEqual(self.v("Receivables", "K9"), self.g["avg_coll"],
                               delta=0.5)

    def test_dashboard_kpis(self):
        self.assertAlmostEqual(self.v("Dashboard", "B3"), self.g["rev"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("Dashboard", "B4"), self.g["net"],
                               delta=0.05)
        self.assertAlmostEqual(self.v("Dashboard", "B6"),
                               self.g["outstanding"], delta=0.05)
        self.assertAlmostEqual(self.v("Dashboard", "B7"), self.g["avg_coll"],
                               delta=0.5)
        self.assertAlmostEqual(self.v("Dashboard", "B8"), self.g["net_cash"],
                               delta=0.05)

    def test_all_checks_true(self):
        for r in range(3, 11):
            self.assertTrue(self.wb["Checks"][f"B{r}"].value,
                            msg=f"check row {r} failed")

    def test_flags_render(self):
        for r in range(7, 15):
            txt = str(self.wb["AI_Report"][f"A{r}"].value or "")
            self.assertTrue(txt.startswith(("✓", "⚠", "ℹ")),
                            msg=f"flag row {r}: {txt[:40]}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
