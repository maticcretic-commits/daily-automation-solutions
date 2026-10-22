"""ReconPulse tests: build the workbook, recalc headless in LibreOffice on a
COPY in a separate outdir (never the source dir), and compare every cached
value against main.py ground truth. Zero formula errors allowed except
expected #N/A in Recon!G (no key match)."""
import os
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import build_workbook
import main

OUTDIR = os.path.join(ROOT, "recalc-out")
os.makedirs(OUTDIR, exist_ok=True)


def recalc_values():
    src = os.path.join(ROOT, "reconpulse.xlsx")
    if not os.path.exists(src):
        build_workbook.build()
    work = os.path.join(OUTDIR, "recon-src")
    os.makedirs(work, exist_ok=True)
    dst = os.path.join(work, "reconpulse.xlsx")
    shutil.copy(src, dst)
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", OUTDIR, dst],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:500]}"
    conv = os.path.join(OUTDIR, "reconpulse.xlsx")
    wb = load_workbook(conv, data_only=True)
    return wb


@pytest.fixture(scope="module")
def wb():
    return recalc_values()


@pytest.fixture(scope="module")
def truth():
    return main.compute()


def val(wb, sheet, cell):
    return wb[sheet][cell].value


def test_sheets_present(wb):
    for s in ("Guide", "Raw_Bank", "Raw_Book", "Recon", "Unmatched",
              "Daily", "Summary", "Dashboard", "Checks"):
        assert s in wb.sheetnames, s


def test_row_counts(wb, truth):
    assert wb["Raw_Bank"].max_row == truth["bank_rows"] + 1
    assert wb["Raw_Book"].max_row == truth["book_rows"] + 1
    assert wb["Recon"].max_row == truth["bank_rows"] + 1


def test_recon_counts(wb, truth):
    assert val(wb, "Summary", "B21") == truth["matched"]
    assert val(wb, "Summary", "B22") == truth["unmatched_bank"]
    assert val(wb, "Summary", "B23") == truth["unmatched_book"]


def test_balances(wb, truth):
    assert abs(val(wb, "Summary", "B4") - truth["bank_close"]) < 0.01
    assert abs(val(wb, "Summary", "B5") - truth["book_close"]) < 0.01
    assert abs(val(wb, "Summary", "B11") - truth["adjusted_bank"]) < 0.01
    assert abs(val(wb, "Summary", "B16") - truth["adjusted_book"]) < 0.01
    assert abs(val(wb, "Summary", "B17") - truth["recon_diff"]) < 0.01


def test_walk_components(wb, truth):
    assert abs(val(wb, "Summary", "B8") - truth["bank_charges"]) < 0.01
    assert abs(val(wb, "Summary", "B9") - truth["interest"]) < 0.01
    assert abs(val(wb, "Summary", "B10") - truth["unidentified"]) < 0.01
    assert abs(val(wb, "Summary", "B14") - truth["unpresented"]) < 0.01
    assert abs(val(wb, "Summary", "B15") - truth["transit"]) < 0.01


def test_kpis(wb, truth):
    assert abs(val(wb, "Summary", "B19") - truth["total_receipts"]) < 0.01
    assert abs(val(wb, "Summary", "B20") - truth["total_payments"]) < 0.01
    assert abs(val(wb, "Summary", "B24") - truth["matched"] / truth["bank_rows"]) < 1e-6
    assert val(wb, "Summary", "B25") == truth["oldest_days"]


def test_daily_ties(wb, truth):
    last = truth["daily"][-1]
    assert abs(val(wb, "Daily", "B22") - last[1]) < 0.01
    assert abs(val(wb, "Daily", "C22") - last[2]) < 0.01
    # every daily EOD matches ground truth
    for i, (d, bbal, kbal) in enumerate(truth["daily"], 2):
        assert abs(val(wb, "Daily", f"B{i}") - bbal) < 0.01, f"bank EOD row {i}"
        assert abs(val(wb, "Daily", f"C{i}") - kbal) < 0.01, f"book EOD row {i}"


def test_all_checks_true(wb):
    for i in range(2, 10):
        assert val(wb, "Checks", f"C{i}") is True, f"check row {i} failed"


def test_unmatched_extraction(wb, truth):
    bank_items = [val(wb, "Unmatched", f"A{r}") for r in range(3, 9)]
    bank_items = [x for x in bank_items if x]
    assert sorted(bank_items) == truth["unmatched_bank_desc"]
    book_items = [val(wb, "Unmatched", f"A{r}") for r in range(12, 20)]
    book_items = [x for x in book_items if x]
    assert sorted(book_items) == truth["unmatched_book_narr"]


def test_formulas_not_values():
    wb = load_workbook(os.path.join(ROOT, "reconpulse.xlsx"), data_only=False)
    assert str(wb["Recon"]["K2"].value).startswith("=IF(")
    assert str(wb["Recon"]["G2"].value).startswith("=MATCH(")
    assert str(wb["Raw_Book"]["G2"].value).startswith("=TEXT(")
    assert str(wb["Summary"]["B17"].value).startswith("=B11-B16")
    assert str(wb["Unmatched"]["A3"].value).startswith('=IF($H3="","",INDEX(')
    assert str(wb["Unmatched"]["H3"].value).startswith("=IFERROR(SMALL(")


def test_no_formula_errors(wb):
    bad = {"#NAME?", "#VALUE!", "#DIV/0!", "#REF!", "Err:502"}
    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v in bad:
                    problems.append(f"{ws.title}!{c.coordinate}={v}")
                if isinstance(v, str) and v == "#N/A" and not (
                        ws.title == "Recon" and c.column == 7):
                    problems.append(f"{ws.title}!{c.coordinate}=#N/A (unexpected)")
    assert not problems, f"formula errors: {problems[:10]}"
