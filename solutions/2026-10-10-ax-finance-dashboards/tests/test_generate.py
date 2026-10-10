"""Tests for the AX-Style Finance Dashboards generator.

Includes a LibreOffice headless recalc: the workbook is converted (forcing a
full formula recalc) and key statement figures are compared against ground
truth computed independently from the CSV inputs.
"""
import csv
import datetime as dt
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent.parent
DATA = HERE / "data"
XLSX = HERE / "ax-finance-dashboards.xlsx"
RECALC_DIR = HERE / "tests" / "_recalc"
PERIOD_START = dt.date(2026, 4, 1)
PERIOD_END = dt.date(2026, 10, 1)
TOL = 0.02


def load_csv(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def net_by_account(rows):
    """per-account net (debit - credit) for period / all-dates / pre-period."""
    per, alld, opened = {}, {}, {}
    for r in rows:
        d = dt.date.fromisoformat(r["date"])
        v = float(r["debit"]) - float(r["credit"])
        alld[r["account_code"]] = alld.get(r["account_code"], 0.0) + v
        if PERIOD_START <= d < PERIOD_END:
            per[r["account_code"]] = per.get(r["account_code"], 0.0) + v
        if d < PERIOD_START:
            opened[r["account_code"]] = opened.get(r["account_code"], 0.0) + v
    return per, alld, opened


@pytest.fixture(scope="module")
def ground():
    gl = load_csv("gl_transactions.csv")
    per, alld, opened = net_by_account(gl)
    g = lambda a: per.get(a, 0.0)
    rev = -(g("4000") + g("4100"))
    cogs = g("5000")
    opex = sum(g(a) for a in ["6000", "6100", "6200", "6300", "6400", "6999"])
    gross = rev - cogs
    op_inc = gross - opex
    ni = op_inc + (-g("7000")) - g("7100")
    return {
        "revenue": rev, "cogs": cogs, "opex": opex, "gross": gross,
        "op_income": op_inc, "net_income": ni,
        "cash_close": alld.get("1000", 0.0),
        "cash_open": opened.get("1000", 0.0),
        "op_cash": ni + g("6400")
                   - ((alld.get("1100", 0) - opened.get("1100", 0)))
                   - ((alld.get("1200", 0) - opened.get("1200", 0)))
                   # 2100 is credit-normal: economic increase = -(net_all - net_open)
                   - ((alld.get("2100", 0) - opened.get("2100", 0))),
        "invest": -(g("1500")),
        "finance": -(g("2500")),
    }


@pytest.fixture(scope="module")
def recalc_wb():
    assert XLSX.exists(), "run generate.py first"
    RECALC_DIR.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         str(XLSX), "--outdir", str(RECALC_DIR)],
        check=True, capture_output=True, timeout=180)
    out = RECALC_DIR / XLSX.name
    assert out.exists(), "libreoffice recalc produced no output"
    return load_workbook(out, data_only=True)


# ---- input data ----
def test_csvs_parse():
    coa = load_csv("chart_of_accounts.csv")
    gl = load_csv("gl_transactions.csv")
    assert len(coa) == 19 and len(gl) == 169
    assert {c for c in coa[0]} == {"account_code", "account_name", "account_group",
                                  "normal_balance", "cashflow_section"}


def test_entries_balanced():
    gl = load_csv("gl_transactions.csv")
    by_entry = {}
    for r in gl:
        d, c = by_entry.setdefault(r["entry_id"], [0.0, 0.0])
        by_entry[r["entry_id"]][0] += float(r["debit"])
        by_entry[r["entry_id"]][1] += float(r["credit"])
    bad = [e for e, (d, c) in by_entry.items() if abs(d - c) > 0.01]
    assert not bad, f"unbalanced entries: {bad}"


# ---- workbook structure ----
def test_sheets_present():
    wb = load_workbook(XLSX)
    assert wb.sheetnames == ["GL_Transactions", "Chart_of_Accounts", "Trial_Balance",
                             "P&L", "Cash_Flow", "Dashboard"]


def test_named_ranges():
    wb = load_workbook(XLSX)
    names = set(wb.defined_names)
    assert {"GL_Date", "GL_Acct", "GL_Net", "GL_Debit", "GL_Credit"} <= names


def test_formulas_are_live_not_hardcoded():
    wb = load_workbook(XLSX)
    pl = wb["P&L"]
    assert str(pl["H19"].value).startswith("=")   # net income total
    assert "SUMIFS" in str(pl["B4"].value)         # monthly revenue line
    cf = wb["Cash_Flow"]
    assert str(cf["B18"].value).startswith("=")    # closing cash


def test_dashboard_has_three_charts():
    wb = load_workbook(XLSX)
    assert len(wb["Dashboard"]._charts) == 3


# ---- recalculated values vs ground truth ----
def test_trial_balance_ok(recalc_wb):
    tb = recalc_wb["Trial_Balance"]
    val = next(c.value for row in tb.iter_rows(min_row=1, max_col=2)
               for c in row if c.value == "Balanced?")
    assert val == "Balanced?"
    # the cell to the right of the label holds the check result
    for row in tb.iter_rows(min_row=1, max_col=2):
        if row[0].value == "Balanced?":
            assert "OK" in str(row[1].value), row[1].value


def test_pl_totals(recalc_wb, ground):
    pl = recalc_wb["P&L"]
    assert abs(pl["H6"].value - ground["revenue"]) < TOL      # total revenue
    assert abs(pl["H7"].value - ground["cogs"]) < TOL         # COGS
    assert abs(pl["H15"].value - ground["opex"]) < TOL        # total opex
    assert abs(pl["H8"].value - ground["gross"]) < TOL        # gross profit
    assert abs(pl["H16"].value - ground["op_income"]) < TOL  # operating income
    assert abs(pl["H19"].value - ground["net_income"]) < TOL  # net income


def test_pl_monthly_adds_to_total(recalc_wb):
    pl = recalc_wb["P&L"]
    for r in (6, 19):  # revenue total, net income
        months = [pl.cell(row=r, column=j).value or 0 for j in range(2, 8)]
        assert abs(sum(months) - pl.cell(row=r, column=8).value) < TOL


def test_cashflow_sections(recalc_wb, ground):
    cf = recalc_wb["Cash_Flow"]
    assert abs(cf["B9"].value - ground["op_cash"]) < TOL
    assert abs(cf["B12"].value - ground["invest"]) < TOL
    assert abs(cf["B15"].value - ground["finance"]) < TOL
    assert abs(cf["B16"].value - (ground["op_cash"] + ground["invest"]
                                  + ground["finance"])) < TOL
    assert abs(cf["B17"].value - ground["cash_open"]) < TOL
    assert abs(cf["B18"].value - ground["cash_close"]) < TOL
    assert "OK" in str(cf["B19"].value)  # ties to trial balance


def test_dashboard_kpis(recalc_wb, ground):
    d = recalc_wb["Dashboard"]
    assert abs(d["B5"].value - ground["revenue"]) < TOL
    assert abs(d["B7"].value - ground["net_income"]) < TOL
    assert abs(d["B8"].value - (ground["gross"] / ground["revenue"])) < 1e-4
    assert abs(d["B10"].value - ground["cash_close"]) < TOL


def test_no_formula_errors(recalc_wb):
    bad = []
    for ws in recalc_wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not bad, f"formula errors: {bad[:10]}"
