#!/usr/bin/env python3
"""WarehousePulse tests: data sanity, formula wiring, LO headless recalc
vs independent Python ground truth, and the AGENTS.md regression guards."""
import csv
import os
import re
import shutil
import subprocess
import sys
import unittest
from datetime import date

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "warehouse-tracker-dashboard.xlsx")
PERIOD_DAYS = 18


def read(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.DictReader(f))


def ground_truth():
    """Independent per-SKU computation straight from the CSVs."""
    skus = read("skus.csv")
    opening = {r["sku"]: int(r["opening_qty"]) for r in read("opening.csv")}
    receipts = read("receipts.csv")
    issues = read("issues.csv")
    adjustments = read("adjustments.csv")
    gt = {}
    for s in skus:
        sku = s["sku"]
        rq = sum(int(r["qty"]) for r in receipts if r["sku"] == sku)
        rv = sum(int(r["qty"]) * float(r["unit_cost"])
                 for r in receipts if r["sku"] == sku)
        iq = sum(int(r["qty"]) for r in issues if r["sku"] == sku)
        aq = sum(int(r["qty"]) for r in adjustments if r["sku"] == sku)
        oq, std = opening[sku], float(s["std_cost"])
        closing = oq + rq - iq + aq
        wac = std if (oq + rq) == 0 else (oq * std + rv) / (oq + rq)
        adi = iq / PERIOD_DAYS
        doc = 999.0 if adi == 0 else closing / adi
        rl = int(s["reorder_level"])
        status = ("STOCKOUT" if closing == 0 else
                  "REORDER" if closing <= rl else
                  "OVERSTOCK" if doc > 60 else "OK")
        gt[sku] = {"closing": closing, "wac": wac,
                   "value": closing * wac, "status": status,
                   "reorder_level": rl, "opening": oq,
                   "receipts_qty": rq, "issues_qty": iq, "adj_qty": aq,
                   "receipts_value": rv, "days_of_cover": doc}
    return gt


def recalc_copy():
    """Headless LibreOffice recalc of a COPY in a separate outdir."""
    tmp = os.path.join(BASE, "tests", "_recalc")
    if os.path.exists(tmp):
        shutil.rmtree(tmp)
    os.makedirs(tmp)
    src = os.path.join(tmp, "src.xlsx")
    shutil.copy(OUT, src)
    outdir = os.path.join(tmp, "out")
    os.makedirs(outdir)
    p = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                        "--outdir", outdir, src],
                       capture_output=True, text=True, timeout=300)
    outs = [f for f in os.listdir(outdir) if f.endswith(".xlsx")]
    assert outs, f"LO convert produced nothing: {p.stderr[-500:]}"
    return os.path.join(outdir, outs[0])


class TestData(unittest.TestCase):
    def test_csv_shapes(self):
        skus = read("skus.csv")
        self.assertEqual(len(skus), 36)
        self.assertEqual(len(read("bins.csv")), 12)
        self.assertGreater(len(read("receipts.csv")), 80)
        self.assertGreater(len(read("issues.csv")), 150)
        self.assertGreater(len(read("adjustments.csv")), 5)
        self.assertEqual(len({s["sku"] for s in skus}), 36)
        self.assertEqual(len({s["category"] for s in skus}), 5)

    def test_sku_bins_valid(self):
        bins = {b["bin"] for b in read("bins.csv")}
        for s in read("skus.csv"):
            self.assertIn(s["bin"], bins)


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import build_workbook
        build_workbook.build()
        assert os.path.exists(OUT), "workbook was not built"

    def test_sheets(self):
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        expect = {"Guide", "SKU_Master", "Raw_Receipts", "Raw_Issues",
                  "Raw_Adjustments", "Stock_Ledger", "Bin_Map", "Entry",
                  "Dashboard", "AI_Report", "Checks", "PowerBI_Model"}
        self.assertEqual(set(wb.sheetnames), expect)

    def test_no_range_literal_without_column_letter(self):
        """AGENTS.md regression: $H$2:$5000 (missing end column) is invalid."""
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        bad = re.compile(r"\$[A-Z]{1,3}\$\d+:\$\d+")
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for m in bad.finditer(c.value):
                            hits.append(f"{ws.title}!{c.coordinate}: {m.group(0)}")
        self.assertEqual(hits, [], f"bad range literals: {hits[:5]}")

    def test_no_blank_guard_in_numeric_columns(self):
        """AGENTS.md regression: '' in a numeric formula column -> #VALUE!."""
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        # numeric formula columns: ledger K(11) closing, L(12) wac, M(13) value,
        # N(14), O(15), Q(17); receipts G(7); bin C(3),E(5)
        numeric = {("Stock_Ledger", 11), ("Stock_Ledger", 12),
                   ("Stock_Ledger", 13), ("Stock_Ledger", 14),
                   ("Stock_Ledger", 15), ("Stock_Ledger", 17),
                   ("Raw_Receipts", 7), ("Bin_Map", 3), ("Bin_Map", 5)}
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if ((ws.title, c.column) in numeric
                            and isinstance(c.value, str)
                            and c.value.startswith("=")
                            and '""' in c.value):
                        hits.append(f"{ws.title}!{c.coordinate}")
        self.assertEqual(hits, [], f'"" in numeric columns: {hits[:5]}')

    def test_powerbi_export(self):
        import powerbi_export
        powerbi_export.main()
        pbi = os.path.join(BASE, "powerbi")
        for name in ("dim_sku.csv", "dim_bin.csv", "dim_date.csv",
                     "fact_inventory_daily.csv", "MEASURES.md"):
            self.assertTrue(os.path.exists(os.path.join(pbi, name)), name)
        with open(os.path.join(pbi, "fact_inventory_daily.csv")) as f:
            fact = list(csv.DictReader(f))
        self.assertEqual(len(fact), 36 * PERIOD_DAYS)
        gt = ground_truth()
        last_day = "2026-10-18"
        for row in fact:
            if row["date"] == last_day:
                self.assertEqual(int(row["closing_qty"]),
                                 gt[row["sku"]]["closing"],
                                 row["sku"])


class TestRecalc(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import build_workbook
        build_workbook.build()
        import powerbi_export
        powerbi_export.main()
        from openpyxl import load_workbook
        cls.wb = load_workbook(recalc_copy(), data_only=True)
        cls.gt = ground_truth()

    def _ledger(self):
        ws = self.wb["Stock_Ledger"]
        out = {}
        for r in range(3, 39):
            sku = ws.cell(row=r, column=1).value
            out[sku] = {
                "closing": ws.cell(row=r, column=11).value,
                "wac": ws.cell(row=r, column=12).value,
                "value": ws.cell(row=r, column=13).value,
                "status": ws.cell(row=r, column=16).value,
                "reorder_qty": ws.cell(row=r, column=17).value,
                "days_of_cover": ws.cell(row=r, column=15).value,
            }
        return out

    def test_no_formula_errors(self):
        errs = {"#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#N/A", "#NUM!",
                "#NULL!"}
        hits = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v in errs:
                        hits.append(f"{ws.title}!{c.coordinate}={v}")
        self.assertEqual(hits, [], f"formula errors: {hits[:8]}")

    def test_ledger_matches_ground_truth(self):
        led = self._ledger()
        self.assertEqual(set(led), set(self.gt))
        for sku, g in self.gt.items():
            w = led[sku]
            self.assertEqual(w["closing"], g["closing"], f"{sku} closing")
            self.assertAlmostEqual(w["wac"], g["wac"], places=2,
                                   msg=f"{sku} wac")
            self.assertAlmostEqual(w["value"], g["value"], places=2,
                                   msg=f"{sku} value")
            self.assertEqual(w["status"], g["status"], f"{sku} status")
            if g["status"] == "REORDER":
                self.assertEqual(w["reorder_qty"],
                                 max(0, 2 * g["reorder_level"] - g["closing"]),
                                 f"{sku} reorder_qty")
            else:
                self.assertEqual(w["reorder_qty"], 0, f"{sku} reorder_qty")
            self.assertAlmostEqual(w["days_of_cover"], g["days_of_cover"],
                                   places=1, msg=f"{sku} days_of_cover")

    def test_dashboard_kpis(self):
        dash = self.wb["Dashboard"]
        gt = self.gt
        self.assertEqual(dash["B3"].value, 36)
        self.assertAlmostEqual(dash["B4"].value,
                               sum(g["value"] for g in gt.values()), places=2)
        self.assertEqual(dash["B5"].value,
                         sum(1 for g in gt.values()
                             if g["status"] == "REORDER"))
        self.assertEqual(dash["B6"].value,
                         sum(1 for g in gt.values()
                             if g["status"] == "STOCKOUT"))
        self.assertEqual(dash["B7"].value,
                         sum(1 for g in gt.values()
                             if g["status"] == "OVERSTOCK"))

    def test_all_checks_true(self):
        ch = self.wb["Checks"]
        for r in range(3, 11):
            self.assertTrue(ch.cell(row=r, column=2).value,
                            f"check row {r} ({ch.cell(row=r, column=1).value}) "
                            "is not TRUE")

    def test_at_least_one_reorder_and_bins_tie(self):
        led = self._ledger()
        statuses = [v["status"] for v in led.values()]
        self.assertIn("REORDER", statuses)
        bm = self.wb["Bin_Map"]
        bin_qty = sum(bm.cell(row=r, column=3).value for r in range(3, 15))
        self.assertEqual(bin_qty, sum(v["closing"] for v in led.values()))

    def test_drillthrough_resolves(self):
        dash = self.wb["Dashboard"]
        # t0=52 -> top-5 SKU rows are 54..58, SKU in column C
        filled = sum(1 for r in range(54, 59)
                     if dash.cell(row=r, column=3).value)
        self.assertGreaterEqual(filled, 5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
