"""CollectionPulse tests: LO headless recalc vs independent Python ground truth."""
import os
import shutil
import subprocess
import sys
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from main import ground_truth, AS_OF  # noqa: E402

ROOT = os.path.join(os.path.dirname(__file__), "..")
XLSX = os.path.join(ROOT, "payment-collections-tracker.xlsx")
RECALC_DIR = os.path.join(ROOT, "recalc-out")
ERRORS = ("#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#REF!", "#NUM!", "Err:509",
          "Err:508", "Err:502")


@pytest.fixture(scope="module")
def recalc():
    if os.path.isdir(RECALC_DIR):
        shutil.rmtree(RECALC_DIR)
    os.makedirs(RECALC_DIR, exist_ok=True)
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
         RECALC_DIR, XLSX],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stderr[-2000:]
    out = os.path.join(RECALC_DIR, "payment-collections-tracker.xlsx")
    assert os.path.exists(out)
    return load_workbook(out, data_only=True)


@pytest.fixture(scope="module")
def gt():
    return ground_truth()


def val(ws, ref):
    return ws[ref].value


def as_date(v):
    return v.date() if isinstance(v, datetime) else v


def test_files_exist():
    for p in ["payment-collections-tracker.xlsx", "README.md", "generate.py",
              "build_workbook.py", "build_part2.py", "main.py",
              "data/schedule.csv", "data/payments.csv",
              "tests/test_build.py"]:
        assert os.path.exists(os.path.join(ROOT, p)), p


def test_no_assistant_mentions():
    for p in ["README.md", "generate.py", "build_workbook.py",
              "build_part2.py", "main.py"]:
        text = open(os.path.join(ROOT, p)).read()
        assert "Muse" not in text and "anthropic" not in text.lower(), p


def test_asof(recalc):
    assert as_date(val(recalc["Guide"], "B3")) == AS_OF


def test_dashboard_kpis(recalc, gt):
    ws = recalc["Dashboard"]
    tot_due = sum(g["due_amt"] for g in gt)
    tot_paid = sum(g["paid"] for g in gt)
    od = [g for g in gt if g["out"] > 0.005]
    assert val(ws, "B3") == pytest.approx(tot_due, abs=0.01)
    assert val(ws, "B4") == pytest.approx(tot_paid, abs=0.01)
    assert val(ws, "B5") == pytest.approx(tot_paid / tot_due, abs=1e-9)
    assert val(ws, "B6") == pytest.approx(sum(g["out"] for g in od), abs=0.01)
    assert val(ws, "B7") == len(od)
    assert val(ws, "B8") == pytest.approx(sum(g["dpd"] for g in od) / len(od), abs=0.05)
    assert val(ws, "B9") in {g["collector"] for g in gt}
    assert val(ws, "B10") in {g["collector"] for g in gt}


def test_aging(recalc, gt):
    ws = recalc["Aging"]
    for j, bucket in enumerate(["COLLECTED", "1-30", "31-60", "61-90", "90+"], start=2):
        g = [x for x in gt if x["bucket"] == bucket]
        assert val(ws, f"B{j}") == len(g), bucket
        assert val(ws, f"C{j}") == pytest.approx(sum(x["due_amt"] for x in g), abs=0.01), bucket
        assert val(ws, f"D{j}") == pytest.approx(sum(x["paid"] for x in g), abs=0.01), bucket
        assert val(ws, f"E{j}") == pytest.approx(sum(x["out"] for x in g), abs=0.01), bucket
    assert val(ws, "B7") == len(gt)
    assert val(ws, "C7") == pytest.approx(sum(g["due_amt"] for g in gt), abs=0.01)


def test_scorecard(recalc, gt):
    ws = recalc["Collector_Scorecard"]
    names = sorted({g["collector"] for g in gt})
    for j, name in enumerate(names, start=2):
        assert val(ws, f"A{j}") == name
        g = [x for x in gt if x["collector"] == name]
        od = [x for x in g if x["out"] > 0.005]
        assert val(ws, f"B{j}") == len(g)
        assert val(ws, f"C{j}") == pytest.approx(sum(x["due_amt"] for x in g), abs=0.01)
        assert val(ws, f"D{j}") == pytest.approx(sum(x["paid"] for x in g), abs=0.01)
        assert val(ws, f"F{j}") == pytest.approx(sum(x["out"] for x in od), abs=0.01)
        if od:
            assert val(ws, f"G{j}") == pytest.approx(sum(x["dpd"] for x in od) / len(od), abs=0.05)


def test_ledger_full(recalc, gt):
    ws = recalc["Receivables_Ledger"]
    for i, g in enumerate(gt, start=2):
        assert val(ws, f"A{i}") == g["due"]
        assert val(ws, f"I{i}") == pytest.approx(g["paid"], abs=0.01), f"row {i}"
        assert val(ws, f"J{i}") == pytest.approx(g["out"], abs=0.01), f"row {i}"
        assert val(ws, f"K{i}") == g["status"], f"row {i}"
        assert val(ws, f"M{i}") == pytest.approx(g["dpd"], abs=0.01), f"row {i}"
        assert val(ws, f"N{i}") == g["bucket"], f"row {i}"
        assert val(ws, f"J{i}") >= -0.01, f"negative outstanding row {i}"
    # LastPayDate spot-checks: PAID dues carry the true max pay date
    paid = [g for g in gt if g["status"] == "PAID"][:5]
    pays = {}
    import csv
    for p in csv.DictReader(open(os.path.join(ROOT, "data/payments.csv"))):
        if p["PayDate"] > pays.get(p["DueID"], ""):
            pays[p["DueID"]] = p["PayDate"]
    for g in paid:
        v = as_date(val(ws, f"L{gt.index(g)+2}"))
        assert v == date.fromisoformat(pays[g["due"]]), g["due"]


def test_followup(recalc, gt):
    ws = recalc["FollowUp"]
    od = [g for g in gt if g["out"] > 0.005]
    n = ws.max_row - 1
    assert n == len(od), f"{n} != {len(od)}"
    seen = set()
    for i in range(2, ws.max_row + 1):
        due = val(ws, f"B{i}")
        seen.add(due)
        g = next(x for x in od if x["due"] == due)
        assert val(ws, f"E{i}") == pytest.approx(g["dpd"], abs=0.01)
        assert val(ws, f"F{i}") == pytest.approx(g["out"], abs=0.01)
        sev = "LOST" if g["dpd"] > 90 else "ESCALATE" if g["dpd"] > 60 else "WATCH" if g["dpd"] > 30 else "MILD"
        assert val(ws, f"G{i}") == sev
        assert as_date(val(ws, f"H{i}")) >= AS_OF
    assert seen == {g["due"] for g in od}


def test_checks_all_true(recalc):
    ws = recalc["Checks"]
    for i in range(2, 10):
        assert val(ws, f"B{i}") is True, f"check row {i}: {val(ws, f'A{i}')}"


def test_no_errors(recalc):
    bad = []
    for ws in recalc.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and (v in ERRORS or v.startswith("Err:")):
                    bad.append(f"{ws.title}!{c.coordinate}={v}")
    assert not bad, bad[:10]


def test_charts_present():
    wb = load_workbook(XLSX)
    assert len(wb["Dashboard"]._charts) >= 3
