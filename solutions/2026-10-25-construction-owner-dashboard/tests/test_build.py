#!/usr/bin/env python3
"""Verify the Owner Dashboard workbook against an independent Python
ground truth computed straight from data/*.csv, plus a LibreOffice headless
recalc pass (recalc in a SEPARATE indir/outdir — soffice keeps the input
basename and refuses to overwrite in place).

Checks: all 6 dashboard KPIs, gross margin, AR/AP bucket ties, all 8 Checks
TRUE, zero formula-error strings anywhere, 3 native charts on Dashboard,
and that key cells actually contain formulas (not pasted values).
"""
import csv
import datetime
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(HERE, "data")
XLSX = os.path.join(HERE, "qbo-owner-dashboard.xlsx")
AS_OF = datetime.date(2026, 10, 25)
ERRORS = {"#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#N/A", "#NUM!"}


def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.DictReader(f))


def ground_truth():
    qbo = read_csv("qbo_pl.csv")
    ar = read_csv("ar.csv")
    ap = read_csv("ap.csv")
    opening = int(read_csv("cash.csv")[0]["Amount"])
    rev = sum(int(r["Amount"]) for r in qbo
            if r["Category"].startswith("Revenue"))
    costs = sum(int(r["Amount"]) for r in qbo
                if not r["Category"].startswith("Revenue"))
    gross = rev + costs
    ar_out = sum(int(r["Amount"]) - int(r["PaidAmount"]) for r in ar)
    ar_paid = sum(int(r["PaidAmount"]) for r in ar)
    ap_out = sum(int(r["Amount"]) - int(r["PaidAmount"]) for r in ap)
    ap_paid = sum(int(r["PaidAmount"]) for r in ap)

    def bucket(due):
        d = datetime.date(*[int(x) for x in due.split("-")])
        if d >= AS_OF:
            return "Current"
        delta = (AS_OF - d).days
        if delta <= 30:
            return "1-30"
        if delta <= 60:
            return "31-60"
        return "61+"

    ar61 = sum(int(r["Amount"]) - int(r["PaidAmount"]) for r in ar
               if bucket(r["DueDate"]) == "61+"
               and int(r["Amount"]) - int(r["PaidAmount"]) > 0)
    cash = opening + ar_paid - ap_paid
    return {
        "revenue": rev, "costs": costs, "gross": gross,
        "margin": gross / rev if rev else 0,
        "ar_out": ar_out, "ap_out": ap_out, "ar61": ar61, "cash": cash,
    }


def recalc_copy():
    """LibreOffice headless recalc into a fresh outdir; return workbook."""
    tmp = tempfile.mkdtemp()
    indir, outdir = os.path.join(tmp, "in"), os.path.join(tmp, "out")
    os.makedirs(indir)
    os.makedirs(outdir)
    shutil.copy(XLSX, os.path.join(indir, "qbo-owner-dashboard.xlsx"))
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         "--outdir", outdir, os.path.join(indir, "qbo-owner-dashboard.xlsx")],
        capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, r.stderr[-500:]
    recalculated = os.path.join(outdir, "qbo-owner-dashboard.xlsx")
    assert os.path.exists(recalculated), "LO produced no output"
    return load_workbook(recalculated, data_only=True)


class TestDashboard(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.gt = ground_truth()
        cls.wb = recalc_copy()
        cls.formulas = load_workbook(XLSX)  # formulas, not values

    def test_kpis(self):
        d = self.wb["Dashboard"]
        self.assertEqual(d["B3"].value, self.gt["revenue"])
        self.assertAlmostEqual(d["B4"].value, self.gt["margin"], places=6)
        self.assertEqual(d["B5"].value, self.gt["cash"])
        self.assertEqual(d["B6"].value, self.gt["ar_out"])
        self.assertEqual(d["B7"].value, self.gt["ap_out"])
        self.assertEqual(d["B8"].value, self.gt["ar61"])

    def test_pl_totals(self):
        pl = self.wb["Monthly_PL"]
        self.assertEqual(pl["H11"].value, self.gt["revenue"])
        self.assertEqual(pl["H12"].value, self.gt["costs"])
        self.assertEqual(pl["H13"].value, self.gt["gross"])
        self.assertAlmostEqual(pl["H14"].value, self.gt["margin"], places=6)

    def test_checks_all_true(self):
        ch = self.wb["Checks"]
        for r in range(3, 11):
            self.assertIs(ch.cell(row=r, column=2).value, True,
                          f"check row {r} failed")

    def test_no_formula_errors(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertNotIn(cell.value, ERRORS,
                                         f"{ws.title}!{cell.coordinate}")

    def test_charts_and_formulas(self):
        dash = self.formulas["Dashboard"]
        self.assertEqual(len(dash._charts), 3)
        self.assertTrue(str(self.formulas["Monthly_PL"]["B3"].value)
                        .startswith("=SUMIFS"))
        self.assertTrue(str(self.formulas["AR_Aging"]["G3"].value)
                        .startswith("=E3-F3"))
        self.assertTrue(str(dash["B3"].value).startswith("=Monthly_PL"))

    def test_bucket_sums(self):
        ar = self.wb["AR_Aging"]
        self.assertEqual(ar["B39"].value, self.gt["ar_out"])
        ap = self.wb["AP_Aging"]
        self.assertEqual(ap["B33"].value, self.gt["ap_out"])


if __name__ == "__main__":
    sys.exit(0 if unittest.main(exit=False).result.wasSuccessful() else 1)
