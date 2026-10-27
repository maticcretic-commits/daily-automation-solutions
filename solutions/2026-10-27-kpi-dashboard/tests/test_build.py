#!/usr/bin/env python3
"""Build the workbook, recalc it headlessly in LibreOffice, and compare
every checked figure against the independent Python ground truth.

LO quirk honoured: --convert-to runs on a COPY in a separate input dir
with --outdir elsewhere (converting in place overwrites the input file).
Formulas are A1-style only; a formula error in the recalculated copy is a
failure, not a warning.
"""
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import generate  # noqa: E402
import build_workbook  # noqa: E402
import main as ground_truth  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, "kpi-dashboard.xlsx")
ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "Err:522")


def close(a, b, tol):
    return abs(a - b) <= tol


def main():
    generate.main()
    build_workbook.main()
    truth = ground_truth.compute()

    tmp = tempfile.mkdtemp(prefix="kpipulse-")
    indir, outdir = os.path.join(tmp, "in"), os.path.join(tmp, "out")
    os.makedirs(indir)
    os.makedirs(outdir)
    shutil.copy(XLSX, os.path.join(indir, "kpi-dashboard.xlsx"))
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir,
         os.path.join(indir, "kpi-dashboard.xlsx")],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:500]}"
    recalc = os.path.join(outdir, "kpi-dashboard.xlsx")
    assert os.path.exists(recalc), "LO produced no output"
    wb = load_workbook(recalc, data_only=True)

    failures = []

    def check(name, got, want, tol):
        if got is None or isinstance(got, str) or not close(got, want, tol):
            failures.append(f"{name}: got {got!r}, want {want}")

    # No formula errors anywhere on the checked sheets.
    for sheet in ("Dashboard", "Monthly", "Channel", "Checks"):
        for row in wb[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    failures.append(f"formula error {sheet}!{c.coordinate}: {c.value}")

    # Monthly sheet: rows 2-10, months Jan-Sep 2026.
    ws = wb["Monthly"]
    months = sorted(truth["months"])
    for i, month in enumerate(months):
        r = i + 2
        t = truth["months"][month]
        assert ws.cell(row=r, column=1).value == month, f"month row {r}"
        check(f"{month} revenue", ws.cell(row=r, column=2).value, t["revenue"], 1.0)
        check(f"{month} variance%", ws.cell(row=r, column=5).value, t["variance_pct"], 1e-3)
        check(f"{month} orders", ws.cell(row=r, column=6).value, t["orders"], 0.01)
        check(f"{month} AOV", ws.cell(row=r, column=7).value, t["aov"], 0.01)
        check(f"{month} margin%", ws.cell(row=r, column=10).value, t["margin_pct"], 1e-3)
        check(f"{month} ontime%", ws.cell(row=r, column=11).value, t["ontime_pct"], 1e-3)
        check(f"{month} CSAT", ws.cell(row=r, column=12).value, t["csat"], 1e-3)

    # Channel sheet: rows 2-4 in alphabetical order (Online, Store, Wholesale).
    ws = wb["Channel"]
    for i, ch in enumerate(["Online", "Store", "Wholesale"]):
        r = i + 2
        t = truth["channels"][ch]
        assert ws.cell(row=r, column=1).value == ch
        check(f"{ch} revenue", ws.cell(row=r, column=2).value, t["revenue"], 1.0)
        check(f"{ch} share", ws.cell(row=r, column=7).value, t["share"], 1e-3)
    check("channel share total", ws.cell(row=5, column=7).value, 1.0, 1e-3)

    # Dashboard cards B4:B9.
    ws = wb["Dashboard"]
    d = truth["dashboard"]
    check("dash latest revenue", ws["B4"].value, d["latest_revenue"], 1.0)
    check("dash YTD revenue", ws["B5"].value, d["ytd_revenue"], 1.0)
    check("dash YTD margin", ws["B6"].value, d["ytd_margin"], 1e-3)
    check("dash ontime", ws["B7"].value, d["latest_ontime"], 1e-3)
    check("dash CSAT", ws["B8"].value, d["latest_csat"], 1e-3)
    check("dash attainment", ws["B9"].value, d["target_attainment"], 1e-3)

    # Checks: all 8 must read TRUE.
    ws = wb["Checks"]
    for r in range(2, 10):
        if ws.cell(row=r, column=2).value is not True:
            failures.append(f"check row {r} ({ws.cell(row=r, column=1).value}): "
                            f"{ws.cell(row=r, column=2).value!r}")

    shutil.rmtree(tmp, ignore_errors=True)
    if failures:
        print(f"FAIL ({len(failures)}):")
        for f in failures:
            print("  -", f)
        return 1
    print("ALL PASS — monthly (9), channel (3), dashboard (6), checks (8), "
          "zero formula errors")
    print(f"YTD revenue {d['ytd_revenue']:,} · YTD margin {d['ytd_margin']:.1%} · "
          f"Sep revenue {d['latest_revenue']:,} · CSAT {d['latest_csat']:.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
