"""Tests for the SalesPulse MIS Factory.

Covers: file generation, workbook structure (sheets/tables/charts), formula
presence, and — via a headless LibreOffice recalculation — that every KPI
matches an independent Python ground truth computed from the source CSVs,
with zero formula errors and the injected anomalies flagged.
"""
import csv
import datetime as dt
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
import generate as G  # noqa: E402

XLSX = BASE / "sales-purchase-mis.xlsx"
RECALC = BASE / "tests" / "_recalc.xlsx"

EXPECTED_SHEETS = {"Dashboard", "Sales", "Purchases", "Products", "Monthly",
                   "Daily", "Checks", "Guide"}


@pytest.fixture(scope="module")
def built():
    G.main()
    assert XLSX.exists()
    return XLSX


@pytest.fixture(scope="module")
def wb(built):
    return load_workbook(built)


@pytest.fixture(scope="module")
def ground_truth(built):
    """Independent recomputation straight from the source CSVs."""
    sales, purchases = [], []
    for name in ("source_sales_apr_jun.csv", "source_sales_jul_sep.csv"):
        with open(BASE / "data" / name) as f:
            for r in csv.DictReader(f):
                r["Qty"] = int(r["Qty"]); r["Rate"] = float(r["Rate"])
                r["Date"] = dt.date.fromisoformat(r["Date"])
                sales.append(r)
    with open(BASE / "data" / "source_purchases.csv") as f:
        for r in csv.DictReader(f):
            r["Qty"] = int(r["Qty"]); r["CostRate"] = float(r["CostRate"])
            purchases.append(r)
    costs = {}
    for p in purchases:
        costs.setdefault(p["Product"], []).append(p["CostRate"])
    stdcost = {k: sum(v) / len(v) for k, v in costs.items()}
    total_sales = sum(s["Qty"] * s["Rate"] for s in sales)
    total_pur = sum(p["Qty"] * p["CostRate"] for p in purchases)
    cogs = sum(s["Qty"] * stdcost[s["Product"]] for s in sales)
    return {
        "n_sales": len(sales), "n_pur": len(purchases),
        "total_sales": total_sales, "total_pur": total_pur,
        "gross_profit": total_sales - cogs,
        "margin": (total_sales - cogs) / total_sales,
    }


@pytest.fixture(scope="module")
def recalc(built):
    RECALC.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", str(RECALC.parent), str(built)],
                   check=True, capture_output=True, timeout=300)
    produced = RECALC.parent / "sales-purchase-mis.xlsx"
    if RECALC.exists():
        RECALC.unlink()
    produced.rename(RECALC)
    return load_workbook(RECALC, data_only=True)


def test_no_malformed_ranges(wb):
    """Regression: every range must repeat the column letter on both ends
    ('$H$2:$H$5000', never '$H$2:$5000') — the latter is not a valid
    reference and breaks recalculation."""
    import re
    bad = re.compile(r"\$[A-Z]{1,3}\$\d+:\$\d+")
    offenders = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and bad.search(v):
                    offenders.append(f"{ws.title}!{c.coordinate}: {v[:80]}")
    assert not offenders, offenders[:5]


def test_csvs_exist(built):
    for n in ("source_sales_apr_jun.csv", "source_sales_jul_sep.csv",
              "source_purchases.csv"):
        assert (BASE / "data" / n).exists()


def test_sheets_present(wb):
    assert EXPECTED_SHEETS <= set(wb.sheetnames)


def test_headers_present(wb):
    assert [c.value for c in wb["Sales"][1]] == [
        "Date", "Invoice", "Customer", "Product", "Category", "Qty",
        "Rate", "Amount", "COGS"]
    assert [c.value for c in wb["Purchases"][1]] == [
        "Date", "Bill", "Supplier", "Product", "Category", "Qty",
        "CostRate", "Amount"]


def test_amount_columns_are_formulas(wb):
    for cell in wb["Sales"]["H"][1:]:
        assert isinstance(cell.value, str) and cell.value.startswith("="), \
            f"Sales!H not a formula at row {cell.row}"
    for cell in wb["Purchases"]["H"][1:]:
        assert isinstance(cell.value, str) and cell.value.startswith("=")


def test_row_counts_match_csvs(wb, ground_truth):
    n_sales = sum(1 for c in wb["Sales"]["B"][1:] if c.value)
    n_pur = sum(1 for c in wb["Purchases"]["B"][1:] if c.value)
    assert n_sales == ground_truth["n_sales"]
    assert n_pur == ground_truth["n_pur"]


def test_charts_present(wb):
    assert len(wb["Dashboard"]._charts) == 3


def _kpi(recalc, row, col):
    return recalc["Dashboard"].cell(row=row, column=col).value


def test_recalc_total_sales(recalc, ground_truth):
    assert _kpi(recalc, 6, 1) == pytest.approx(ground_truth["total_sales"],
                                              abs=0.01)


def test_recalc_total_purchases(recalc, ground_truth):
    assert _kpi(recalc, 6, 3) == pytest.approx(ground_truth["total_pur"],
                                              abs=0.01)


def test_recalc_gross_profit(recalc, ground_truth):
    assert _kpi(recalc, 6, 5) == pytest.approx(ground_truth["gross_profit"],
                                              abs=0.05)


def test_recalc_margin(recalc, ground_truth):
    assert _kpi(recalc, 6, 7) == pytest.approx(ground_truth["margin"],
                                              abs=1e-4)


def test_no_formula_errors(recalc):
    bad = 0
    for ws in ("Dashboard", "Monthly", "Daily", "Checks", "Products"):
        for row in recalc[ws].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad += 1
    assert bad == 0, f"{bad} formula errors after recalc"


def test_anomaly_days_flagged(recalc):
    flags = {}
    for row in recalc["Daily"].iter_rows(min_row=2, values_only=True):
        d = row[0]
        if isinstance(d, dt.datetime):
            d = d.date()
        flags[d] = row[6]
    assert flags.get(dt.date(2026, 5, 18)) == "ANOMALY"   # bulk-order spike
    assert flags.get(dt.date(2026, 8, 2)) == "ANOMALY"     # near-zero day


def test_checks_all_ok(recalc):
    statuses = [recalc["Checks"].cell(row=r, column=3).value
                for r in range(3, 9)]
    assert statuses == ["OK"] * 6


def test_monthly_reconciles(recalc, ground_truth):
    m = recalc["Monthly"]
    sales_m = sum(m.cell(row=r, column=2).value for r in range(2, 8))
    assert sales_m == pytest.approx(ground_truth["total_sales"], abs=0.05)
