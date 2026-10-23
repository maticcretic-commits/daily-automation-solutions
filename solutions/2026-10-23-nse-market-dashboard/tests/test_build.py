#!/usr/bin/env python3
"""NSEPulse tests: build + LibreOffice headless recalc vs Python ground truth."""
import csv
import os
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
import generate  # noqa: E402

XLSX = os.path.join(BASE, "nsepulse.xlsx")
N = 60


def load_rows(name):
    with open(os.path.join(BASE, "data", name), newline="") as f:
        return list(csv.DictReader(f))


def ground_truth():
    """Independent Python computation straight from the CSVs."""
    t1 = {r["Ticker"]: r for r in load_rows("snapshots_t1.csv")}
    rows = []
    for r in load_rows("snapshots_t.csv"):
        p = t1[r["Ticker"]]
        ltp, ltp_prev = float(r["LTP"]), float(p["LTP"])
        oi, oi_prev = int(r["OI"]), int(p["OI"])
        pct = (ltp - ltp_prev) / ltp_prev
        oi_chg = oi - oi_prev
        oi_chg_pct = oi_chg / oi_prev
        vol_mult = int(r["Volume"]) / int(p["Vol20d_Avg"])
        if pct > 0 and oi_chg > 0:
            sig = "Long Buildup"
        elif pct < 0 and oi_chg > 0:
            sig = "Short Buildup"
        elif pct < 0 and oi_chg < 0:
            sig = "Long Unwinding"
        elif pct > 0 and oi_chg < 0:
            sig = "Short Covering"
        else:
            sig = "Neutral"
        rows.append({"ticker": r["Ticker"], "sector": r["Sector"], "pct": pct,
                     "oi_chg": oi_chg, "oi_chg_pct": oi_chg_pct,
                     "vol_mult": vol_mult, "signal": sig,
                     "volume": int(r["Volume"])})
    return rows


@pytest.fixture(scope="session")
def recalc(tmp_path_factory):
    """LibreOffice-recalculated workbook (data_only), computed once."""
    base = tmp_path_factory.mktemp("lo")
    indir, outdir = base / "in", base / "out"
    indir.mkdir()
    outdir.mkdir()
    src = indir / "nsepulse.xlsx"
    shutil.copy(XLSX, src)
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         "--outdir", str(outdir), str(src)],
        capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:500]}"
    return load_workbook(outdir / "nsepulse.xlsx", data_only=True)


def test_data_deterministic(tmp_path):
    a1 = [dict(r) for r in load_rows("snapshots_t.csv")]
    generate.main()
    a2 = [dict(r) for r in load_rows("snapshots_t.csv")]
    assert a1 == a2 and len(a1) == N


def test_workbook_builds():
    wb = load_workbook(XLSX)
    for s in ["Guide", "Dashboard", "Raw_Today", "Raw_PrevDay",
              "Top_Movers", "OI_Signals", "Sector_Heatmap", "Checks"]:
        assert s in wb.sheetnames, s
    assert wb["Raw_Today"].max_row == N + 1


def _cell(recalc, sheet, ref):
    return recalc[sheet][ref].value


def test_lo_top_gainer(recalc):
    gt = max(ground_truth(), key=lambda r: r["pct"])
    assert _cell(recalc, "Dashboard", "B4") == gt["ticker"]
    assert _cell(recalc, "Dashboard", "C4") == pytest.approx(gt["pct"], rel=1e-9)


def test_lo_top_loser(recalc):
    gt = min(ground_truth(), key=lambda r: r["pct"])
    assert _cell(recalc, "Dashboard", "B6") == gt["ticker"]
    assert _cell(recalc, "Dashboard", "C6") == pytest.approx(gt["pct"], rel=1e-9)


def test_lo_long_buildup(recalc):
    lbs = [r for r in ground_truth() if r["signal"] == "Long Buildup"]
    gt = max(lbs, key=lambda r: r["oi_chg_pct"])
    assert _cell(recalc, "Dashboard", "B8") == gt["ticker"]
    assert _cell(recalc, "Dashboard", "C8") == pytest.approx(gt["oi_chg_pct"], rel=1e-9)


def test_lo_volume_shocker(recalc):
    gt = max(ground_truth(), key=lambda r: r["vol_mult"])
    assert _cell(recalc, "Dashboard", "B10") == gt["ticker"]
    assert _cell(recalc, "Dashboard", "C10") == pytest.approx(gt["vol_mult"], rel=1e-9)


def test_lo_breadth(recalc):
    gt = ground_truth()
    adv = sum(1 for r in gt if r["pct"] > 0)
    dec = sum(1 for r in gt if r["pct"] < 0)
    assert _cell(recalc, "Dashboard", "B12") == f"{adv} / {dec}"


def test_lo_gainers_table(recalc):
    gt = sorted(ground_truth(), key=lambda r: r["pct"], reverse=True)[:10]
    tm = recalc["Top_Movers"]
    got = {(tm.cell(row=r, column=2).value,
            round(tm.cell(row=r, column=3).value, 9)) for r in range(3, 13)}
    exp = {(r["ticker"], round(r["pct"], 9)) for r in gt}
    assert got == exp


def test_lo_sector_volume_ties(recalc):
    gt = ground_truth()
    sh = recalc["Sector_Heatmap"]
    total = sum(sh.cell(row=r, column=5).value for r in range(2, 10))
    assert total == sum(r["volume"] for r in gt)


def test_lo_all_checks_true(recalc):
    ck = recalc["Checks"]
    vals = [ck.cell(row=r, column=2).value for r in range(2, 10)]
    assert vals == [True] * 8, vals


def test_lo_no_formula_errors(recalc):
    bad = []
    for ws in recalc.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not bad, bad[:10]
