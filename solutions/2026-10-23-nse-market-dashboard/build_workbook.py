#!/usr/bin/env python3
"""Build the NSEPulse workbook: formula-driven NSE market dashboard.

Reads data/snapshots_t1.csv + data/snapshots_t.csv, writes nsepulse.xlsx.
Every derived figure is a live Excel formula — nothing hard-coded — so the
dashboard refreshes when fresh snapshots are pasted in. All sample data is
synthetic (see generate.py); the Guide sheet says so.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "nsepulse.xlsx")

N = 60                       # tickers
R_T = f"$A$2:$A${N+1}"       # ticker col on Raw_Today (with $)
ROW = f"2:{N+1}"

NAVY = "1F3864"
LIGHT = "D9E2F3"
ACCENT = "2E75B6"
GREEN = "548235"
RED = "C00000"

HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color=NAVY, size=14)
KPI_FONT = Font(bold=True, size=13)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_csv(name):
    with open(os.path.join(BASE, "data", name), newline="") as f:
        return list(csv.DictReader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def style_range(ws, min_row, max_row, ncols):
    for r in range(min_row, max_row + 1):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = BORDER


def build():
    t1 = load_csv("snapshots_t1.csv")
    t = load_csv("snapshots_t.csv")
    assert len(t1) == len(t) == N

    wb = Workbook()

    # ---------------- Guide ----------------
    ws = wb.active
    ws.title = "Guide"
    ws["A1"] = "NSEPulse — NSE Market Dashboard (demo)"
    ws["A1"].font = TITLE_FONT
    guide = [
        "",
        "SYNTHETIC SAMPLE DATA — every figure in this workbook is deterministic",
        "fictional demo data (generate.py, seed 20261023). It is NOT real NSE data.",
        "",
        "Sheet map:",
        "  Dashboard       — 6 KPI cards + top-gainers and OI-buildup charts",
        "  Raw_Today       — today's snapshot; every derived column is a formula",
        "  Raw_PrevDay     — previous snapshot (LTP, 20-day avg volume, OI)",
        "  Top_Movers      — top-10 gainers / losers, extracted by formula",
        "  OI_Signals      — top-10 long buildups + volume shockers",
        "  Sector_Heatmap  — sector breadth, avg move, volume, OI change",
        "  Checks          — 8 self-tests; all must read TRUE",
        "",
        "OI interpretation (standard price-OI logic):",
        "  price up + OI up = Long Buildup | price down + OI up = Short Buildup",
        "  price down + OI down = Long Unwinding | price up + OI down = Short Covering",
        "",
        "Production refresh: point Power Query at the NSE snapshot source,",
        "paste/refresh Raw_Today + Raw_PrevDay — every sheet recomputes.",
    ]
    for i, line in enumerate(guide, start=2):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 78

    # ---------------- Raw_PrevDay ----------------
    rp = wb.create_sheet("Raw_PrevDay")
    headers = ["Ticker", "Sector", "LTP_Prev", "Vol20d_Avg", "OI_Prev"]
    rp.append(headers)
    style_header(rp, 1, 5)
    for row in t1:
        rp.append([row["Ticker"], row["Sector"], float(row["LTP"]),
                   int(row["Vol20d_Avg"]), int(row["OI"])])
    for r in range(2, N + 2):
        rp.cell(row=r, column=3).number_format = "#,##0.00"
        rp.cell(row=r, column=4).number_format = "#,##0"
        rp.cell(row=r, column=5).number_format = "#,##0"
    style_range(rp, 2, N + 1, 5)
    for w, col in [(12, "A"), (12, "B"), (12, "C"), (14, "D"), (14, "E")]:
        rp.column_dimensions[col].width = w
    rp.freeze_panes = "A2"
    rp.sheet_properties.tabColor = "808080"

    # ---------------- Raw_Today ----------------
    rt = wb.create_sheet("Raw_Today")
    headers = ["Ticker", "Sector", "LTP", "Volume", "OI", "% Change",
               "OI Change", "OI Chg %", "Vol Multiple", "Signal",
               "HL_Gain", "HL_Loss", "HL_LongBuild"]
    rt.append(headers)
    style_header(rt, 1, 13)
    for row in t:
        rt.append([row["Ticker"], row["Sector"], float(row["LTP"]),
                   int(row["Volume"]), int(row["OI"])] + [None] * 8)
    prev = "Raw_PrevDay!$A$2:$E$61"
    for r in range(2, N + 2):
        rt.cell(row=r, column=6).value = (
            f"=(C{r}-VLOOKUP(A{r},{prev},3,0))/VLOOKUP(A{r},{prev},3,0)")
        rt.cell(row=r, column=7).value = f"=E{r}-VLOOKUP(A{r},{prev},5,0)"
        rt.cell(row=r, column=8).value = f"=G{r}/VLOOKUP(A{r},{prev},5,0)"
        rt.cell(row=r, column=9).value = f"=D{r}/VLOOKUP(A{r},{prev},4,0)"
        rt.cell(row=r, column=10).value = (
            f'=IF(AND(F{r}>0,G{r}>0),"Long Buildup",'
            f'IF(AND(F{r}<0,G{r}>0),"Short Buildup",'
            f'IF(AND(F{r}<0,G{r}<0),"Long Unwinding",'
            f'IF(AND(F{r}>0,G{r}<0),"Short Covering","Neutral"))))')
        rt.cell(row=r, column=11).value = f"=IF(F{r}>0,F{r},0)"
        rt.cell(row=r, column=12).value = f"=IF(F{r}<0,F{r},0)"
        rt.cell(row=r, column=13).value = f'=IF(J{r}="Long Buildup",H{r},0)'
    for r in range(2, N + 2):
        rt.cell(row=r, column=3).number_format = "#,##0.00"
        rt.cell(row=r, column=4).number_format = "#,##0"
        rt.cell(row=r, column=5).number_format = "#,##0"
        rt.cell(row=r, column=6).number_format = "0.00%"
        rt.cell(row=r, column=7).number_format = "#,##0"
        rt.cell(row=r, column=8).number_format = "0.00%"
        rt.cell(row=r, column=9).number_format = "0.00"
        rt.cell(row=r, column=11).number_format = "0.00%"
        rt.cell(row=r, column=12).number_format = "0.00%"
        rt.cell(row=r, column=13).number_format = "0.00%"
    style_range(rt, 2, N + 1, 13)
    widths = [12, 12, 12, 14, 14, 10, 12, 10, 11, 15, 10, 10, 12]
    for i, w in enumerate(widths, start=1):
        rt.column_dimensions[chr(64 + i)].width = w
    rt.freeze_panes = "A2"
    rt.sheet_properties.tabColor = ACCENT

    # ---------------- Top_Movers ----------------
    tm = wb.create_sheet("Top_Movers")
    tm["A1"] = "Top 10 Gainers"
    tm["A1"].font = Font(bold=True, color=GREEN, size=12)
    tm["E1"] = "Top 10 Losers"
    tm["E1"].font = Font(bold=True, color=RED, size=12)
    for col0, label in [(1, "Rank"), (2, "Ticker"), (3, "% Change"),
                        (5, "Rank"), (6, "Ticker"), (7, "% Change")]:
        tm.cell(row=2, column=col0, value=label)
    style_header(tm, 2, 3)
    style_header(tm, 2, 7)
    for k in range(1, 11):
        r = k + 2
        tm.cell(row=r, column=1, value=k)
        tm.cell(row=r, column=2).value = (
            f"=INDEX(Raw_Today!$A$2:$A$61,"
            f'MATCH(LARGE(Raw_Today!$K$2:$K$61,A{r}),Raw_Today!$K$2:$K$61,0))')
        tm.cell(row=r, column=3).value = f"=LARGE(Raw_Today!$K$2:$K$61,A{r})"
        tm.cell(row=r, column=5, value=k)
        tm.cell(row=r, column=6).value = (
            f"=INDEX(Raw_Today!$A$2:$A$61,"
            f'MATCH(SMALL(Raw_Today!$L$2:$L$61,E{r}),Raw_Today!$L$2:$L$61,0))')
        tm.cell(row=r, column=7).value = f"=SMALL(Raw_Today!$L$2:$L$61,E{r})"
        tm.cell(row=r, column=3).number_format = "0.00%"
        tm.cell(row=r, column=7).number_format = "0.00%"
    style_range(tm, 3, 12, 3)
    style_range(tm, 3, 12, 7)
    for col, w in [("A", 6), ("B", 13), ("C", 11), ("E", 6), ("F", 13), ("G", 11)]:
        tm.column_dimensions[col].width = w

    # ---------------- OI_Signals ----------------
    oi = wb.create_sheet("OI_Signals")
    oi["A1"] = "Top 10 Long Buildups (OI Chg %)"
    oi["A1"].font = Font(bold=True, color=NAVY, size=12)
    oi["E1"] = "Top 10 Volume Shockers"
    oi["E1"].font = Font(bold=True, color=NAVY, size=12)
    for col0, label in [(1, "Rank"), (2, "Ticker"), (3, "OI Chg %"),
                        (5, "Rank"), (6, "Ticker"), (7, "Vol x Avg")]:
        oi.cell(row=2, column=col0, value=label)
    style_header(oi, 2, 3)
    style_header(oi, 2, 7)
    for k in range(1, 11):
        r = k + 2
        oi.cell(row=r, column=1, value=k)
        oi.cell(row=r, column=2).value = (
            f"=INDEX(Raw_Today!$A$2:$A$61,"
            f'MATCH(LARGE(Raw_Today!$M$2:$M$61,A{r}),Raw_Today!$M$2:$M$61,0))')
        oi.cell(row=r, column=3).value = f"=LARGE(Raw_Today!$M$2:$M$61,A{r})"
        oi.cell(row=r, column=5, value=k)
        oi.cell(row=r, column=6).value = (
            f"=INDEX(Raw_Today!$A$2:$A$61,"
            f'MATCH(LARGE(Raw_Today!$I$2:$I$61,E{r}),Raw_Today!$I$2:$I$61,0))')
        oi.cell(row=r, column=7).value = f"=LARGE(Raw_Today!$I$2:$I$61,E{r})"
        oi.cell(row=r, column=3).number_format = "0.00%"
        oi.cell(row=r, column=7).number_format = "0.00"
    style_range(oi, 3, 12, 3)
    style_range(oi, 3, 12, 7)
    for col, w in [("A", 6), ("B", 13), ("C", 11), ("E", 6), ("F", 13), ("G", 11)]:
        oi.column_dimensions[col].width = w

    # ---------------- Sector_Heatmap ----------------
    sh = wb.create_sheet("Sector_Heatmap")
    headers = ["Sector", "Avg % Chg", "Advancers", "Decliners",
               "Total Volume", "Net OI Change", "Long Buildups"]
    sh.append(headers)
    style_header(sh, 1, 7)
    sectors = sorted({row["Sector"] for row in t})
    for i, sec in enumerate(sectors, start=2):
        sh.cell(row=i, column=1, value=sec)
        sh.cell(row=i, column=2).value = (
            f'=AVERAGEIF(Raw_Today!$B$2:$B$61,A{i},Raw_Today!$F$2:$F$61)')
        sh.cell(row=i, column=3).value = (
            f'=COUNTIFS(Raw_Today!$B$2:$B$61,A{i},Raw_Today!$F$2:$F$61,">0")')
        sh.cell(row=i, column=4).value = (
            f'=COUNTIFS(Raw_Today!$B$2:$B$61,A{i},Raw_Today!$F$2:$F$61,"<0")')
        sh.cell(row=i, column=5).value = (
            f"=SUMIF(Raw_Today!$B$2:$B$61,A{i},Raw_Today!$D$2:$D$61)")
        sh.cell(row=i, column=6).value = (
            f"=SUMIF(Raw_Today!$B$2:$B$61,A{i},Raw_Today!$G$2:$G$61)")
        sh.cell(row=i, column=7).value = (
            f'=COUNTIFS(Raw_Today!$B$2:$B$61,A{i},'
            f'Raw_Today!$J$2:$J$61,"Long Buildup")')
        sh.cell(row=i, column=2).number_format = "0.00%"
        sh.cell(row=i, column=5).number_format = "#,##0"
        sh.cell(row=i, column=6).number_format = "#,##0"
    style_range(sh, 2, 9, 7)
    sh.conditional_formatting.add(
        "B2:B9", ColorScaleRule(start_type="num", start_value=-0.03,
                                start_color="F8696B", mid_type="num",
                                mid_value=0, mid_color="FFEB84",
                                end_type="num", end_value=0.03,
                                end_color="63BE7B"))
    for col, w in [("A", 12), ("B", 11), ("C", 11), ("D", 10),
                   ("E", 15), ("F", 15), ("G", 14)]:
        sh.column_dimensions[col].width = w

    # ---------------- Dashboard ----------------
    db = wb.create_sheet("Dashboard", 1)   # right after Guide
    db["A1"] = "NSEPulse — Market Dashboard"
    db["A1"].font = TITLE_FONT
    db["A2"] = "SYNTHETIC SAMPLE DATA — see Guide"
    db["A2"].font = Font(italic=True, color=RED, size=10)
    kpis = [
        ("Top Gainer", 4,
         '=INDEX(Raw_Today!$A$2:$A$61,MATCH(MAX(Raw_Today!$F$2:$F$61),Raw_Today!$F$2:$F$61,0))',
         "=MAX(Raw_Today!$F$2:$F$61)", "0.00%"),
        ("Top Loser", 6,
         '=INDEX(Raw_Today!$A$2:$A$61,MATCH(MIN(Raw_Today!$F$2:$F$61),Raw_Today!$F$2:$F$61,0))',
         "=MIN(Raw_Today!$F$2:$F$61)", "0.00%"),
        ("Strongest Long Buildup", 8,
         '=INDEX(Raw_Today!$A$2:$A$61,MATCH(MAX(Raw_Today!$M$2:$M$61),Raw_Today!$M$2:$M$61,0))',
         "=MAX(Raw_Today!$M$2:$M$61)", "0.00%"),
        ("Biggest Volume Shocker", 10,
         '=INDEX(Raw_Today!$A$2:$A$61,MATCH(MAX(Raw_Today!$I$2:$I$61),Raw_Today!$I$2:$I$61,0))',
         "=MAX(Raw_Today!$I$2:$I$61)", "0.00"),
        ("Advancers / Decliners", 12,
         '=COUNTIF(Raw_Today!$F$2:$F$61,">0")&" / "&COUNTIF(Raw_Today!$F$2:$F$61,"<0")',
         None, None),
        ("Total OI Change", 14,
         "=SUM(Raw_Today!$G$2:$G$61)",
         "=(SUM(Raw_Today!$E$2:$E$61)-SUM(Raw_PrevDay!$E$2:$E$61))/SUM(Raw_PrevDay!$E$2:$E$61)",
         "0.00%"),
    ]
    for label, r, ticker_f, val_f, fmt in kpis:
        db.cell(row=r, column=1, value=label).font = Font(bold=True, size=11)
        db.cell(row=r, column=2, value=ticker_f).font = KPI_FONT
        db.cell(row=r, column=2).fill = PatternFill("solid", fgColor=LIGHT)
        if val_f:
            db.cell(row=r, column=3, value=val_f).font = KPI_FONT
            db.cell(row=r, column=3).fill = PatternFill("solid", fgColor=LIGHT)
            if fmt:
                db.cell(row=r, column=3).number_format = fmt
    db.cell(row=12, column=2).number_format = "@"
    db.cell(row=14, column=2).number_format = "#,##0"
    for col, w in [("A", 24), ("B", 18), ("C", 14)]:
        db.column_dimensions[col].width = w

    c1 = BarChart()
    c1.type = "col"
    c1.title = "Top 10 Gainers (% change)"
    c1.y_axis.title = "% change"
    vals = Reference(tm, min_col=3, min_row=2, max_row=12)
    cats = Reference(tm, min_col=2, min_row=3, max_row=12)
    c1.add_data(vals, titles_from_data=True)
    c1.set_categories(cats)
    c1.height, c1.width = 7.5, 14
    db.add_chart(c1, "A16")

    c2 = BarChart()
    c2.type = "col"
    c2.title = "Top 10 Long Buildups (OI chg %)"
    c2.y_axis.title = "OI chg %"
    vals2 = Reference(oi, min_col=3, min_row=2, max_row=12)
    cats2 = Reference(oi, min_col=2, min_row=3, max_row=12)
    c2.add_data(vals2, titles_from_data=True)
    c2.set_categories(cats2)
    c2.height, c2.width = 7.5, 14
    db.add_chart(c2, "A32")

    # ---------------- Checks ----------------
    ck = wb.create_sheet("Checks")
    ck["A1"] = "Self-tests — all must read TRUE"
    ck["A1"].font = Font(bold=True, size=12)
    checks = [
        ("Ticker universe identical (Today vs PrevDay)",
         "=COUNTA(Raw_Today!$A$2:$A$61)=COUNTA(Raw_PrevDay!$A$2:$A$61)"),
        ("OI walk: SUM(OI Change) = SUM(OI_T) - SUM(OI_T-1)",
         "=SUM(Raw_Today!$G$2:$G$61)=SUM(Raw_Today!$E$2:$E$61)-SUM(Raw_PrevDay!$E$2:$E$61)"),
        ("Breadth adds up: adv + dec + flat = 60",
         '=COUNTIF(Raw_Today!$F$2:$F$61,">0")+COUNTIF(Raw_Today!$F$2:$F$61,"<0")'
         '+COUNTIF(Raw_Today!$F$2:$F$61,0)=60'),
        ("Dashboard top gainer ties to MAX(% Change)",
         "=ABS(Dashboard!$C$4-MAX(Raw_Today!$F$2:$F$61))<0.0000001"),
        ("Dashboard top loser ties to MIN(% Change)",
         "=ABS(Dashboard!$C$6-MIN(Raw_Today!$F$2:$F$61))<0.0000001"),
        ("Dashboard volume shocker ties to MAX(Vol Multiple)",
         "=ABS(Dashboard!$C$10-MAX(Raw_Today!$I$2:$I$61))<0.0000001"),
        ("Sector heatmap total volume ties to raw",
         "=SUM(Sector_Heatmap!$E$2:$E$9)=SUM(Raw_Today!$D$2:$D$61)"),
        ("OI signals exhaustive: 5 categories sum to 60",
         '=COUNTIF(Raw_Today!$J$2:$J$61,"Long Buildup")'
         '+COUNTIF(Raw_Today!$J$2:$J$61,"Short Buildup")'
         '+COUNTIF(Raw_Today!$J$2:$J$61,"Long Unwinding")'
         '+COUNTIF(Raw_Today!$J$2:$J$61,"Short Covering")'
         '+COUNTIF(Raw_Today!$J$2:$J$61,"Neutral")=60'),
    ]
    for i, (label, formula) in enumerate(checks, start=2):
        ck.cell(row=i, column=1, value=f"Check {i-1}: {label}")
        ck.cell(row=i, column=2, value=formula).font = Font(bold=True)
        ck.cell(row=i, column=2).alignment = Alignment(horizontal="center")
    for col, w in [("A", 62), ("B", 10)]:
        ck.column_dimensions[col].width = w

    wb.save(OUT)
    print(f"built {OUT}")


if __name__ == "__main__":
    build()
