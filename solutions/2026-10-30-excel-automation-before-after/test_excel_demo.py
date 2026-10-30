#!/usr/bin/env python3
"""Tests for the Excel automation demo.

Asserts (all must pass):
 1. after.xlsx opens and has Register + Summary sheets.
 2. Line Total cells contain formulas (not hard-coded values).
 3. Summary count/total cells contain formulas (COUNTIF / SUMIF / SUM).
 4. No duplicate Asset IDs remain in after.xlsx.
 5. Every populated Purchase Date cell in after.xlsx is a real date.
 6. Populated Quantity / Unit Price cells in after.xlsx are numbers.
 7. Flagged-row count in after.xlsx matches build_counts.json / BEFORE_AFTER.md.
 8. before.xlsx still contains the original mess (duplicates, text numbers,
    text dates, mixed currency symbols, blank quantities, invalid status)
    and has no Line Total column — the script did not clean it in place.
 9. Register sheet has frozen panes, an auto-filter and print titles;
    conditional formatting exists for flagged rows.
10. LibreOffice headless recalc (if soffice is available): converted
    Summary totals match the independently computed expected totals.

Usage:
    python3 test_excel_demo.py
Exit code 0 only if every check passes. Output is written to VERIFY.txt
by the caller (tee), and the recalc section is printed here.
"""

import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BEFORE_PATH = HERE / "before.xlsx"
AFTER_PATH = HERE / "after.xlsx"
COUNTS_PATH = HERE / "build_counts.json"
REPORT_PATH = HERE / "BEFORE_AFTER.md"

FAILURES = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"[{status}] {name}" + (f" — {detail}" if detail else ""))
    if not condition:
        FAILURES.append(name)


def main():
    counts = json.loads(COUNTS_PATH.read_text(encoding="utf-8"))

    # ---- before.xlsx still messy -------------------------------------------
    wb_before = load_workbook(BEFORE_PATH)
    ws_b = wb_before["Register"]
    rows_b = list(ws_b.iter_rows(min_row=4, max_col=8, values_only=True))
    check("before.xlsx opens with 60 data rows", len(rows_b) == 60, f"{len(rows_b)} rows")
    ids_b = [r[0] for r in rows_b]
    dupes = [k for k, v in Counter(ids_b).items() if v > 1]
    check("before.xlsx still contains duplicate Asset IDs", len(dupes) == 5,
          f"duplicated IDs: {dupes}")
    check("before.xlsx has quantities stored as text",
          any(isinstance(r[4], str) and r[4].strip() != "" for r in rows_b))
    check("before.xlsx has blank quantities",
          any(r[4] is None for r in rows_b))
    check("before.xlsx has prices stored as text with currency symbols",
          any(isinstance(r[5], str) and any(s in r[5] for s in ("\u20b9", "Rs.", "INR", "$"))
              for r in rows_b))
    check("before.xlsx has dates stored as text",
          any(isinstance(r[3], str) for r in rows_b))
    check("before.xlsx has the invalid status value 'In Usee'",
          any(r[6] == "In Usee" for r in rows_b))
    check("before.xlsx has no Line Total column (8 columns only)",
          ws_b.max_column == 8, f"max_column={ws_b.max_column}")

    # ---- after.xlsx structure ------------------------------------------------
    wb = load_workbook(AFTER_PATH)  # formulas, not cached values
    check("after.xlsx opens with Register + Summary sheets",
          set(wb.sheetnames) == {"Register", "Summary"}, str(wb.sheetnames))
    ws = wb["Register"]
    rows = list(ws.iter_rows(min_row=4, max_col=11, values_only=False))
    check("after.xlsx has 55 data rows", len(rows) == counts["after_rows"],
          f"{len(rows)} rows")
    ids = [r[0].value for r in rows]
    check("no duplicate Asset IDs remain in after.xlsx", len(set(ids)) == len(ids))

    # ---- dates / numbers are real types --------------------------------------
    bad_dates = [(r[0].value, r[3].value) for r in rows
                 if r[3].value is not None and not isinstance(r[3].value, (date, datetime))]
    check("every populated Purchase Date in after.xlsx is a real date",
          not bad_dates, f"non-date cells: {bad_dates}")
    bad_qty = [(r[0].value, r[4].value) for r in rows
               if r[4].value is not None and not isinstance(r[4].value, (int, float))]
    check("every populated Quantity in after.xlsx is a number", not bad_qty)
    bad_price = [(r[0].value, r[5].value) for r in rows
                 if r[5].value is not None and not isinstance(r[5].value, (int, float))]
    check("every populated Unit Price in after.xlsx is a number", not bad_price)

    # ---- formulas, not values -------------------------------------------------
    line_totals = [r[6] for r in rows]
    check("every Line Total cell contains a formula",
          all(isinstance(c.value, str) and c.value.startswith("=") for c in line_totals))
    check("Line Total formulas multiply Quantity by Unit Price",
          all("*" in c.value and "E" in c.value and "F" in c.value for c in line_totals),
          f"example: {line_totals[0].value}")
    check("Line Total formulas return blank when Quantity or Price is blank",
          all("OR(" in c.value and '""' in c.value for c in line_totals))

    sm = wb["Summary"]
    summary_formulas = {f"B{r}": sm.cell(row=r, column=2).value for r in range(5, 9)}
    summary_formulas.update({f"C{r}": sm.cell(row=r, column=3).value for r in range(5, 9)})
    check("Summary category counts are COUNTIF formulas",
          all(isinstance(v, str) and "COUNTIF" in v for k, v in summary_formulas.items()
              if k.startswith("B")))
    check("Summary category totals are SUMIF formulas",
          all(isinstance(v, str) and "SUMIF" in v for k, v in summary_formulas.items()
              if k.startswith("C")))
    grand_row = 9
    check("Summary Grand Total row uses SUM formulas",
          "SUM(" in str(sm.cell(row=grand_row, column=2).value)
          and "SUM(" in str(sm.cell(row=grand_row, column=3).value))
    check("Summary flagged-row count is a COUNTIF formula",
          "COUNTIF" in str(sm.cell(row=12, column=2).value),
          str(sm.cell(row=12, column=2).value))

    # ---- flags -----------------------------------------------------------------
    flagged = sum(1 for r in rows if r[9].value)
    check("flagged-row count matches build_counts.json",
          flagged == counts["flagged_rows"], f"{flagged} flagged")
    report_text = REPORT_PATH.read_text(encoding="utf-8")
    m = re.search(r"Flagged rows in after\.xlsx: (\d+)", report_text)
    check("flagged-row count matches BEFORE_AFTER.md",
          m is not None and int(m.group(1)) == flagged)

    # ---- sheet presentation ------------------------------------------------------
    check("Register sheet has frozen panes", ws.freeze_panes == "A4",
          str(ws.freeze_panes))
    check("Register sheet has an auto-filter", ws.auto_filter.ref is not None,
          str(ws.auto_filter.ref))
    check("Register sheet has print titles",
          str(ws.print_title_rows).replace("$", "") == "1:3",
          str(ws.print_title_rows))
    check("conditional formatting exists for flagged rows",
          len(list(ws.conditional_formatting._cf_rules)) > 0)

    # ---- LibreOffice recalc --------------------------------------------------------
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("[SKIP] LibreOffice recalc check — soffice not available; "
              "formulas verified by inspection above.")
    else:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "xlsx",
                 "--outdir", tmp, str(AFTER_PATH)],
                capture_output=True, text=True, timeout=180, check=True)
            converted = Path(tmp) / "after.xlsx"
            wb_calc = load_workbook(converted, data_only=True)
            sm_calc = wb_calc["Summary"]
            expected = counts["totals_by_category"]
            cats = ["IT Equipment", "Furniture", "Vehicle", "Tools"]
            for i, cat in enumerate(cats):
                got = sm_calc.cell(row=5 + i, column=3).value
                check(f"recalc total for {cat} matches expected",
                      got is not None and abs(got - expected[cat]) < 0.01,
                      f"got {got}, expected {expected[cat]:,.2f}")
            got_grand = sm_calc.cell(row=9, column=3).value
            check("recalc grand total matches expected",
                  got_grand is not None and abs(got_grand - counts["grand_total"]) < 0.01,
                  f"got {got_grand}, expected {counts['grand_total']:,.2f}")
            got_flagged = sm_calc.cell(row=12, column=2).value
            check("recalc flagged-row count matches expected",
                  got_flagged == counts["flagged_rows"], f"got {got_flagged}")
            sample = wb_calc["Register"].cell(row=4, column=7).value
            check("recalc Line Total computes a number in row 4",
                  isinstance(sample, (int, float)), f"got {sample!r}")

    print()
    if FAILURES:
        print(f"RESULT: {len(FAILURES)} check(s) FAILED: {FAILURES}")
        return 1
    print("RESULT: all checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
