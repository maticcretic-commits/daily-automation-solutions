#!/usr/bin/env python3
"""Build lease-disparity.xlsx from data/leases.csv — every figure a live formula.

Sheets: Guide, Dashboard, Leases, Schedule, Assumptions, Checks.
Rules honoured (learned the hard way): plain A1 ranges only, range
literals repeat the column letter on both ends ($E$2:$E$13), no Table
objects, chart series titles come from header cells (titles_from_data),
never plain strings. The Leases sheet holds values in A-G and formulas in
H-L, so every derived figure is auditable in place. Valuation is a monthly
annuity in arrears at ONE yield rate (Assumptions!$B$3, monthly), applied
identically to all 12 leases: PV = disparity * (1-(1+r)^-n)/r.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "lease-disparity.xlsx")

NAVY = "1F3864"
TEAL = "2E75B6"
LIGHT = "D9E2F3"

N_LEASES = 12
LAST = N_LEASES + 1  # data rows 2..13
TOTAL_ROW = LAST + 1  # 14
MAX_MONTHS = 60
SCHED_LAST = MAX_MONTHS + 1  # schedule rows 2..61


def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.reader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def main():
    rows = read_csv("leases.csv")
    header, data = rows[0], rows[1:]
    assert len(data) == N_LEASES, f"expected {N_LEASES} leases, got {len(data)}"

    wb = Workbook()

    # ---- Guide ----
    ws = wb.active
    ws.title = "Guide"
    guide_lines = [
        ("LeaseDisparity — Contracted vs Market Rent: Disparity Valuation", True),
        ("Demo workbook built around a lease rental disparity brief: compare", False),
        ("contracted rent with current market rent for up to 12 leases, and value", False),
        ("the disparity over each remaining term at a single yield rate.", False),
        ("", False),
        ("How to use: replace the values in Leases columns A-G with your own", False),
        ("leases (up to 12 rows; extend the ranges if you add rows). Every", False),
        ("derived figure — disparity, present value, dashboard KPIs, the monthly", False),
        ("run-off schedule and the Checks sheet — is a live formula. Change the", False),
        ("yield on the Assumptions sheet and the whole valuation reprices.", False),
        ("", False),
        ("Valuation: monthly annuity in arrears. Monthly disparity = market", False),
        ("rent minus contracted rent. PV = disparity x (1-(1+r)^-n)/r, where", False),
        ("r is the monthly yield and n is the remaining term in months.", False),
        ("Positive disparity = under-rented (market above contract). Negative", False),
        ("disparity = over-rented. Within $25/month of market = At market.", False),
        ("", False),
        ("All data are fictional and invented for this demo. No real client", False),
        ("data. Rents are USD per month; yield 8.5% per year (see Assumptions).", False),
    ]
    for i, (text, bold) in enumerate(guide_lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11,
                      color=NAVY if i == 1 else "000000")
    widths(ws, {"A": 78})

    # ---- Assumptions ----
    ws = wb.create_sheet("Assumptions")
    ws.cell(row=1, column=1, value="Assumptions").font = Font(bold=True, size=13, color=NAVY)
    ws.cell(row=2, column=1, value="Annual yield").font = Font(bold=True)
    ws.cell(row=2, column=2, value=0.085).number_format = "0.0%"
    ws.cell(row=3, column=1, value="Monthly yield").font = Font(bold=True)
    ws.cell(row=3, column=2, value="=$B$2/12").number_format = "0.0000%"
    ws.cell(row=4, column=1, value="As-of date").font = Font(bold=True)
    ws.cell(row=4, column=2, value="2026-10-30")
    ws.cell(row=5, column=1, value="Valuation basis").font = Font(bold=True)
    ws.cell(row=5, column=2, value="Monthly annuity, payments in arrears")
    ws.cell(row=6, column=1, value="At-market tolerance ($/month)").font = Font(bold=True)
    ws.cell(row=6, column=2, value=25).number_format = "#,##0"
    ws.cell(row=7, column=1, value="Max leases in this workbook").font = Font(bold=True)
    ws.cell(row=7, column=2, value=N_LEASES)
    widths(ws, {"A": 30, "B": 38})

    # ---- Leases ----
    ws = wb.create_sheet("Leases")
    cols = ["LeaseID", "Property", "Tenant", "Area", "Contract Rent ($/mo)",
            "Market Rent ($/mo)", "Remaining Term (months)",
            "Monthly Disparity ($)", "Annual Disparity ($)",
            "Undiscounted Disparity ($)", "PV of Disparity ($)", "Status"]
    ws.append(cols)
    style_header(ws, 1, len(cols))
    for i, row in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=row[0])
        ws.cell(row=i, column=2, value=row[1])
        ws.cell(row=i, column=3, value=row[2])
        ws.cell(row=i, column=4, value=row[3])
        ws.cell(row=i, column=5, value=float(row[4])).number_format = "#,##0"
        ws.cell(row=i, column=6, value=float(row[5])).number_format = "#,##0"
        ws.cell(row=i, column=7, value=int(row[6])).number_format = "#,##0"
        ws.cell(row=i, column=8, value=f"=$F{i}-$E{i}").number_format = "#,##0"
        ws.cell(row=i, column=9, value=f"=$H{i}*12").number_format = "#,##0"
        ws.cell(row=i, column=10, value=f"=$H{i}*$G{i}").number_format = "#,##0"
        ws.cell(row=i, column=11,
                value=f"=$H{i}*(1-(1+Assumptions!$B$3)^-$G{i})/Assumptions!$B$3"
                ).number_format = "#,##0"
        ws.cell(row=i, column=12,
                value=(f"=IF(ABS($H{i})<=Assumptions!$B$6,\"At market\","
                       f"IF($H{i}>0,\"Under-rented\",\"Over-rented\"))"))
    ws.cell(row=TOTAL_ROW, column=1, value="TOTAL").font = Font(bold=True)
    for col in ("E", "F", "H", "I", "J", "K"):
        c = ws.cell(row=TOTAL_ROW, column=ord(col) - 64,
                    value=f"=SUM(${col}$2:${col}${LAST})")
        c.font = Font(bold=True)
        c.number_format = "#,##0"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.auto_filter.ref = f"A1:L{LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 9, "B": 30, "C": 24, "D": 14, "E": 18, "F": 17,
                "G": 16, "H": 18, "I": 17, "J": 20, "K": 18, "L": 13})

    # ---- Schedule: monthly run-off, months 1..60 ----
    ws = wb.create_sheet("Schedule")
    ws.append(["Month", "Monthly Disparity ($)", "Discount Factor",
               "Discounted Disparity ($)", "Cumulative PV ($)"])
    style_header(ws, 1, 5)
    for m in range(1, MAX_MONTHS + 1):
        r = m + 1
        ws.cell(row=r, column=1, value=m)
        ws.cell(row=r, column=2,
                value=(f"=SUMPRODUCT((Leases!$G$2:$G${LAST}>=$A{r})"
                       f"*Leases!$H$2:$H${LAST})")).number_format = "#,##0"
        ws.cell(row=r, column=3,
                value=f"=(1+Assumptions!$B$3)^-$A{r}").number_format = "0.0000"
        ws.cell(row=r, column=4, value=f"=$B{r}*$C{r}").number_format = "#,##0"
        ws.cell(row=r, column=5, value=f"=SUM($D$2:$D{r})").number_format = "#,##0"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 8, "B": 22, "C": 16, "D": 24, "E": 18})

    # ---- Checks ----
    ws = wb.create_sheet("Checks")
    ws.append(["Check", "Expected", "Got", "Diff", "Status"])
    checks = [
        ("Lease count", N_LEASES, f"=COUNTA(Leases!$A$2:$A${LAST})", "0"),
        ("Dashboard contract total ties to Leases",
         f"=SUM(Leases!$E$2:$E${LAST})", "=Dashboard!$B$4", "0.01"),
        ("Dashboard market total ties to Leases",
         f"=SUM(Leases!$F$2:$F${LAST})", "=Dashboard!$B$5", "0.01"),
        ("Dashboard monthly disparity ties to Leases",
         f"=SUM(Leases!$H$2:$H${LAST})", "=Dashboard!$B$6", "0.01"),
        ("Dashboard PV ties to Leases",
         f"=SUM(Leases!$K$2:$K${LAST})", "=Dashboard!$B$9", "0.01"),
        ("Schedule month 1 ties to dashboard monthly disparity",
         "=Schedule!$B$2", "=Dashboard!$B$6", "0.01"),
        ("Schedule cumulative PV ties to dashboard PV",
         f"=Schedule!$E${SCHED_LAST}", "=Dashboard!$B$9", "1"),
        ("Leases with negative remaining term", 0,
         f"=COUNTIF(Leases!$G$2:$G${LAST},\"<0\")", "0"),
        ("Leases with blank market rent", 0,
         f"=COUNTBLANK(Leases!$F$2:$F${LAST})", "0"),
    ]
    for i, (label, expected, got, tol) in enumerate(checks, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=expected)
        ws.cell(row=i, column=3, value=got)
        ws.cell(row=i, column=4, value=f"=$C{i}-$B{i}")
        ws.cell(row=i, column=5,
                value=f"=IF(ABS($D{i})<={tol},\"PASS\",\"FAIL\")")
    style_header(ws, 1, 5)
    ws.freeze_panes = "A2"
    for r in range(2, 2 + len(checks)):
        for col in ("B", "C", "D"):
            ws[f"{col}{r}"].number_format = "#,##0.00"
    widths(ws, {"A": 52, "B": 16, "C": 16, "D": 12, "E": 10})

    # ---- Dashboard (built last so it can reference the other sheets) ----
    ws = wb.create_sheet("Dashboard", 1)
    ws.cell(row=1, column=1,
            value="LeaseDisparity — Portfolio Dashboard (as of 2026-10-30)").font = \
        Font(bold=True, size=14, color=NAVY)
    ws.cell(row=2, column=1,
            value="All values are live formulas over Leases / Assumptions.").font = \
        Font(italic=True, color="595959")
    kpis = [
        ("Leases in portfolio", f"=COUNTA(Leases!$A$2:$A${LAST})", "#,##0"),
        ("Total contracted rent ($/month)", f"=SUM(Leases!$E$2:$E${LAST})", "#,##0"),
        ("Total market rent ($/month)", f"=SUM(Leases!$F$2:$F${LAST})", "#,##0"),
        ("Monthly disparity — market minus contract ($)",
         f"=SUM(Leases!$H$2:$H${LAST})", "#,##0"),
        ("Annual disparity ($)", f"=SUM(Leases!$I$2:$I${LAST})", "#,##0"),
        ("Undiscounted disparity over remaining terms ($)",
         f"=SUM(Leases!$J$2:$J${LAST})", "#,##0"),
        ("Present value of disparity ($)", f"=SUM(Leases!$K$2:$K${LAST})", "#,##0"),
        ("Under-rented leases",
         f"=COUNTIF(Leases!$L$2:$L${LAST},\"Under-rented\")", "#,##0"),
        ("Over-rented leases",
         f"=COUNTIF(Leases!$L$2:$L${LAST},\"Over-rented\")", "#,##0"),
        ("At-market leases",
         f"=COUNTIF(Leases!$L$2:$L${LAST},\"At market\")", "#,##0"),
        ("Average remaining term (months)",
         f"=AVERAGE(Leases!$G$2:$G${LAST})", "0.0"),
        ("Annual yield", "=Assumptions!$B$2", "0.0%"),
    ]
    # KPI rows start at row 3; value cells are B3..B14.
    for i, (label, formula, fmt) in enumerate(kpis, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format = fmt
        c.font = Font(bold=True, size=12)
    ws.cell(row=15, column=1,
            value="Charts below are wired to Leases / Schedule — they move with the data.").font = \
        Font(italic=True, color="595959")
    widths(ws, {"A": 46, "B": 18, "C": 20, "D": 20, "E": 20, "F": 20})

    col_chart = BarChart()
    col_chart.type = "col"
    col_chart.title = "PV of Disparity by Lease ($)"
    col_chart.y_axis.title = "$"
    col_chart.style = 10
    data = Reference(wb["Leases"], min_col=11, min_row=1, max_row=LAST)
    cats = Reference(wb["Leases"], min_col=1, min_row=2, max_row=LAST)
    col_chart.add_data(data, titles_from_data=True)
    col_chart.set_categories(cats)
    col_chart.height = 7.5
    col_chart.width = 15
    ws.add_chart(col_chart, "A17")

    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.title = "Contracted vs Market Rent by Lease ($/month)"
    bar_chart.x_axis.title = "$/month"
    bar_chart.style = 10
    data = Reference(wb["Leases"], min_col=5, min_row=1, max_row=LAST, max_col=6)
    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    bar_chart.height = 7.5
    bar_chart.width = 15
    ws.add_chart(bar_chart, "A32")

    line = LineChart()
    line.title = "Cumulative PV of Disparity by Month ($)"
    line.y_axis.title = "$"
    line.style = 10
    ldata = Reference(wb["Schedule"], min_col=5, min_row=1, max_row=SCHED_LAST)
    lcats = Reference(wb["Schedule"], min_col=1, min_row=2, max_row=SCHED_LAST)
    line.add_data(ldata, titles_from_data=True)
    line.set_categories(lcats)
    line.height = 7.5
    line.width = 15
    ws.add_chart(line, "A47")

    wb.save(OUT)
    print(f"wrote {OUT} sheets={wb.sheetnames}")


if __name__ == "__main__":
    main()
