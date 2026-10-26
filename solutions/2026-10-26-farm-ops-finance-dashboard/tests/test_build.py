#!/usr/bin/env python3
"""FarmPulse tests: build, LibreOffice-recalc, compare vs ground truth.

Run from the build dir: python3 tests/test_build.py
Recalc uses a COPY in a separate input dir (soffice overwrites the input
basename) with --outdir elsewhere.
"""
import os
import shutil
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from main import ground_truth  # noqa: E402

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(HERE, "farmpulse.xlsx")
WORK = "/tmp/farmpulse-test"
INDIR = os.path.join(WORK, "in")
OUTDIR = os.path.join(WORK, "out")


def recalc():
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(INDIR)
    os.makedirs(OUTDIR)
    shutil.copy(SRC, os.path.join(INDIR, "farmpulse.xlsx"))
    r = subprocess.run(["soffice", "--headless", "--calc", "--convert-to",
                        "xlsx", "--outdir", OUTDIR,
                        os.path.join(INDIR, "farmpulse.xlsx")],
                       capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stdout + r.stderr
    return load_workbook(os.path.join(OUTDIR, "farmpulse.xlsx"),
                         data_only=True)


def val(wb, sheet, cell):
    return wb[sheet][cell].value


def close(a, b, tol=0.02):
    assert a is not None and b is not None, f"missing value {a} vs {b}"
    assert abs(float(a) - float(b)) <= tol, f"{a} != {b}"


def main():
    gt = ground_truth()
    wb = recalc()
    # Formula-error scan over all value sheets.
    errors = []
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("#"):
                    errors.append(f"{name}!{c.coordinate}={v}")
    assert not errors, f"formula errors: {errors[:10]}"

    d = wb["Dashboard"]
    close(d["B4"].value, gt["total_income"])
    close(d["B5"].value, gt["total_expense"])
    close(d["B6"].value, gt["net_profit"])
    close(d["B7"].value, gt["margin"], tol=0.0005)
    assert d["B8"].value == gt["best_enterprise"], d["B8"].value
    close(d["B9"].value, gt["total_expense"] / gt["total_income"],
          tol=0.0005)

    ep = wb["Enterprise_Profit"]
    for i, e in enumerate(["Dairy", "Wheat", "Vegetables", "Poultry",
                           "Sheep"], start=2):
        assert ep.cell(row=i, column=1).value == e
        for j, q in enumerate(["2025-Q3", "2025-Q4", "2026-Q1",
                               "2026-Q2"], start=2):
            close(ep.cell(row=i, column=j).value, gt["ent_profit"][e][q])
        close(ep.cell(row=i, column=6).value,
              sum(gt["ent_profit"][e].values()))

    ia = wb["Income_Activity"]
    for i, a in enumerate(["Milk Sales", "Calf Sales", "Grain Sales",
                           "Straw Sales", "Vegetable Sales",
                           "Market Stall Sales", "Egg Sales",
                           "Broiler Sales", "Wool Sales", "Lamb Sales"],
                          start=2):
        close(ia.cell(row=i, column=6).value,
              sum(gt["inc_activity"].get(a, {}).values()))

    ec = wb["Expense_Category"]
    for i, c in enumerate(["Feed", "Labour", "Veterinary",
                           "Fuel & Machinery", "Fertiliser & Seed",
                           "Transport", "Utilities", "Other"], start=2):
        close(ec.cell(row=i, column=6).value,
              sum(gt["exp_category"].get(c, {}).values()))

    q = wb["Dashboard"]
    for i, quarter in enumerate(["2025-Q3", "2025-Q4", "2026-Q1",
                                 "2026-Q2"], start=12):
        close(q.cell(row=i, column=2).value, gt["q_income"][quarter])
        close(q.cell(row=i, column=3).value, gt["q_expense"][quarter])

    ck = wb["Checks"]
    for r in range(2, 10):
        assert ck.cell(row=r, column=3).value is True, \
            f"check row {r} failed: {ck.cell(row=r, column=1).value}"

    print(f"ALL TESTS PASS: income {gt['total_income']}, "
          f"expenses {gt['total_expense']}, net {gt['net_profit']}, "
          f"margin {gt['margin']:.1%}, best {gt['best_enterprise']}")


if __name__ == "__main__":
    main()
