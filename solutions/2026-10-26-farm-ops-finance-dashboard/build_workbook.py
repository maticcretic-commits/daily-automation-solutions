#!/usr/bin/env python3
"""Build farmpulse.xlsx: formula-driven farm ops & finance dashboard.

Every figure is a live A1-style formula over the Raw sheets (headroom to
row 1000, so pasting a new quarter's exports refreshes everything). No
Table objects, no structured references — LibreOffice-recalc safe.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
QUARTERS = ["2025-Q3", "2025-Q4", "2026-Q1", "2026-Q2"]
ENTERPRISES = ["Dairy", "Wheat", "Vegetables", "Poultry", "Sheep"]
ACTIVITIES = ["Milk Sales", "Calf Sales", "Grain Sales", "Straw Sales",
              "Vegetable Sales", "Market Stall Sales", "Egg Sales",
              "Broiler Sales", "Wool Sales", "Lamb Sales"]
EXP_CATS = ["Feed", "Labour", "Veterinary", "Fuel & Machinery",
            "Fertiliser & Seed", "Transport", "Utilities", "Other"]
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
MONEY = '#,##0.00'
PCT = '0.0%'

INC = "Raw_Income"
EXP = "Raw_Expense"


def read_csv(name):
    with open(os.path.join(HERE, "data", name), newline="") as f:
        return list(csv.DictReader(f))


def header_row(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = BOLD
        c.fill = HDR_FILL


def style_money(ws, col, r1, r2):
    for r in range(r1, r2 + 1):
        ws.cell(row=r, column=col).number_format = MONEY


def main():
    inc = read_csv("income.csv")
    exp = read_csv("expense.csv")
    wb = Workbook()

    # ---- Guide ----
    ws = wb.active
    ws.title = "Guide"
    ws["A1"] = "FarmPulse — Farm Ops & Finance Dashboard (SYNTHETIC DEMO)"
    ws["A1"].font = TITLE
    guide = [
        "Fictional farm data (seed 20261026). No client data is used anywhere.",
        "Built around a real brief shape: one interactive dashboard from",
        "Excel/CSV farm data — income by activity, expenses by category,",
        "profit by enterprise over time, with drill-through to transactions.",
        "",
        "HOW TO REFRESH: paste new income rows into Raw_Income and new",
        "expense rows into Raw_Expense (same columns, rows 2-1000). Every",
        "summary, KPI and chart updates by formula — nothing is pasted",
        "as values. Keep the Quarter label format YYYY-Qn.",
        "",
        "HOW TO FILTER / DRILL: use the AutoFilter on any Raw sheet column",
        "(Season/Quarter, Enterprise, Paddock/Herd, Cost Centre) to drill",
        "from a summary figure down to its transactions. The Dashboard",
        "quarterly table is the overview; Raw sheets are the detail.",
        "",
        "TRAFFIC LIGHTS: Enterprise_Profit Margin is green above 20%,",
        "red below 0%. A negative enterprise margin means that enterprise",
        "lost money over the four quarters — check its expense mix.",
        "",
        "POWER BI PATH: the same model maps 1:1 to Power BI — see",
        "powerbi/MEASURES.md for the DAX measures (Total Income,",
        "Total Expenses, Net Profit, Margin %, profit by enterprise).",
        "Import Raw_Income + Raw_Expense as two fact tables and the",
        "measures reproduce this workbook exactly.",
    ]
    for i, line in enumerate(guide, start=3):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 110

    # ---- Raw sheets (built first so formulas can reference them) ----
    def raw_sheet(name, rows, dim4):
        s = wb.create_sheet(name)
        headers = ["Date", "Quarter", "Enterprise", dim4,
                   "Paddock/Herd", "Cost Centre", "Amount"]
        header_row(s, 1, headers)
        for i, r in enumerate(rows, start=2):
            s.cell(row=i, column=1, value=r["date"])
            s.cell(row=i, column=2, value=r["quarter"])
            s.cell(row=i, column=3, value=r["enterprise"])
            s.cell(row=i, column=4, value=r[dim4.lower().replace(" ", "_")
                                             if dim4 == "Activity"
                                             else "category"])
            s.cell(row=i, column=5, value=r["paddock_herd"])
            s.cell(row=i, column=6, value=r["cost_centre"])
            c = s.cell(row=i, column=7, value=float(r["amount"]))
            c.number_format = MONEY
        s.auto_filter.ref = f"A1:G{len(rows) + 1}"
        s.freeze_panes = "A2"
        widths = [12, 10, 13, 20, 13, 20, 13]
        for j, w in enumerate(widths, start=1):
            s.column_dimensions[get_column_letter(j)].width = w
        return s

    # NOTE: income rows carry "activity", expense rows carry "category".
    def raw_sheet2(name, rows, key4, label4):
        s = wb.create_sheet(name)
        header_row(s, 1, ["Date", "Quarter", "Enterprise", label4,
                          "Paddock/Herd", "Cost Centre", "Amount"])
        for i, r in enumerate(rows, start=2):
            s.cell(row=i, column=1, value=r["date"])
            s.cell(row=i, column=2, value=r["quarter"])
            s.cell(row=i, column=3, value=r["enterprise"])
            s.cell(row=i, column=4, value=r[key4])
            s.cell(row=i, column=5, value=r["paddock_herd"])
            s.cell(row=i, column=6, value=r["cost_centre"])
            c = s.cell(row=i, column=7, value=float(r["amount"]))
            c.number_format = MONEY
        s.auto_filter.ref = f"A1:G{len(rows) + 1}"
        s.freeze_panes = "A2"
        for j, w in enumerate([12, 10, 13, 20, 13, 20, 13], start=1):
            s.column_dimensions[get_column_letter(j)].width = w
        return s

    raw_sheet2(INC, inc, "activity", "Activity")
    raw_sheet2(EXP, exp, "category", "Category")

    # ---- Enterprise_Profit ----
    ws = wb.create_sheet("Enterprise_Profit")
    header_row(ws, 1, ["Enterprise"] + QUARTERS +
               ["Total Profit", "Total Income", "Margin"])
    for i, e in enumerate(ENTERPRISES, start=2):
        ws.cell(row=i, column=1, value=e)
        for j, q in enumerate(QUARTERS, start=2):
            col = get_column_letter(j)
            ws.cell(row=i, column=j).value = (
                f"=SUMIFS({INC}!$G$2:$G$1000,{INC}!$C$2:$C$1000,$A{i},"
                f"{INC}!$B$2:$B$1000,{col}$1)"
                f"-SUMIFS({EXP}!$G$2:$G$1000,{EXP}!$C$2:$C$1000,$A{i},"
                f"{EXP}!$B$2:$B$1000,{col}$1)")
        ws.cell(row=i, column=6).value = f"=SUM(B{i}:E{i})"
        ws.cell(row=i, column=7).value = (
            f"=SUMIFS({INC}!$G$2:$G$1000,{INC}!$C$2:$C$1000,$A{i})")
        ws.cell(row=i, column=8).value = f"=IF(G{i}=0,0,F{i}/G{i})"
    ws.cell(row=7, column=1, value="TOTAL").font = BOLD
    for j in range(2, 8):
        col = get_column_letter(j)
        ws.cell(row=7, column=j).value = f"=SUM({col}2:{col}6)"
        ws.cell(row=7, column=j).font = BOLD
    ws.cell(row=7, column=8).value = "=IF(G7=0,0,F7/G7)"
    ws.cell(row=7, column=8).font = BOLD
    for j in range(2, 8):
        style_money(ws, j, 2, 7)
    for r in range(2, 8):
        ws.cell(row=r, column=8).number_format = PCT
    ws.conditional_formatting.add("H2:H6", CellIsRule(
        operator="greaterThan", formula=["0.2"],
        fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add("H2:H6", CellIsRule(
        operator="lessThan", formula=["0"],
        fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.column_dimensions["A"].width = 14
    for j in range(2, 9):
        ws.column_dimensions[get_column_letter(j)].width = 14

    # ---- Income_Activity ----
    ws = wb.create_sheet("Income_Activity")
    header_row(ws, 1, ["Activity"] + QUARTERS + ["Total"])
    for i, a in enumerate(ACTIVITIES, start=2):
        ws.cell(row=i, column=1, value=a)
        for j, _q in enumerate(QUARTERS, start=2):
            col = get_column_letter(j)
            ws.cell(row=i, column=j).value = (
                f"=SUMIFS({INC}!$G$2:$G$1000,{INC}!$D$2:$D$1000,$A{i},"
                f"{INC}!$B$2:$B$1000,{col}$1)")
        ws.cell(row=i, column=6).value = f"=SUM(B{i}:E{i})"
    ws.cell(row=12, column=1, value="TOTAL").font = BOLD
    for j in range(2, 7):
        col = get_column_letter(j)
        ws.cell(row=12, column=j).value = f"=SUM({col}2:{col}11)"
        ws.cell(row=12, column=j).font = BOLD
    for j in range(2, 7):
        style_money(ws, j, 2, 12)
    ws.column_dimensions["A"].width = 20
    for j in range(2, 7):
        ws.column_dimensions[get_column_letter(j)].width = 14

    # ---- Expense_Category ----
    ws = wb.create_sheet("Expense_Category")
    header_row(ws, 1, ["Category"] + QUARTERS + ["Total"])
    for i, c in enumerate(EXP_CATS, start=2):
        ws.cell(row=i, column=1, value=c)
        for j, _q in enumerate(QUARTERS, start=2):
            col = get_column_letter(j)
            ws.cell(row=i, column=j).value = (
                f"=SUMIFS({EXP}!$G$2:$G$1000,{EXP}!$D$2:$D$1000,$A{i},"
                f"{EXP}!$B$2:$B$1000,{col}$1)")
        ws.cell(row=i, column=6).value = f"=SUM(B{i}:E{i})"
    ws.cell(row=10, column=1, value="TOTAL").font = BOLD
    for j in range(2, 7):
        col = get_column_letter(j)
        ws.cell(row=10, column=j).value = f"=SUM({col}2:{col}9)"
        ws.cell(row=10, column=j).font = BOLD
    for j in range(2, 7):
        style_money(ws, j, 2, 10)
    ws.column_dimensions["A"].width = 20
    for j in range(2, 7):
        ws.column_dimensions[get_column_letter(j)].width = 14

    # ---- Dashboard ----
    ws = wb.create_sheet("Dashboard", 1)  # right after Guide
    ws["A1"] = "FarmPulse Dashboard — income, expenses & profit by enterprise"
    ws["A1"].font = TITLE
    header_row(ws, 3, ["KPI", "Value"])
    kpis = [
        ("Total Income", f"=SUM({INC}!$G$2:$G$1000)"),
        ("Total Expenses", f"=SUM({EXP}!$G$2:$G$1000)"),
        ("Net Profit", "=B4-B5"),
        ("Margin", "=IF(B4=0,0,B6/B4)"),
        ("Best Enterprise (profit)",
         "=INDEX(Enterprise_Profit!$A$2:$A$6,"
         "MATCH(MAX(Enterprise_Profit!$F$2:$F$6),"
         "Enterprise_Profit!$F$2:$F$6,0))"),
        ("Expense Ratio", "=IF(B4=0,0,B5/B4)"),
    ]
    for i, (label, formula) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=formula)
    for r in (4, 5, 6):
        ws.cell(row=r, column=2).number_format = MONEY
    for r in (7, 9):
        ws.cell(row=r, column=2).number_format = PCT
    header_row(ws, 11, ["Quarter", "Income", "Expense", "Profit"])
    for i, q in enumerate(QUARTERS, start=12):
        ws.cell(row=i, column=1, value=q)
        ws.cell(row=i, column=2).value = (
            f"=SUMIFS({INC}!$G$2:$G$1000,{INC}!$B$2:$B$1000,$A{i})")
        ws.cell(row=i, column=3).value = (
            f"=SUMIFS({EXP}!$G$2:$G$1000,{EXP}!$B$2:$B$1000,$A{i})")
        ws.cell(row=i, column=4).value = f"=B{i}-C{i}"
    ws.cell(row=16, column=1, value="TOTAL").font = BOLD
    for j in (2, 3, 4):
        col = get_column_letter(j)
        ws.cell(row=16, column=j).value = f"=SUM({col}12:{col}15)"
        ws.cell(row=16, column=j).font = BOLD
    for j in (2, 3, 4):
        style_money(ws, j, 12, 16)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    ch1 = BarChart()
    ch1.type = "col"
    ch1.title = "Profit by Enterprise"
    data = Reference(wb["Enterprise_Profit"], min_col=6, min_row=1,
                     max_row=6)
    cats = Reference(wb["Enterprise_Profit"], min_col=1, min_row=2,
                     max_row=6)
    ch1.add_data(data, titles_from_data=True)
    ch1.set_categories(cats)
    ch1.series[0].title = SeriesLabel(v="Profit")
    ch1.height, ch1.width = 7.5, 14
    ws.add_chart(ch1, "A18")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "Income vs Expense by Quarter"
    data = Reference(ws, min_col=2, min_row=11, max_row=15, max_col=3)
    cats = Reference(ws, min_col=1, min_row=12, max_row=15)
    ch2.add_data(data, titles_from_data=True)
    ch2.set_categories(cats)
    ch2.series[0].title = SeriesLabel(v="Income")
    ch2.series[1].title = SeriesLabel(v="Expense")
    ch2.height, ch2.width = 7.5, 14
    ws.add_chart(ch2, "A33")

    ch3 = PieChart()
    ch3.title = "Expenses by Category"
    data = Reference(wb["Expense_Category"], min_col=6, min_row=1,
                     max_row=9)
    cats = Reference(wb["Expense_Category"], min_col=1, min_row=2,
                     max_row=9)
    ch3.add_data(data, titles_from_data=True)
    ch3.set_categories(cats)
    ch3.series[0].title = SeriesLabel(v="Expenses")
    ch3.height, ch3.width = 7.5, 14
    ws.add_chart(ch3, "A48")

    # ---- Checks ----
    ws = wb.create_sheet("Checks")
    header_row(ws, 1, ["Check", "Difference / Value", "Pass?"])
    checks = [
        ("Dashboard income = Income_Activity total",
         "=Dashboard!B4-Income_Activity!F12"),
        ("Dashboard expenses = Expense_Category total",
         "=Dashboard!B5-Expense_Category!F10"),
        ("Dashboard net = Enterprise_Profit total profit",
         "=Dashboard!B6-Enterprise_Profit!F7"),
        ("Dashboard income = quarterly table total",
         "=Dashboard!B4-Dashboard!B16"),
        ("Dashboard expenses = quarterly table total",
         "=Dashboard!B5-Dashboard!C16"),
        ("Income rows = 172", f"=COUNTA({INC}!$A$2:$A$1000)-172"),
        ("Expense rows = 164", f"=COUNTA({EXP}!$A$2:$A$1000)-164"),
        ("No negative amounts (min income >= 0)",
         f"=MIN({INC}!$G$2:$G$1000)"),
    ]
    for i, (label, formula) in enumerate(checks, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=formula)
        if i <= 6:
            ws.cell(row=i, column=3).value = f"=ABS(B{i})<0.01"
        elif i <= 8:
            ws.cell(row=i, column=3).value = f"=B{i}=0"
        else:
            ws.cell(row=i, column=3).value = f"=B{i}>=0"
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 10

    out = os.path.join(HERE, "farmpulse.xlsx")
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
