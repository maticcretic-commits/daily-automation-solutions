#!/usr/bin/env python3
"""Build statementpulse.xlsx from data/*.csv — every figure a live formula.

Sheets: Guide, Dashboard, Statements, Ratios, Trend, Raw_Statements,
Raw_Monthly, Checks.
Rules honoured (learned the hard way): plain A1 ranges only, range
literals repeat the column letter on both ends ($D$2:$D$72), no Table
objects, chart series titles come from header cells (titles_from_data),
never plain strings. FY2024 change-in-working-capital formulas compare
against the "Opening" rows stored in Raw_Statements.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart
from openpyxl.chart.data_source import NumDataSource  # noqa: F401
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "statementpulse.xlsx")

NAVY = "1F3864"
TEAL = "2E75B6"
RED = "C00000"

YEARS = ["FY2024", "FY2025", "FY2026"]
YEAR_COLS = ["B", "C", "D"]


def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.reader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)


def sumifs(raw_last, line_item, year, row):
    """SUMIFS over Raw_Statements for one line item + year, at given row.

    The line-item label lives in Statements column A; the year label is a
    literal so the formula stays readable in the sheet.
    """
    return (f"=SUMIFS(Raw_Statements!$D$2:$D${raw_last},"
            f"Raw_Statements!$B$2:$B${raw_last},$A{row},"
            f"Raw_Statements!$C$2:$C${raw_last},\"{year}\")")


def main():
    st = read_csv("statements.csv")
    mo = read_csv("monthly.csv")
    st_header, st_rows = st[0], st[1:]
    mo_header, mo_rows = mo[0], mo[1:]
    raw_last = len(st_rows) + 1
    mo_last = len(mo_rows) + 1

    wb = Workbook()

    # ---- Guide ----
    ws = wb.active
    ws.title = "Guide"
    guide = [
        "StatementPulse — Financial Statement Analysis (DEMO)",
        "Built around: Financial Statement Analysis Needed ASAP (Freelancer, Brisbane client).",
        "Data: SYNTHETIC — fictional Sharma Retail Pvt. Ltd., Patna; FY2024-FY2026, seeded.",
        "",
        "What this answers: is the business profitable, liquid, and generating cash?",
        "Statements rebuilds the P&L, balance sheet and indirect cash flow from the raw",
        "line items — every subtotal is a formula, nothing is hard-coded. Ratios benchmarks",
        "each ratio against a plain threshold and flags FY2026 exceptions. Trend shows the",
        "FY2026 monthly build-up; its totals tie exactly to the annual statements.",
        "",
        "How to use:",
        "1. Replace the Amount column in Raw_Statements / rows in Raw_Monthly with real data,",
        "   keeping the same Section / LineItem / Year labels.",
        "2. Statements, Ratios, Trend, Dashboard and Checks update by formula.",
        "3. Power BI version: see powerbi/MEASURES.md for the DAX mapping.",
        "",
        "Sheets: Dashboard (KPIs + charts) | Statements | Ratios | Trend |",
        "Raw_Statements | Raw_Monthly | Checks (must all read TRUE).",
    ]
    for i, line in enumerate(guide, start=1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 95
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)

    # ---- Raw_Statements ----
    ws = wb.create_sheet("Raw_Statements")
    ws.append(st_header)
    style_header(ws, 1, len(st_header))
    for r in st_rows:
        ws.append([r[0], r[1], r[2], int(r[3])])
    ws.auto_filter.ref = f"A1:D{raw_last}"
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["D"].number_format = "#,##0"

    # ---- Raw_Monthly ----
    ws = wb.create_sheet("Raw_Monthly")
    ws.append(mo_header)
    style_header(ws, 1, len(mo_header))
    for r in mo_rows:
        ws.append([r[0]] + [int(v) for v in r[1:]])
    ws.auto_filter.ref = f"A1:G{mo_last}"
    ws.freeze_panes = "A2"
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 10
    for col in "BCDEFG":
        ws.column_dimensions[col].number_format = "#,##0"

    # ---- Statements ----
    ws = wb.create_sheet("Statements")
    ws.append(["Line item", "FY2024", "FY2025", "FY2026"])
    style_header(ws, 1, 4)
    ws.column_dimensions["A"].width = 42
    for col in "BCD":
        ws.column_dimensions[col].width = 18
        ws.column_dimensions[col].number_format = "#,##0"

    # (label, kind) — kind: "value" (SUMIFS), "formula" (filled below),
    # "section" (header row). Row numbers are recorded as we go.
    layout = [
        ("PROFIT & LOSS", "section"),
        ("Revenue", "value"),
        ("Cost of goods sold", "value"),
        ("Gross profit", "formula"),
        ("Operating expenses", "value"),
        ("EBITDA", "formula"),
        ("Depreciation", "value"),
        ("EBIT", "formula"),
        ("Interest expense", "value"),
        ("Profit before tax", "formula"),
        ("Tax expense", "value"),
        ("Profit after tax", "formula"),
        ("BALANCE SHEET", "section"),
        ("Cash and bank", "value"),
        ("Trade receivables", "value"),
        ("Inventories", "value"),
        ("Other current assets", "value"),
        ("Total current assets", "formula"),
        ("Property plant and equipment (net)", "value"),
        ("Other non-current assets", "value"),
        ("Total non-current assets", "formula"),
        ("TOTAL ASSETS", "formula"),
        ("Trade payables", "value"),
        ("Short-term borrowings", "value"),
        ("Other current liabilities", "value"),
        ("Total current liabilities", "formula"),
        ("Long-term borrowings", "value"),
        ("Share capital", "value"),
        ("Retained earnings", "value"),
        ("Total equity", "formula"),
        ("TOTAL LIABILITIES + EQUITY", "formula"),
        ("Balance check (assets - L+E; must be 0)", "formula"),
        ("CASH FLOW (indirect)", "section"),
        ("Profit after tax", "ref"),
        ("Add: Depreciation", "ref"),
        ("Less: increase in receivables", "delta"),
        ("Less: increase in inventories", "delta"),
        ("Add: increase in payables", "delta"),
        ("Working capital: other current items", "delta2"),
        ("Cash flow from operations", "formula"),
        ("Capital expenditure", "cfcapex"),
        ("Change in borrowings", "delta3"),
        ("Dividends paid", "cfdiv"),
        ("Cash flow from financing", "formula"),
        ("Net change in cash", "formula"),
        ("Opening cash", "opencash"),
        ("Closing cash", "formula"),
        ("Closing vs balance-sheet cash (must be 0)", "formula"),
    ]
    row_of = {}
    kind_by_row = {}
    for label, kind in layout:
        r = ws.max_row + 1
        row_of[label] = r
        kind_by_row[r] = kind
        ws.cell(row=r, column=1, value=label)
        if kind == "section":
            for c in range(1, 5):
                ws.cell(row=r, column=c).font = Font(bold=True, color=NAVY)
    R = row_of
    # "Profit after tax" occurs twice: the P&L total (first) and the
    # cash-flow ref row (last). Keep them distinct for formula targets.
    pl_pat_row = next(r for r, (lab, k) in
                      zip(kind_by_row, layout) if lab == "Profit after tax")
    cf_pat_row = max(r for r, (lab, k) in
                     zip(kind_by_row, layout) if lab == "Profit after tax")

    def col_formula(row, col_idx):
        """Formula for a value/derived cell in Statements at (row, col)."""
        col = YEAR_COLS[col_idx]
        label = ws.cell(row=row, column=1).value
        kind = kind_by_row[row]
        prev = YEAR_COLS[col_idx - 1] if col_idx > 0 else None

        def val(line, year):
            return (f"=SUMIFS(Raw_Statements!$D$2:$D${raw_last},"
                    f"Raw_Statements!$B$2:$B${raw_last},\"{line}\","
                    f"Raw_Statements!$C$2:$C${raw_last},\"{year}\")")

        def delta(line, sign):
            # Change vs prior column; FY2024 compares to "Opening" rows.
            if col_idx == 0:
                return f"={sign}({col}{row_of[line]}-{val(line, 'Opening')[1:]})"
            return f"={sign}({col}{row_of[line]}-{prev}{row_of[line]})"

        if kind == "value":
            return val(label, YEARS[col_idx])
        if kind == "ref":
            if label == "Profit after tax":
                return f"={col}{pl_pat_row}"
            return f"={col}{row_of['Depreciation']}"
        if kind == "delta":
            line = {"Less: increase in receivables": "Trade receivables",
                    "Less: increase in inventories": "Inventories",
                    "Add: increase in payables": "Trade payables"}[label]
            return delta(line, "-") if "Less" in label else delta(line, "+")
        if kind == "delta2":
            # -(ΔOther CA) + (ΔOther CL)
            def d2(line, sign):
                if col_idx == 0:
                    return f"{sign}({col}{row_of[line]}-{val(line, 'Opening')[1:]})"
                return f"{sign}({col}{row_of[line]}-{prev}{row_of[line]})"
            return ("=" + d2("Other current assets", "-")
                    + "+" + d2("Other current liabilities", "+"))
        if kind == "cfcapex":
            return f"=-{val('Capital expenditure', YEARS[col_idx])[1:]}"
        if kind == "cfdiv":
            return f"=-{val('Dividends paid', YEARS[col_idx])[1:]}"
        if kind == "delta3":
            # Δ short-term + Δ long-term borrowings
            def d3(line):
                if col_idx == 0:
                    return f"({col}{row_of[line]}-{val(line, 'Opening')[1:]})"
                return f"({col}{row_of[line]}-{prev}{row_of[line]})"
            return "=" + d3("Short-term borrowings") + "+" + d3("Long-term borrowings")
        if kind == "opencash":
            if col_idx == 0:
                return val("Cash and bank", "Opening")
            return f"={prev}{row_of['Closing cash']}"
        # formula kinds: subtotals / totals
        f = {
            "Gross profit": f"={col}{R['Revenue']}-{col}{R['Cost of goods sold']}",
            "EBITDA": f"={col}{R['Gross profit']}-{col}{R['Operating expenses']}",
            "EBIT": f"={col}{R['EBITDA']}-{col}{R['Depreciation']}",
            "Profit before tax": f"={col}{R['EBIT']}-{col}{R['Interest expense']}",
            "Profit after tax": f"={col}{R['Profit before tax']}-{col}{R['Tax expense']}",
            "Total current assets": f"=SUM({col}{R['Cash and bank']}:{col}{R['Other current assets']})",
            "Total non-current assets": f"={col}{R['Property plant and equipment (net)']}+{col}{R['Other non-current assets']}",
            "TOTAL ASSETS": f"={col}{R['Total current assets']}+{col}{R['Total non-current assets']}",
            "Total current liabilities": f"=SUM({col}{R['Trade payables']}:{col}{R['Other current liabilities']})",
            "Total equity": f"={col}{R['Share capital']}+{col}{R['Retained earnings']}",
            "TOTAL LIABILITIES + EQUITY": f"={col}{R['Total current liabilities']}+{col}{R['Long-term borrowings']}+{col}{R['Total equity']}",
            "Balance check (assets - L+E; must be 0)": f"={col}{R['TOTAL ASSETS']}-{col}{R['TOTAL LIABILITIES + EQUITY']}",
            "Cash flow from operations": f"=SUM({col}{R['Profit after tax']}:{col}{R['Working capital: other current items']})",
            "Cash flow from financing": f"={col}{R['Change in borrowings']}+{col}{R['Dividends paid']}",
            "Net change in cash": f"={col}{R['Cash flow from operations']}+{col}{R['Capital expenditure']}+{col}{R['Cash flow from financing']}",
            "Closing cash": f"={col}{R['Opening cash']}+{col}{R['Net change in cash']}",
            "Closing vs balance-sheet cash (must be 0)": f"={col}{R['Closing cash']}-{col}{R['Cash and bank']}",
        }
        return f[label]

    for idx, (label, kind) in enumerate(layout):
        if kind == "section":
            continue
        r = 2 + idx  # header row 1, first layout row 2 (label-keyed maps
        # lose the first "Profit after tax" row to its cash-flow namesake)
        for ci in range(3):
            ws.cell(row=r, column=2 + ci, value=col_formula(r, ci))
    ws.freeze_panes = "A2"

    # ---- Ratios ----
    ws = wb.create_sheet("Ratios")
    ws.append(["Ratio", "FY2024", "FY2025", "FY2026", "Benchmark", "FY2026 status"])
    style_header(ws, 1, 6)
    ws.column_dimensions["A"].width = 34
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 14
    S = "Statements!"
    ratio_defs = [
        ("Current ratio (CA / CL)", lambda c: f"={S}{c}{R['Total current assets']}/{S}{c}{R['Total current liabilities']}", 1.5, "ge", "0.00"),
        ("Quick ratio ((CA - inventory) / CL)", lambda c: f"=({S}{c}{R['Total current assets']}-{S}{c}{R['Inventories']})/{S}{c}{R['Total current liabilities']}", 1.0, "ge", "0.00"),
        ("Debt-to-equity ((ST + LT debt) / equity)", lambda c: f"=({S}{c}{R['Short-term borrowings']}+{S}{c}{R['Long-term borrowings']})/{S}{c}{R['Total equity']}", 1.0, "le", "0.00"),
        ("Interest coverage (EBIT / interest)", lambda c: f"={S}{c}{R['EBIT']}/{S}{c}{R['Interest expense']}", 3.0, "ge", "0.00"),
        ("Gross margin", lambda c: f"={S}{c}{R['Gross profit']}/{S}{c}{R['Revenue']}", None, None, "0.0%"),
        ("Net margin (PAT / revenue)", lambda c: f"={S}{c}{pl_pat_row}/{S}{c}{R['Revenue']}", 0.08, "ge", "0.0%"),
        ("Return on assets (PAT / total assets)", lambda c: f"={S}{c}{pl_pat_row}/{S}{c}{R['TOTAL ASSETS']}", None, None, "0.0%"),
        ("Return on equity (PAT / equity)", lambda c: f"={S}{c}{pl_pat_row}/{S}{c}{R['Total equity']}", 0.15, "ge", "0.0%"),
        ("Receivables days", lambda c: f"={S}{c}{R['Trade receivables']}/{S}{c}{R['Revenue']}*365", 60, "le", "0"),
        ("Inventory days", lambda c: f"={S}{c}{R['Inventories']}/{S}{c}{R['Cost of goods sold']}*365", 60, "le", "0"),
        ("Payables days", lambda c: f"={S}{c}{R['Trade payables']}/{S}{c}{R['Cost of goods sold']}*365", 45, "le", "0"),
        ("Operating cash conversion (CFO / PAT)", lambda c: f"={S}{c}{R['Cash flow from operations']}/{S}{c}{pl_pat_row}", 1.0, "ge", "0.00"),
        ("Working capital (CA - CL, Rs)", lambda c: f"={S}{c}{R['Total current assets']}-{S}{c}{R['Total current liabilities']}", None, None, "#,##0"),
    ]
    for i, (label, fn, bench, direction, fmt) in enumerate(ratio_defs):
        r = i + 2
        ws.cell(row=r, column=1, value=label)
        for ci, col in enumerate(YEAR_COLS):
            ws.cell(row=r, column=2 + ci, value=fn(col))
            ws.cell(row=r, column=2 + ci).number_format = fmt
        if bench is None:
            ws.cell(row=r, column=5, value="—")
            ws.cell(row=r, column=6, value="—")
        else:
            ws.cell(row=r, column=5, value=bench)
            ws.cell(row=r, column=5).number_format = fmt
            op = ">=" if direction == "ge" else "<="
            ws.cell(row=r, column=6, value=f'=IF(D{r}{op}E{r},"OK","FLAG")')
    flag_row = len(ratio_defs) + 2
    ws.cell(row=flag_row, column=1, value="FY2026 flags (count of FLAG above)")
    ws.cell(row=flag_row, column=2,
            value=f'=COUNTIF(F2:F{len(ratio_defs) + 1},"FLAG")')
    ws.cell(row=flag_row, column=1).font = Font(bold=True, color=NAVY)
    ws.freeze_panes = "A2"

    # ---- Trend ----
    ws = wb.create_sheet("Trend")
    headers = ["Month", "Revenue", "COGS", "Gross profit", "Opex",
               "Operating surplus", "Margin%", "Receivables", "Inventory",
               "Payables", "Working capital"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.column_dimensions["A"].width = 10
    for col in "BCDEFHIJK":
        ws.column_dimensions[col].width = 15
    for i, mrow in enumerate(mo_rows):
        r = i + 2
        ws.cell(row=r, column=1, value=mrow[0])

        def msum(col_letter, rr):
            return (f"=SUMIFS(Raw_Monthly!${col_letter}$2:${col_letter}${mo_last},"
                    f"Raw_Monthly!$A$2:$A${mo_last},$A{rr})")
        ws.cell(row=r, column=2, value=msum("B", r))
        ws.cell(row=r, column=3, value=msum("C", r))
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
        ws.cell(row=r, column=5, value=msum("D", r))
        ws.cell(row=r, column=6, value=f"=B{r}-C{r}-E{r}")
        ws.cell(row=r, column=7, value=f"=IF(B{r}=0,0,F{r}/B{r})")
        ws.cell(row=r, column=8, value=msum("E", r))
        ws.cell(row=r, column=9, value=msum("F", r))
        ws.cell(row=r, column=10, value=msum("G", r))
        ws.cell(row=r, column=11, value=f"=H{r}+I{r}-J{r}")
    tot = len(mo_rows) + 2
    ws.cell(row=tot, column=1, value="TOTAL / Mar-26")
    for ci, col in enumerate(["B", "C", "D", "E", "F"]):
        ws.cell(row=tot, column=2 + ci,
                value=f"=SUM({col}2:{col}{tot - 1})")
    ws.cell(row=tot, column=7, value=f"=IF(B{tot}=0,0,F{tot}/B{tot})")
    for ci, col in enumerate(["H", "I", "J", "K"]):
        ws.cell(row=tot, column=8 + ci, value=f"={col}{tot - 1}")
    for col in "BCDEFHIJK":
        ws.column_dimensions[col].number_format = "#,##0"
    ws.column_dimensions["G"].number_format = "0.0%"
    ws.freeze_panes = "A2"

    # ---- Dashboard ----
    ws = wb.create_sheet("Dashboard", 1)  # after Guide
    ws["A1"] = "StatementPulse — Sharma Retail Pvt. Ltd. (SYNTHETIC DEMO, FY2024-FY2026)"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A3"] = "KPI (FY2026 unless stated)"
    ws["B3"] = "Value"
    style_header(ws, 3, 2)
    cards = [
        ("Revenue (FY2026)", f"={S}D{R['Revenue']}", "#,##0"),
        ("Profit after tax (FY2026)", f"={S}D{pl_pat_row}", "#,##0"),
        ("Net margin (FY2026)", "=Ratios!D7", "0.0%"),
        ("Current ratio (FY2026)", "=Ratios!D2", "0.00"),
        ("Receivables days (FY2026)", "=Ratios!D10", "0"),
        ("Operating cash conversion (FY2026)", "=Ratios!D13", "0.00"),
    ]
    for i, (label, formula, fmt) in enumerate(cards):
        r = 4 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=formula)
        ws.cell(row=r, column=2).number_format = fmt
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    ws.cell(row=11, column=1, value="FY2026 ratio flags (see Ratios sheet)")
    ws.cell(row=11, column=2, value=f"=Ratios!B{flag_row}")
    ws.cell(row=11, column=1).font = Font(bold=True, color=NAVY)

    # Helper table for the year chart (titles from header cells).
    ws["A13"] = "Year"
    ws["B13"] = "Revenue"
    ws["C13"] = "Profit after tax"
    style_header(ws, 13, 3)
    for i, (year, col) in enumerate(zip(YEARS, YEAR_COLS)):
        r = 14 + i
        ws.cell(row=r, column=1, value=year)
        ws.cell(row=r, column=2, value=f"={S}{col}{R['Revenue']}")
        ws.cell(row=r, column=3, value=f"={S}{col}{R['Profit after tax']}")
    bar = BarChart()
    bar.type = "col"
    bar.title = "Revenue vs profit after tax by year"
    bar.add_data(__import__("openpyxl.chart").chart.Reference(
        ws, min_col=2, min_row=13, max_col=3, max_row=16), titles_from_data=True)
    bar.set_categories(__import__("openpyxl.chart").chart.Reference(
        ws, min_col=1, min_row=14, max_row=16))
    bar.height, bar.width = 7.5, 14
    ws.add_chart(bar, "A18")

    line = LineChart()
    line.title = "Monthly revenue, FY2026"
    line.add_data(__import__("openpyxl.chart").chart.Reference(
        wb["Trend"], min_col=2, min_row=1, max_row=13), titles_from_data=True)
    line.set_categories(__import__("openpyxl.chart").chart.Reference(
        wb["Trend"], min_col=1, min_row=2, max_row=13))
    line.height, line.width = 7.5, 14
    ws.add_chart(line, "A34")

    # Helper table for the cost-mix pie (FY2026).
    ws["A50"] = "Cost item (FY2026)"
    ws["B50"] = "Amount"
    style_header(ws, 50, 2)
    pie_items = [("Cost of goods sold", R["Cost of goods sold"]),
                 ("Operating expenses", R["Operating expenses"]),
                 ("Depreciation", R["Depreciation"]),
                 ("Interest expense", R["Interest expense"]),
                 ("Tax expense", R["Tax expense"])]
    for i, (label, srow) in enumerate(pie_items):
        r = 51 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f"={S}D{srow}")
    pie = PieChart()
    pie.title = "FY2026 cost mix"
    pie.add_data(__import__("openpyxl.chart").chart.Reference(
        ws, min_col=2, min_row=50, max_row=55), titles_from_data=True)
    pie.set_categories(__import__("openpyxl.chart").chart.Reference(
        ws, min_col=1, min_row=51, max_row=55))
    pie.height, pie.width = 7.5, 14
    ws.add_chart(pie, "A57")

    # ---- Checks ----
    ws = wb.create_sheet("Checks")
    ws.append(["Check", "Result"])
    style_header(ws, 1, 2)
    div_fy26 = (f"SUMIFS(Raw_Statements!$D$2:$D${raw_last},"
                f"Raw_Statements!$B$2:$B${raw_last},\"Dividends paid\","
                f"Raw_Statements!$C$2:$C${raw_last},\"FY2026\")")
    checks = [
        ("Balance sheet balances, FY2024",
         f"=ABS({S}B{R['Balance check (assets - L+E; must be 0)']})<1"),
        ("Balance sheet balances, FY2025",
         f"=ABS({S}C{R['Balance check (assets - L+E; must be 0)']})<1"),
        ("Balance sheet balances, FY2026",
         f"=ABS({S}D{R['Balance check (assets - L+E; must be 0)']})<1"),
        ("Cash-flow closing cash = balance-sheet cash, FY2024",
         f"=ABS({S}B{R['Closing vs balance-sheet cash (must be 0)']})<1"),
        ("Cash-flow closing cash = balance-sheet cash, FY2025",
         f"=ABS({S}C{R['Closing vs balance-sheet cash (must be 0)']})<1"),
        ("Cash-flow closing cash = balance-sheet cash, FY2026",
         f"=ABS({S}D{R['Closing vs balance-sheet cash (must be 0)']})<1"),
        ("Monthly revenue total = FY2026 revenue",
         f"=ABS(Trend!B{tot}-{S}D{R['Revenue']})<1"),
        ("Retained earnings roll-forward, FY2026",
         f"=ABS({S}D{R['Retained earnings']}-({S}C{R['Retained earnings']}"
         f"+{S}D{R['Profit after tax']}-{div_fy26}))<1"),
    ]
    for label, formula in checks:
        ws.append([label, formula])
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 12

    wb.save(OUT)
    print(f"wrote {OUT} ({len(st_rows)} statement rows, {len(mo_rows)} months)")


if __name__ == "__main__":
    main()
