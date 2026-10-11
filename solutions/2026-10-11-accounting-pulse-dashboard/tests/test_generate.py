"""Tests for the AccountingPulse generator.

Includes a LibreOffice headless recalc: the workbook is converted (forcing a
full formula recalc) and key figures are compared against ground truth
computed independently from the CSV inputs.
"""
import csv
import datetime as dt
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent.parent
DATA = HERE / "data"
XLSX = HERE / "accounting-pulse.xlsx"
RECALC_DIR = HERE / "tests" / "_recalc"
PERIOD_START = dt.date(2026, 4, 1)
PERIOD_END = dt.date(2026, 10, 1)
AS_OF = dt.date(2026, 9, 30)
BOOK_OPEN = 500000.0
TOL = 0.02
BUCKETS = ["Current", "1-30", "31-60", "61-90", "90+"]


def load_csv(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def bucket(dd):
    return ("Current" if dd <= 0 else "1-30" if dd <= 30 else "31-60"
            if dd <= 60 else "61-90" if dd <= 90 else "90+")


@pytest.fixture(scope="module")
def ground():
    ap = load_csv("ap_ledger.csv")
    ar = load_csv("ar_ledger.csv")
    pay = load_csv("payroll.csv")
    st = load_csv("bank_statement.csv")
    in_p = lambda rows, dk: [r for r in rows
                             if PERIOD_START <= dt.date.fromisoformat(r[dk]) < PERIOD_END]
    ap, ar, pay, st = in_p(ap, "date"), in_p(ar, "date"), in_p(pay, "pay_date"), in_p(st, "date")

    ap_age = {b: 0.0 for b in BUCKETS}
    for r in ap:
        if r["paid_date"]:
            continue
        dd = (AS_OF - dt.date.fromisoformat(r["due_date"])).days
        ap_age[bucket(dd)] += float(r["amount"])
    ar_age = {b: 0.0 for b in BUCKETS}
    for r in ar:
        if r["paid_date"]:
            continue
        dd = (AS_OF - dt.date.fromisoformat(r["due_date"])).days
        ar_age[bucket(dd)] += float(r["amount"])

    receipts = sum(float(r["amount"]) for r in ar if r["paid_date"])
    ap_paid = sum(float(r["amount"]) for r in ap if r["paid_date"])
    pay_net = sum(float(r["net"]) for r in pay)
    book = BOOK_OPEN + receipts - ap_paid - pay_net
    interest = sum(float(r["deposit"]) for r in st if r["ref"].startswith("INT-"))
    charges = sum(float(r["withdrawal"]) for r in st if r["ref"].startswith("CHG-"))
    adj_book = book + interest - charges
    stmt_bal = sum(float(r["deposit"]) for r in st) - sum(float(r["withdrawal"]) for r in st)
    dit = sum(float(r["amount"]) for r in ar
              if r["paid_date"] and r["in_transit"] == "1")
    oc = sum(float(r["amount"]) for r in ap
             if r["paid_date"] and r["pay_mode"] == "cheque" and r["cleared"] == "0")
    adj_bank = stmt_bal + dit - oc

    rev = sum(float(r["amount"]) for r in ar)
    pur = sum(float(r["amount"]) for r in ap if r["category"] == "Purchases")
    oth = sum(float(r["amount"]) for r in ap if r["category"] != "Purchases")
    ni = rev - pur - pay_net - oth
    depts = {}
    for r in pay:
        depts[r["department"]] = depts.get(r["department"], 0.0) + float(r["net"])
    return {"ap_age": ap_age, "ar_age": ar_age,
            "ap_out": sum(ap_age.values()), "ar_out": sum(ar_age.values()),
            "book": book, "adj_book": adj_book, "stmt_bal": stmt_bal,
            "dit": dit, "oc": oc, "adj_bank": adj_bank,
            "interest": interest, "charges": charges,
            "rev": rev, "pur": pur, "pay_net": pay_net, "oth": oth, "ni": ni,
            "depts": depts}


@pytest.fixture(scope="module")
def recalc_wb():
    assert XLSX.exists(), "run generate.py first"
    RECALC_DIR.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         str(XLSX), "--outdir", str(RECALC_DIR)],
        check=True, capture_output=True, timeout=180)
    out = RECALC_DIR / XLSX.name
    assert out.exists(), "libreoffice recalc produced no output"
    return load_workbook(out, data_only=True)


# ---- input data ----
def test_csvs_parse():
    assert len(load_csv("ap_ledger.csv")) == 39
    assert len(load_csv("ar_ledger.csv")) == 36
    assert len(load_csv("payroll.csv")) == 48
    assert len(load_csv("bank_statement.csv")) == 107


def test_dates_sane():
    for r in load_csv("ap_ledger.csv") + load_csv("ar_ledger.csv"):
        d = dt.date.fromisoformat(r["date"])
        assert PERIOD_START <= d < PERIOD_END
        assert dt.date.fromisoformat(r["due_date"]) >= d
        if r["paid_date"]:
            assert dt.date.fromisoformat(r["paid_date"]) >= d


def test_rec_items_present():
    ar = load_csv("ar_ledger.csv")
    ap = load_csv("ap_ledger.csv")
    assert sum(1 for r in ar if r["in_transit"] == "1") >= 2
    assert sum(1 for r in ap if r["paid_date"] and r["pay_mode"] == "cheque"
               and r["cleared"] == "0") >= 2


# ---- workbook structure ----
def test_sheets_present():
    wb = load_workbook(XLSX)
    assert wb.sheetnames == ["Raw_AP", "Raw_AR", "Payroll", "Bank", "Bank_Rec",
                             "AP_Aging", "AR_Aging", "P&L", "Dashboard"]


def test_named_ranges():
    wb = load_workbook(XLSX)
    names = set(wb.defined_names)
    assert {"AsOf", "BookOpen", "PeriodStart", "PeriodEnd"} <= names


def test_formulas_are_live_not_hardcoded():
    wb = load_workbook(XLSX)
    assert str(wb["Raw_AP"]["J2"].value).startswith("=")     # days overdue
    assert "IF(" in str(wb["Raw_AP"]["K2"].value)            # aging bucket
    assert str(wb["Bank_Rec"]["B6"].value).startswith("=")   # adjusted bank
    assert str(wb["Bank_Rec"]["B11"].value).startswith("=")  # adjusted book
    assert "SUMIFS" in str(wb["AP_Aging"]["B3"].value)
    assert "SUMIFS" in str(wb["P&L"]["B4"].value)            # monthly revenue
    assert str(wb["P&L"]["H10"].value).startswith("=")       # net income total


def test_dashboard_has_three_charts():
    wb = load_workbook(XLSX)
    assert len(wb["Dashboard"]._charts) == 3


# ---- recalculated values vs ground truth ----
def test_aging_totals(recalc_wb, ground):
    for sheet, key in (("AP_Aging", "ap_age"), ("AR_Aging", "ar_age")):
        ws = recalc_wb[sheet]
        for i, b in enumerate(BUCKETS):
            assert abs(ws.cell(row=3 + i, column=2).value
                       - ground[key][b]) < TOL, (sheet, b)
        assert abs(ws["B8"].value - sum(ground[key].values())) < TOL
        assert "OK" in str(ws["B9"].value)


def test_bank_rec(recalc_wb, ground):
    ws = recalc_wb["Bank_Rec"]
    assert abs(ws["B3"].value - ground["stmt_bal"]) < TOL
    assert abs(ws["B4"].value - ground["dit"]) < TOL
    assert abs(ws["B5"].value - ground["oc"]) < TOL
    assert abs(ws["B6"].value - ground["adj_bank"]) < TOL
    assert abs(ws["B8"].value - ground["book"]) < TOL
    assert abs(ws["B9"].value - ground["interest"]) < TOL
    assert abs(ws["B10"].value - ground["charges"]) < TOL
    assert abs(ws["B11"].value - ground["adj_book"]) < TOL
    assert abs(ws["B13"].value) < TOL
    assert "OK" in str(ws["B14"].value)


def test_pl_totals(recalc_wb, ground):
    pl = recalc_wb["P&L"]
    assert abs(pl["H4"].value - ground["rev"]) < TOL
    assert abs(pl["H5"].value - ground["pur"]) < TOL
    assert abs(pl["H7"].value - ground["pay_net"]) < TOL
    assert abs(pl["H8"].value - ground["oth"]) < TOL
    assert abs(pl["H10"].value - ground["ni"]) < TOL
    assert abs(pl["H6"].value - (ground["rev"] - ground["pur"])) < TOL


def test_pl_monthly_adds_to_total(recalc_wb):
    pl = recalc_wb["P&L"]
    for r in (4, 10):  # revenue, net income
        months = [pl.cell(row=r, column=j).value or 0 for j in range(2, 8)]
        assert abs(sum(months) - pl.cell(row=r, column=8).value) < TOL


def test_dashboard_kpis(recalc_wb, ground):
    d = recalc_wb["Dashboard"]
    assert abs(d["B7"].value - ground["ap_out"]) < TOL
    assert abs(d["B8"].value - ground["ar_out"]) < TOL
    assert abs(d["B9"].value - ground["adj_bank"]) < TOL
    assert abs(d["B10"].value) < TOL
    assert "OK" in str(d["B11"].value)
    assert abs(d["B12"].value - ground["pay_net"] / 6) < TOL
    assert abs(d["B13"].value - ground["ni"]) < TOL
    assert abs(d["B14"].value - ground["ar_age"]["90+"]) < TOL


def test_dept_payroll_adds_up(recalc_wb, ground):
    d = recalc_wb["Dashboard"]
    got = sum(d.cell(row=r, column=8).value or 0
              for r in range(7, 7 + len(ground["depts"])))
    assert abs(got - ground["pay_net"]) < TOL


def test_no_formula_errors(recalc_wb):
    bad = []
    for ws in recalc_wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not bad, f"formula errors: {bad[:10]}"
