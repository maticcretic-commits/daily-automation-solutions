#!/usr/bin/env python3
"""FlowForge verification: LibreOffice headless recalc of the built workbook
vs independent Python ground truth, plus structural checks on the raw file.

Run: python3 tests/test_build.py   (from /tmp/cashflow-forge)
"""
import csv
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).parent.parent
DATA = ROOT / "data"
XLSX = ROOT / "cashflow-forecast.xlsx"

from openpyxl import load_workbook  # noqa: E402


def ground_truth():
    def rc(name):
        with open(DATA / name) as f:
            return list(csv.DictReader(f))
    inflows, outflows = rc("history_inflows.csv"), rc("history_outflows.csv")
    inc = ["sales_receipts", "gst_refund", "other_income"]
    outc = ["supplier_payments", "payroll", "rent", "logistics", "gst_tax", "loan_emi"]
    tin = [sum(int(r[c]) for c in inc) for r in inflows]
    tout = [sum(int(r[c]) for c in outc) for r in outflows]
    bal, closes = 4_250_000, []
    for a, b in zip(tin, tout):
        bal += a - b
        closes.append(bal)
    avg4_in = sum(tin[-4:]) / 4
    avg4_out = sum(tout[-4:]) / 4
    fc, b = [], closes[-1]
    for n in range(1, 14):
        fi = round(avg4_in * 1.02 ** n)
        fo = round(avg4_out * 1.015 ** n)
        b += fi - fo
        fc.append(b)
    return {"hist_closing": closes, "fc_closing": fc,
            "min_cash": min(fc), "fc_open": closes[-1]}


GT = ground_truth()


def recalc_values():
    """Headless LibreOffice recalc; read back with data_only."""
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "src.xlsx"
    shutil.copy(XLSX, src)  # never convert the source in place (basename quirk)
    outdir = tmp / "out"
    outdir.mkdir()
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
         str(outdir), str(src)],
        capture_output=True, text=True, timeout=180)
    if r.returncode != 0:
        raise RuntimeError(f"soffice failed: {r.stderr[-2000:]}")
    got = list(outdir.glob("*.xlsx"))
    assert got, "no converted file"
    wb = load_workbook(got[0], data_only=True)
    vals = {}
    h, f = wb["History"], wb["Forecast"]
    vals["hist_closing"] = [h.cell(row=i + 2, column=15).value for i in range(26)]
    vals["fc_closing"] = [f.cell(row=i + 2, column=6).value for i in range(13)]
    vals["checks"] = [wb["Checks"].cell(row=i + 3, column=2).value for i in range(8)]
    vals["kpi_min"] = wb["Dashboard"].cell(row=6, column=2).value
    vals["kpi_close"] = wb["Dashboard"].cell(row=5, column=2).value
    vals["kpi_below"] = wb["Dashboard"].cell(row=7, column=2).value
    return vals


class TestStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(XLSX)

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames,
                         ["Guide", "Assumptions", "History", "Forecast",
                          "Dashboard", "Checks", "Report"])

    def test_charts_present(self):
        self.assertEqual(len(self.wb["Dashboard"]._charts), 3)

    def test_history_formulas_live(self):
        h = self.wb["History"]
        # totals, net, opening, closing, variance, alert are all formulas
        self.assertTrue(str(h["E2"].value).startswith("="))
        self.assertTrue(str(h["L2"].value).startswith("="))
        self.assertTrue(str(h["M2"].value).startswith("="))
        self.assertTrue(str(h["N3"].value).startswith("="))
        self.assertTrue(str(h["O2"].value).startswith("="))
        self.assertTrue(str(h["R2"].value).startswith("="))
        self.assertTrue(str(h["T2"].value).startswith("="))

    def test_forecast_formulas_live(self):
        f = self.wb["Forecast"]
        self.assertTrue(str(f["B2"].value).startswith("="))
        self.assertTrue(str(f["F2"].value).startswith("="))
        self.assertTrue(str(f["G2"].value).startswith("="))

    def test_no_hardcoded_money_on_forecast(self):
        f = self.wb["Forecast"]
        for r in range(2, 15):
            for c in "BCDEF":
                v = f[f"{c}{r}"].value
                self.assertTrue(isinstance(v, str) and v.startswith("="),
                                f"{c}{r} not a formula: {v!r}")


class TestRecalc(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.v = recalc_values()

    def test_no_errors(self):
        for name, seq in (("hist", self.v["hist_closing"]),
                          ("fc", self.v["fc_closing"])):
            for i, v in enumerate(seq):
                self.assertIsInstance(v, (int, float), f"{name}[{i}]={v!r}")

    def test_history_closing_matches(self):
        for i, (got, exp) in enumerate(zip(self.v["hist_closing"],
                                           GT["hist_closing"])):
            self.assertAlmostEqual(got, exp, delta=1, msg=f"week {i}")

    def test_forecast_closing_matches(self):
        for i, (got, exp) in enumerate(zip(self.v["fc_closing"],
                                           GT["fc_closing"])):
            self.assertAlmostEqual(got, exp, delta=1, msg=f"F{i + 1}")

    def test_all_checks_pass(self):
        self.assertEqual(self.v["checks"], ["PASS"] * 8)

    def test_dashboard_kpis(self):
        self.assertAlmostEqual(self.v["kpi_min"], GT["min_cash"], delta=1)
        self.assertAlmostEqual(self.v["kpi_close"], GT["fc_closing"][-1], delta=1)
        self.assertEqual(self.v["kpi_below"],
                         sum(1 for c in GT["fc_closing"] if c < 500_000))


if __name__ == "__main__":
    unittest.main(verbosity=2)
