#!/usr/bin/env python3
"""Build the workbook, recalc it headlessly in LibreOffice, and compare
every checked figure against the independent Python ground truth.

LO quirk honoured: --convert-to runs on a COPY in a separate input dir
with --outdir elsewhere (converting in place overwrites the input file).
Formulas are A1-style only; a formula error in the recalculated copy is a
failure, not a warning. A second regression scan rejects short-form range
literals ($H$2:$5000) in every formula cell of the builder.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import generate  # noqa: E402
import build_workbook  # noqa: E402
import main as ground_truth  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, "ops-pulse.xlsx")
ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "Err:522")
BAD_RANGE = re.compile(r"\$[A-Z]+\$\d+:\$\d+")


def close(a, b, tol):
    return abs(a - b) <= tol


def main():
    generate.main()
    build_workbook.main()
    truth = ground_truth.compute()

    with open(os.path.join(HERE, "build_workbook.py")) as f:
        src = f.read()
    bad = [m.group(0) for m in BAD_RANGE.finditer(src)]
    assert not bad, f"short-form range literals in builder: {bad[:5]}"

    tmp = tempfile.mkdtemp(prefix="opspulse-")
    indir, outdir = os.path.join(tmp, "in"), os.path.join(tmp, "out")
    os.makedirs(indir)
    os.makedirs(outdir)
    shutil.copy(XLSX, os.path.join(indir, "ops-pulse.xlsx"))
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir,
         os.path.join(indir, "ops-pulse.xlsx")],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:500]}"
    recalc = os.path.join(outdir, "ops-pulse.xlsx")
    assert os.path.exists(recalc), "LO produced no output"
    wb = load_workbook(recalc, data_only=True)

    failures = []

    def check(name, got, want, tol):
        if got is None or isinstance(got, str) or not close(got, want, tol):
            failures.append(f"{name}: got {got!r}, want {want}")

    for sheet in ("Dashboard", "Inventory", "Movements", "Raw_Shifts", "Staff",
                  "Orders", "Weekly", "Checks", "Assumptions"):
        for row in wb[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    failures.append(f"formula error {sheet}!{c.coordinate}: {c.value}")

    ws = wb["Dashboard"]
    check("Dashboard SKUs", ws["B3"].value, truth["count"], 0)
    check("Dashboard reorder", ws["B4"].value, truth["reorder_count"], 0)
    check("Dashboard stock value", ws["B5"].value, truth["stock_value"], 0.01)
    check("Dashboard stock units", ws["B6"].value, truth["stock_units"], 0.01)
    check("Dashboard orders", ws["B7"].value, truth["order_count"], 0)
    check("Dashboard order value", ws["B8"].value, truth["order_value"], 0.01)
    check("Dashboard on-time %", ws["B9"].value, truth["on_time_pct"], 1e-6)
    check("Dashboard avg days", ws["B10"].value, truth["avg_days"], 1e-6)
    check("Dashboard overtime", ws["B11"].value, truth["overtime"], 0.01)
    check("Dashboard utilisation", ws["B12"].value, truth["utilisation"], 1e-6)
    check("Dashboard reorder share", ws["B13"].value,
          truth["reorder_count"] / truth["count"], 1e-6)
    check("Dashboard standard hours", ws["B14"].value, 8.0, 1e-9)
    check("Dashboard target days", ws["B15"].value, 3, 0)

    ws = wb["Inventory"]
    for i, item in enumerate(truth["inventory"], start=2):
        check(f"Inventory {item['id']} stock", ws.cell(row=i, column=9).value,
              item["stock"], 0.01)
        check(f"Inventory {item['id']} value", ws.cell(row=i, column=11).value,
              item["value"], 0.01)
        got_status = ws.cell(row=i, column=13).value
        if got_status != item["status"]:
            failures.append(f"Inventory {item['id']} status: got {got_status!r}, "
                            f"want {item['status']!r}")

    ws = wb["Staff"]
    for i, emp in enumerate(truth["staff"], start=2):
        check(f"Staff {emp['name']} overtime", ws.cell(row=i, column=6).value,
              emp["overtime"], 0.01)
        check(f"Staff {emp['name']} utilisation", ws.cell(row=i, column=7).value,
              emp["utilisation"], 1e-6)

    ws = wb["Orders"]
    for i, o in enumerate(truth["orders"], start=2):
        check(f"Orders {o['id']} days", ws.cell(row=i, column=8).value,
              o["days"], 0)
        check(f"Orders {o['id']} value", ws.cell(row=i, column=6).value,
              o["value"], 0.01)

    ws = wb["Weekly"]
    for i, w in enumerate(truth["weekly"], start=2):
        check(f"Weekly {w['week']} orders", ws.cell(row=i, column=2).value,
              w["orders"], 0)
        check(f"Weekly {w['week']} value", ws.cell(row=i, column=3).value,
              w["value"], 0.01)
        check(f"Weekly {w['week']} on-time", ws.cell(row=i, column=4).value,
              w["on_time"], 0)
        check(f"Weekly {w['week']} on-time %", ws.cell(row=i, column=5).value,
              w["on_time_pct"], 1e-6)
        check(f"Weekly {w['week']} avg days", ws.cell(row=i, column=6).value,
              w["avg_days"], 1e-6)
        check(f"Weekly {w['week']} receipts", ws.cell(row=i, column=7).value,
              w["receipts"], 0.01)
        check(f"Weekly {w['week']} sales", ws.cell(row=i, column=8).value,
              w["sales"], 0.01)
        check(f"Weekly {w['week']} overtime", ws.cell(row=i, column=9).value,
              w["overtime"], 0.01)

    ws = wb["Checks"]
    for r in range(2, 12):
        if ws.cell(row=r, column=5).value != "PASS":
            failures.append(f"Checks row {r} not PASS: "
                            f"expected={ws.cell(row=r, column=2).value!r} "
                            f"got={ws.cell(row=r, column=3).value!r}")

    print(f"truth: skus={truth['count']} reorder={truth['reorder_count']} "
          f"stock_value=Rs {truth['stock_value']:,.0f} "
          f"orders={truth['order_count']} order_value=Rs {truth['order_value']:,.0f} "
          f"on_time={truth['on_time_pct']:.1%} avg_days={truth['avg_days']:.2f} "
          f"overtime={truth['overtime']:.1f}h utilisation={truth['utilisation']:.1%}")
    if failures:
        print(f"{len(failures)} FAILURES:")
        for f in failures:
            print(" -", f)
        sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
