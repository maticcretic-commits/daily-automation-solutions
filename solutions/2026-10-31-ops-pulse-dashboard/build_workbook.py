#!/usr/bin/env python3
"""Build ops-pulse.xlsx from data/*.csv — every figure a live formula.

Sheets: Guide, Dashboard, Inventory, Movements, Raw_Shifts, Staff, Orders,
Weekly, Assumptions, Checks.
Rules honoured (learned the hard way): plain A1 ranges only, range
literals repeat the column letter on both ends ($E$2:$E$169), no Table
objects, chart series titles come from header cells (titles_from_data),
never plain strings. Dates are stored as real date values so fulfilment
days (ShipDate - OrderDate) calculate. Standards live on Assumptions:
8-hour standard shift, on-time = 3 days or fewer.
"""
import csv
import datetime
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter  # noqa: F401  (kept for parity)

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "ops-pulse.xlsx")

NAVY = "1F3864"
TEAL = "2E75B6"  # noqa: F841
LIGHT = "D9E2F3"

N_PRODUCTS = 24
PROD_LAST = N_PRODUCTS + 1          # 25
PROD_TOTAL = PROD_LAST + 1          # 26
N_MOVEMENTS = 168
MOV_LAST = N_MOVEMENTS + 1          # 169
N_SHIFTS = 120
SHIFT_LAST = N_SHIFTS + 1           # 121
N_STAFF = 10
STAFF_LAST = N_STAFF + 1            # 11
STAFF_TOTAL = STAFF_LAST + 1        # 12
N_ORDERS = 96
ORDER_LAST = N_ORDERS + 1           # 97
ORDER_TOTAL = ORDER_LAST + 1        # 98
N_WEEKS = 8
WEEK_LAST = N_WEEKS + 1             # 9
WEEK_TOTAL = WEEK_LAST + 1          # 10
WEEKS = [datetime.date(2026, 9, 7) + datetime.timedelta(weeks=i)
         for i in range(N_WEEKS)]


def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as f:
        return list(csv.reader(f))


def dval(text):
    return datetime.date.fromisoformat(text)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def main():
    prod_rows = read_csv("products.csv")
    mov_rows = read_csv("movements.csv")
    shift_rows = read_csv("shifts.csv")
    order_rows = read_csv("orders.csv")
    assert len(prod_rows) - 1 == N_PRODUCTS
    assert len(mov_rows) - 1 == N_MOVEMENTS
    assert len(shift_rows) - 1 == N_SHIFTS
    assert len(order_rows) - 1 == N_ORDERS

    wb = Workbook()

    # ---- Guide ----
    ws = wb.active
    ws.title = "Guide"
    guide_lines = [
        ("OpsPulse — Weekly Operations Dashboard (Mehta Trading Co., Mumbai)", True),
        ("Demo workbook built around an operations-dashboard brief: weekly CSV", False),
        ("exports for product movements, shift schedules and order fulfilment,", False),
        ("turned into one refreshable dashboard for inventory, staff and orders.", False),
        ("", False),
        ("How to use: drop each week's new CSV rows into Movements / Raw_Shifts", False),
        ("(extend the formula ranges) / Orders, keeping the same columns. Every", False),
        ("derived figure — stock on hand, reorder flags, overtime, fulfilment", False),
        ("days, weekly summaries, dashboard KPIs and Checks — is a live formula.", False),
        ("Change the standards on the Assumptions sheet and the whole workbook", False),
        ("recomputes. No hard-coded totals anywhere.", False),
        ("", False),
        ("Standards: standard shift 8 hours; overtime = worked hours above 8;", False),
        ("on-time = shipped within 3 days of ordering; REORDER when stock on", False),
        ("hand is at or below the product's reorder level.", False),
        ("", False),
        ("All data are fictional and invented for this demo. No real client", False),
        ("data. Money is INR (Rs); period 7 Sep – 31 Oct 2026 (8 weeks).", False),
    ]
    for i, (text, bold) in enumerate(guide_lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11,
                      color=NAVY if i == 1 else "000000")
    widths(ws, {"A": 78})

    # ---- Dashboard (formulas reference sheets built below; strings only) ----
    ws = wb.create_sheet("Dashboard")
    ws.cell(row=1, column=1,
            value="OpsPulse — Operations Dashboard (as of 2026-10-31)").font = \
        Font(bold=True, size=14, color=NAVY)
    ws.cell(row=2, column=1,
            value="All values are live formulas over Inventory / Movements / "
                  "Raw_Shifts / Staff / Orders / Weekly / Assumptions.").font = \
        Font(italic=True, color="595959")
    kpis = [
        ("SKUs in range", f"=COUNTA(Inventory!$A$2:$A${PROD_LAST})", "#,##0"),
        ("SKUs needing reorder", f"=COUNTIF(Inventory!$M$2:$M${PROD_LAST},\"REORDER\")", "#,##0"),
        ("Total stock value (Rs)", f"=SUM(Inventory!$K$2:$K${PROD_LAST})", "#,##0"),
        ("Total stock on hand (units)", f"=SUM(Inventory!$I$2:$I${PROD_LAST})", "#,##0"),
        ("Orders (8 weeks)", f"=SUM(Weekly!$B$2:$B${WEEK_LAST})", "#,##0"),
        ("Order value (Rs)", f"=SUM(Weekly!$C$2:$C${WEEK_LAST})", "#,##0"),
        ("On-time fulfilment", "=IF($B$7=0,0,SUM(Weekly!$D$2:$D$9)/$B$7)", "0.0%"),
        ("Average fulfilment (days)", f"=AVERAGE(Orders!$H$2:$H${ORDER_LAST})", "0.00"),
        ("Overtime hours", f"=SUM(Weekly!$I$2:$I${WEEK_LAST})", "0.0"),
        ("Staff utilisation", f"=SUM(Staff!$E$2:$E${STAFF_LAST})/SUM(Staff!$D$2:$D${STAFF_LAST})", "0.0%"),
        ("Share of SKUs needing reorder", "=IF($B$3=0,0,$B$4/$B$3)", "0.0%"),
        ("Standard shift (hours)", "=Assumptions!$B$2", "0.0"),
        ("On-time target (days)", "=Assumptions!$B$3", "0"),
    ]
    for i, (label, formula, fmt) in enumerate(kpis, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format = fmt
        c.font = Font(bold=True, size=12)
    ws.cell(row=16, column=1,
            value="Charts below are wired to Inventory / Staff / Weekly — they move with the data.").font = \
        Font(italic=True, color="595959")
    widths(ws, {"A": 46, "B": 18, "C": 20, "D": 20, "E": 20, "F": 20})

    # ---- Inventory ----
    ws = wb.create_sheet("Inventory")
    cols = ["ProductID", "Product", "Category", "Unit", "OpeningStock",
            "Receipts", "Sales", "Adjustments", "StockOnHand", "UnitCost (Rs)",
            "StockValue (Rs)", "ReorderLevel", "Status"]
    ws.append(cols)
    style_header(ws, 1, len(cols))
    for i, row in enumerate(prod_rows[1:], start=2):
        ws.cell(row=i, column=1, value=row[0])
        ws.cell(row=i, column=2, value=row[1])
        ws.cell(row=i, column=3, value=row[2])
        ws.cell(row=i, column=4, value=row[3])
        ws.cell(row=i, column=5, value=int(row[4])).number_format = "#,##0"
        ws.cell(row=i, column=6,
                value=f"=SUMIFS(Movements!$E$2:$E${MOV_LAST},Movements!$B$2:$B${MOV_LAST},$A{i},Movements!$C$2:$C${MOV_LAST},\"Receipt\")").number_format = "#,##0"
        ws.cell(row=i, column=7,
                value=f"=SUMIFS(Movements!$E$2:$E${MOV_LAST},Movements!$B$2:$B${MOV_LAST},$A{i},Movements!$C$2:$C${MOV_LAST},\"Sale\")").number_format = "#,##0"
        ws.cell(row=i, column=8,
                value=f"=SUMIFS(Movements!$E$2:$E${MOV_LAST},Movements!$B$2:$B${MOV_LAST},$A{i},Movements!$C$2:$C${MOV_LAST},\"Adjustment\")").number_format = "#,##0"
        ws.cell(row=i, column=9, value=f"=$E{i}+$F{i}-$G{i}+$H{i}").number_format = "#,##0"
        ws.cell(row=i, column=10, value=float(row[6])).number_format = "#,##0"
        ws.cell(row=i, column=11, value=f"=$I{i}*$J{i}").number_format = "#,##0"
        ws.cell(row=i, column=12, value=int(row[5])).number_format = "#,##0"
        ws.cell(row=i, column=13, value=f"=IF($I{i}<=$L{i},\"REORDER\",\"OK\")")
    ws.cell(row=PROD_TOTAL, column=1, value="TOTAL").font = Font(bold=True)
    for col in ("E", "F", "G", "H", "I", "K"):
        c = ws.cell(row=PROD_TOTAL, column=ord(col) - 64,
                    value=f"=SUM(${col}$2:${col}${PROD_LAST})")
        c.font = Font(bold=True)
        c.number_format = "#,##0"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.auto_filter.ref = f"A1:M{PROD_LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 24, "C": 12, "D": 8, "E": 13, "F": 10, "G": 9,
                "H": 12, "I": 13, "J": 13, "K": 15, "L": 13, "M": 10})

    # ---- Movements (raw + NetQuantity formula) ----
    ws = wb.create_sheet("Movements")
    ws.append(["MovementDate", "ProductID", "MovementType", "WeekStart",
               "Quantity", "NetQuantity"])
    style_header(ws, 1, 6)
    for i, row in enumerate(mov_rows[1:], start=2):
        c = ws.cell(row=i, column=1, value=dval(row[0])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=row[1])
        ws.cell(row=i, column=3, value=row[2])
        c = ws.cell(row=i, column=4, value=dval(row[3])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5, value=int(row[4])).number_format = "#,##0"
        ws.cell(row=i, column=6, value=f"=IF($C{i}=\"Sale\",-$E{i},$E{i})").number_format = "#,##0"
    ws.auto_filter.ref = f"A1:F{MOV_LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 14, "B": 10, "C": 13, "D": 12, "E": 10, "F": 12})

    # ---- Raw_Shifts (raw + Overtime formula) ----
    ws = wb.create_sheet("Raw_Shifts")
    ws.append(["ShiftDate", "Employee", "Team", "WeekStart",
               "ScheduledHours", "WorkedHours", "OvertimeHours"])
    style_header(ws, 1, 7)
    for i, row in enumerate(shift_rows[1:], start=2):
        c = ws.cell(row=i, column=1, value=dval(row[0])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=row[1])
        ws.cell(row=i, column=3, value=row[2])
        c = ws.cell(row=i, column=4, value=dval(row[3])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5, value=float(row[4])).number_format = "0.0"
        ws.cell(row=i, column=6, value=float(row[5])).number_format = "0.0"
        ws.cell(row=i, column=7, value=f"=MAX(0,$F{i}-Assumptions!$B$2)").number_format = "0.0"
    ws.auto_filter.ref = f"A1:G{SHIFT_LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 12, "B": 14, "C": 11, "D": 12, "E": 15, "F": 12, "G": 14})

    # ---- Staff summary ----
    ws = wb.create_sheet("Staff")
    ws.append(["Employee", "Team", "Shifts", "ScheduledHours", "WorkedHours",
               "OvertimeHours", "Utilisation", "Status"])
    style_header(ws, 1, 8)
    employees = []
    for row in shift_rows[1:]:
        if (row[1], row[2]) not in employees:
            employees.append((row[1], row[2]))
    assert len(employees) == N_STAFF
    for i, (name, team) in enumerate(employees, start=2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=team)
        ws.cell(row=i, column=3,
                value=f"=COUNTIF(Raw_Shifts!$B$2:$B${SHIFT_LAST},$A{i})").number_format = "#,##0"
        ws.cell(row=i, column=4,
                value=f"=SUMIF(Raw_Shifts!$B$2:$B${SHIFT_LAST},$A{i},Raw_Shifts!$E$2:$E${SHIFT_LAST})").number_format = "0.0"
        ws.cell(row=i, column=5,
                value=f"=SUMIF(Raw_Shifts!$B$2:$B${SHIFT_LAST},$A{i},Raw_Shifts!$F$2:$F${SHIFT_LAST})").number_format = "0.0"
        ws.cell(row=i, column=6,
                value=f"=SUMIF(Raw_Shifts!$B$2:$B${SHIFT_LAST},$A{i},Raw_Shifts!$G$2:$G${SHIFT_LAST})").number_format = "0.0"
        ws.cell(row=i, column=7, value=f"=IF($D{i}=0,0,$E{i}/$D{i})").number_format = "0.0%"
        ws.cell(row=i, column=8, value=f"=IF($F{i}>10,\"Review\",\"OK\")")
    ws.cell(row=STAFF_TOTAL, column=1, value="TOTAL").font = Font(bold=True)
    for col in ("C", "D", "E", "F"):
        c = ws.cell(row=STAFF_TOTAL, column=ord(col) - 64,
                    value=f"=SUM(${col}$2:${col}${STAFF_LAST})")
        c.font = Font(bold=True)
        c.number_format = "0.0" if col != "C" else "#,##0"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.auto_filter.ref = f"A1:H{STAFF_LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 14, "B": 11, "C": 8, "D": 15, "E": 12, "F": 14, "G": 12, "H": 9})

    # ---- Orders (raw + value / fulfilment / on-time formulas) ----
    ws = wb.create_sheet("Orders")
    ws.append(["OrderID", "OrderDate", "ProductID", "Quantity", "UnitPrice (Rs)",
               "OrderValue (Rs)", "ShipDate", "FulfilmentDays", "WeekStart", "OnTime"])
    style_header(ws, 1, 10)
    for i, row in enumerate(order_rows[1:], start=2):
        ws.cell(row=i, column=1, value=row[0])
        c = ws.cell(row=i, column=2, value=dval(row[1])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3, value=row[2])
        ws.cell(row=i, column=4, value=int(row[3])).number_format = "#,##0"
        ws.cell(row=i, column=5, value=float(row[4])).number_format = "#,##0"
        ws.cell(row=i, column=6, value=f"=$D{i}*$E{i}").number_format = "#,##0"
        c = ws.cell(row=i, column=7, value=dval(row[5])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=8, value=f"=$G{i}-$B{i}").number_format = "0"
        c = ws.cell(row=i, column=9, value=dval(row[6])); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=10, value=f"=IF($H{i}<=Assumptions!$B$3,\"Yes\",\"No\")")
    ws.cell(row=ORDER_TOTAL, column=1, value="TOTAL").font = Font(bold=True)
    for col in ("D", "F"):
        c = ws.cell(row=ORDER_TOTAL, column=ord(col) - 64,
                    value=f"=SUM(${col}$2:${col}${ORDER_LAST})")
        c.font = Font(bold=True)
        c.number_format = "#,##0"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.auto_filter.ref = f"A1:J{ORDER_LAST}"
    ws.freeze_panes = "A2"
    widths(ws, {"A": 9, "B": 12, "C": 10, "D": 9, "E": 14, "F": 15, "G": 11,
                "H": 15, "I": 12, "J": 8})

    # ---- Weekly summary ----
    ws = wb.create_sheet("Weekly")
    ws.append(["WeekStart", "Orders", "OrderValue (Rs)", "OnTimeOrders",
               "OnTimePct", "AvgFulfilmentDays", "ReceiptQty", "SaleQty",
               "OvertimeHours"])
    style_header(ws, 1, 9)
    for i, wk in enumerate(WEEKS, start=2):
        c = ws.cell(row=i, column=1, value=wk); c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2,
                value=f"=COUNTIF(Orders!$I$2:$I${ORDER_LAST},$A{i})").number_format = "#,##0"
        ws.cell(row=i, column=3,
                value=f"=SUMIF(Orders!$I$2:$I${ORDER_LAST},$A{i},Orders!$F$2:$F${ORDER_LAST})").number_format = "#,##0"
        ws.cell(row=i, column=4,
                value=f"=COUNTIFS(Orders!$I$2:$I${ORDER_LAST},$A{i},Orders!$J$2:$J${ORDER_LAST},\"Yes\")").number_format = "#,##0"
        ws.cell(row=i, column=5, value=f"=IF($B{i}=0,0,$D{i}/$B{i})").number_format = "0.0%"
        ws.cell(row=i, column=6,
                value=f"=AVERAGEIF(Orders!$I$2:$I${ORDER_LAST},$A{i},Orders!$H$2:$H${ORDER_LAST})").number_format = "0.00"
        ws.cell(row=i, column=7,
                value=f"=SUMIFS(Movements!$E$2:$E${MOV_LAST},Movements!$D$2:$D${MOV_LAST},$A{i},Movements!$C$2:$C${MOV_LAST},\"Receipt\")").number_format = "#,##0"
        ws.cell(row=i, column=8,
                value=f"=SUMIFS(Movements!$E$2:$E${MOV_LAST},Movements!$D$2:$D${MOV_LAST},$A{i},Movements!$C$2:$C${MOV_LAST},\"Sale\")").number_format = "#,##0"
        ws.cell(row=i, column=9,
                value=f"=SUMIF(Raw_Shifts!$D$2:$D${SHIFT_LAST},$A{i},Raw_Shifts!$G$2:$G${SHIFT_LAST})").number_format = "0.0"
    ws.cell(row=WEEK_TOTAL, column=1, value="TOTAL").font = Font(bold=True)
    for col in ("B", "C", "D", "G", "H", "I"):
        c = ws.cell(row=WEEK_TOTAL, column=ord(col) - 64,
                    value=f"=SUM(${col}$2:${col}${WEEK_LAST})")
        c.font = Font(bold=True)
        c.number_format = "0.0" if col == "I" else "#,##0"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.cell(row=WEEK_TOTAL, column=5,
            value=f"=IF($B${WEEK_TOTAL}=0,0,$D${WEEK_TOTAL}/$B${WEEK_TOTAL})").number_format = "0.0%"
    ws.cell(row=WEEK_TOTAL, column=5).font = Font(bold=True)
    ws.cell(row=WEEK_TOTAL, column=6,
            value=f"=AVERAGE(Orders!$H$2:$H${ORDER_LAST})").number_format = "0.00"
    ws.cell(row=WEEK_TOTAL, column=6).font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths(ws, {"A": 12, "B": 8, "C": 15, "D": 13, "E": 11, "F": 17, "G": 12,
                "H": 9, "I": 14})

    # ---- Assumptions ----
    ws = wb.create_sheet("Assumptions")
    ws.cell(row=1, column=1, value="Assumptions").font = Font(bold=True, size=13, color=NAVY)
    ws.cell(row=2, column=1, value="Standard shift (hours)").font = Font(bold=True)
    ws.cell(row=2, column=2, value=8).number_format = "0.0"
    ws.cell(row=3, column=1, value="On-time target (days)").font = Font(bold=True)
    ws.cell(row=3, column=2, value=3).number_format = "0"
    ws.cell(row=4, column=1, value="As-of date").font = Font(bold=True)
    c = ws.cell(row=4, column=2, value=datetime.date(2026, 10, 31))
    c.number_format = "yyyy-mm-dd"
    ws.cell(row=5, column=1, value="Reporting period").font = Font(bold=True)
    ws.cell(row=5, column=2, value="2026-09-07 to 2026-10-31 (8 weeks)")
    ws.cell(row=6, column=1, value="Weeks in period").font = Font(bold=True)
    ws.cell(row=6, column=2, value=N_WEEKS)
    ws.cell(row=7, column=1, value="Currency").font = Font(bold=True)
    ws.cell(row=7, column=2, value="INR (Rs)")
    widths(ws, {"A": 30, "B": 38})

    # ---- Checks ----
    ws = wb.create_sheet("Checks")
    ws.append(["Check", "Expected", "Got", "Diff", "Status"])
    checks = [
        ("SKU count", N_PRODUCTS, f"=COUNTA(Inventory!$A$2:$A${PROD_LAST})", "0"),
        ("Dashboard stock value ties to Inventory",
         f"=SUM(Inventory!$K$2:$K${PROD_LAST})", "=Dashboard!$B$5", "0.01"),
        ("Dashboard reorder count ties to Inventory",
         f"=COUNTIF(Inventory!$M$2:$M${PROD_LAST},\"REORDER\")", "=Dashboard!$B$4", "0"),
        ("Dashboard order count ties to Orders",
         f"=COUNTA(Orders!$A$2:$A${ORDER_LAST})", "=Dashboard!$B$7", "0"),
        ("Dashboard order value ties to Weekly",
         f"=SUM(Weekly!$C$2:$C${WEEK_LAST})", "=Dashboard!$B$8", "0.01"),
        ("Weekly on-time total ties to Orders",
         f"=SUM(Weekly!$D$2:$D${WEEK_LAST})",
         f"=COUNTIF(Orders!$J$2:$J${ORDER_LAST},\"Yes\")", "0"),
        ("Net movements tie to stock change",
         f"=SUM(Inventory!$I$2:$I${PROD_LAST})-SUM(Inventory!$E$2:$E${PROD_LAST})",
         f"=SUM(Movements!$F$2:$F${MOV_LAST})", "0.01"),
        ("Weekly overtime ties to Staff",
         f"=SUM(Staff!$F$2:$F${STAFF_LAST})", f"=SUM(Weekly!$I$2:$I${WEEK_LAST})", "0.01"),
        ("Products with negative stock", 0,
         f"=COUNTIF(Inventory!$I$2:$I${PROD_LAST},\"<0\")", "0"),
        ("Orders with blank ship date", 0,
         f"=COUNTBLANK(Orders!$G$2:$G${ORDER_LAST})", "0"),
    ]
    for i, (label, expected, got, tol) in enumerate(checks, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=expected)
        ws.cell(row=i, column=3, value=got)
        ws.cell(row=i, column=4, value=f"=$C{i}-$B{i}")
        ws.cell(row=i, column=5,
                value=f"=IF(ABS($D{i})<={tol},\"PASS\",\"FAIL\")")
    style_header(ws, 1, 5)
    ws.freeze_panes = "A2"
    for r in range(2, 2 + len(checks)):
        for col in ("B", "C", "D"):
            ws[f"{col}{r}"].number_format = "#,##0.00"
    widths(ws, {"A": 52, "B": 16, "C": 16, "D": 12, "E": 10})

    # ---- Charts on Dashboard ----
    ws = wb["Dashboard"]
    col_chart = BarChart()
    col_chart.type = "col"
    col_chart.title = "Stock Value by Product (Rs)"
    col_chart.y_axis.title = "Rs"
    col_chart.style = 10
    data = Reference(wb["Inventory"], min_col=11, min_row=1, max_row=PROD_LAST)
    cats = Reference(wb["Inventory"], min_col=1, min_row=2, max_row=PROD_LAST)
    col_chart.add_data(data, titles_from_data=True)
    col_chart.set_categories(cats)
    col_chart.height = 7.5
    col_chart.width = 15
    ws.add_chart(col_chart, "A17")

    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.title = "Overtime Hours by Employee"
    bar_chart.x_axis.title = "Hours"
    bar_chart.style = 10
    data = Reference(wb["Staff"], min_col=6, min_row=1, max_row=STAFF_LAST)
    cats = Reference(wb["Staff"], min_col=1, min_row=2, max_row=STAFF_LAST)
    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    bar_chart.height = 7.5
    bar_chart.width = 15
    ws.add_chart(bar_chart, "A32")

    line = LineChart()
    line.title = "Order Value by Week (Rs)"
    line.y_axis.title = "Rs"
    line.style = 10
    ldata = Reference(wb["Weekly"], min_col=3, min_row=1, max_row=WEEK_LAST)
    lcats = Reference(wb["Weekly"], min_col=1, min_row=2, max_row=WEEK_LAST)
    line.add_data(ldata, titles_from_data=True)
    line.set_categories(lcats)
    line.height = 7.5
    line.width = 15
    ws.add_chart(line, "A47")

    wb.save(OUT)
    print(f"wrote {OUT} sheets={wb.sheetnames}")


if __name__ == "__main__":
    main()
