"""Tests for the Excel Finance Tracker generator."""
import csv
import os
import subprocess
import unittest
from datetime import date, timedelta

import openpyxl

XLSX = "Excel_Finance_Tracker.xlsx"
RECALC_DIR = "/tmp/ft-20261009/recalc"
AS_OF = date(2026, 10, 9)

EXPECTED_SHEETS = ["Dashboard", "Transactions", "Budget vs Actual",
                   "Invoices", "Tax Prep", "ChartData"]
EXPENSE_CATS = {"Rent", "Salaries", "Marketing", "Travel", "Utilities",
                "Software", "Office Supplies", "Insurance", "Professional Fees"}


def read_csv(name):
    with open(name, newline="") as f:
        return list(csv.DictReader(f))


def recalc_values():
    """Roundtrip through LibreOffice so formulas get computed values."""
    os.makedirs(RECALC_DIR, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                    "--outdir", RECALC_DIR, XLSX],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(os.path.join(RECALC_DIR, XLSX), data_only=True)


class TestSampleData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tx = read_csv("sample_transactions.csv")
        cls.inv = read_csv("sample_invoices.csv")
        cls.bud = read_csv("sample_budgets.csv")

    def test_transaction_columns(self):
        self.assertEqual(list(self.tx[0].keys()),
                         ["Date", "Type", "Category", "Description", "Amount"])

    def test_transaction_count(self):
        # 9 months x (2 income + 7 fixed expenses) + 6 quarterly extras
        self.assertEqual(len(self.tx), 9 * 9 + 6)

    def test_invoice_columns(self):
        self.assertEqual(list(self.inv[0].keys()),
                         ["InvoiceNo", "Client", "IssueDate", "DueDate",
                          "Amount", "Status"])

    def test_invoice_count_and_statuses(self):
        self.assertEqual(len(self.inv), 20)
        statuses = {r["Status"] for r in self.inv}
        self.assertEqual(statuses, {"Paid", "Unpaid"})

    def test_due_dates_are_issue_plus_30(self):
        for r in self.inv:
            issued = date.fromisoformat(r["IssueDate"])
            due = date.fromisoformat(r["DueDate"])
            self.assertEqual(due - issued, timedelta(days=30))

    def test_budget_covers_all_expense_categories(self):
        self.assertEqual({r["Category"] for r in self.bud}, EXPENSE_CATS)

    def test_income_is_positive(self):
        inc = sum(float(r["Amount"]) for r in self.tx if r["Type"] == "Income")
        self.assertGreater(inc, 0)

    def test_three_overdue_and_three_outstanding(self):
        unpaid = [r for r in self.inv if r["Status"] == "Unpaid"]
        overdue = [r for r in unpaid
                   if date.fromisoformat(r["DueDate"]) < AS_OF]
        outst = [r for r in unpaid
                 if date.fromisoformat(r["DueDate"]) >= AS_OF]
        self.assertEqual(len(overdue), 3)
        self.assertEqual(len(outst), 3)

    def test_deterministic_regeneration(self):
        import make_sample_data
        before_tx = read_csv("sample_transactions.csv")
        before_inv = read_csv("sample_invoices.csv")
        make_sample_data.main()
        self.assertEqual(read_csv("sample_transactions.csv"), before_tx)
        self.assertEqual(read_csv("sample_invoices.csv"), before_inv)


class TestWorkbookStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(XLSX)

    def test_all_sheets_present_in_order(self):
        self.assertEqual(self.wb.sheetnames, EXPECTED_SHEETS)

    def test_transactions_row_count(self):
        ws = self.wb["Transactions"]
        self.assertEqual(ws.max_row, 1 + 87)  # header + 87 data rows

    def test_month_helper_formula(self):
        ws = self.wb["Transactions"]
        self.assertEqual(ws["F2"].value, '=TEXT(A2,"mmm-yy")')

    def test_taxcat_vlookup_formula(self):
        ws = self.wb["Transactions"]
        f = ws["G2"].value
        self.assertIn("VLOOKUP", f)
        self.assertIn("'Tax Prep'!$A$21:$B$29", f)
        self.assertIn('IF(B2="Income",""', f)

    def test_kpi_formulas(self):
        ws = self.wb["Dashboard"]
        self.assertEqual(ws["A3"].value, "Total Income")
        self.assertIn('SUMIF(Transactions!B:B,"Income"', ws["B3"].value)
        self.assertEqual(ws["A5"].value, "Net Profit")
        self.assertEqual(ws["B5"].value, "=B3-B4")
        self.assertEqual(ws["A7"].value, "Outstanding Receivables")
        self.assertIn("SUMIFS(Invoices!E:E", ws["B7"].value)
        self.assertIn('"Outstanding"', ws["B7"].value)
        self.assertEqual(ws["A8"].value, "Overdue Amount")
        self.assertIn('"Overdue"', ws["B8"].value)

    def test_budget_actual_formulas(self):
        ws = self.wb["Budget vs Actual"]
        self.assertEqual(ws["C2"].value, "=B2*9")
        f = ws["D2"].value
        self.assertIn("SUMIFS(Transactions!$E:$E", f)
        self.assertIn('Transactions!$B:$B,"Expense"', f)
        self.assertEqual(ws["E2"].value, "=C2-D2")
        self.assertEqual(ws["B11"].value, "=SUM(B2:B10)")
        self.assertIn("TOTAL", ws["A11"].value)

    def test_conditional_formatting_on_variance(self):
        ws = self.wb["Budget vs Actual"]
        rules = list(ws.conditional_formatting._cf_rules.keys())
        self.assertTrue(any("E2:E10" in str(r) for r in rules),
                        f"no variance CF rule found: {rules}")

    def test_invoice_aging_formulas(self):
        ws = self.wb["Invoices"]
        self.assertEqual(ws["B1"].value.date(), AS_OF)
        self.assertEqual(ws["G4"].value, '=IF(F4="Paid","",MAX(0,$B$1-D4))')
        f = ws["H4"].value
        self.assertIn('IF(F4="Paid","Paid"', f)
        self.assertIn('"Overdue"', f)
        self.assertIn('"Outstanding"', f)

    def test_tax_prep_formulas(self):
        ws = self.wb["Tax Prep"]
        f = ws["B3"].value
        self.assertIn("SUMIFS(Transactions!$E:$E", f)
        self.assertIn("Transactions!$G:$G", f)
        self.assertEqual(ws["A21"].value, "Insurance")
        self.assertEqual(ws["B21"].value, "Insurance")

    def test_chartdata_formulas(self):
        ws = self.wb["ChartData"]
        self.assertEqual(ws["A2"].value, "Jan-26")
        self.assertIn("SUMIFS(Transactions!$E:$E", ws["B2"].value)
        self.assertEqual(ws["D2"].value, "=B2-C2")
        self.assertEqual(ws["D3"].value, "=B3-C3+D2")

    def test_three_charts_on_dashboard(self):
        ws = self.wb["Dashboard"]
        self.assertEqual(len(ws._charts), 3)
        kinds = sorted(type(c).__name__ for c in ws._charts)
        self.assertEqual(kinds, ["BarChart", "LineChart", "PieChart"])

    def test_bar_chart_has_two_series(self):
        ws = self.wb["Dashboard"]
        bar = next(c for c in ws._charts if type(c).__name__ == "BarChart")
        self.assertEqual(len(bar.series), 2)

    def test_no_broken_formulas_anywhere(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        self.assertNotIn("#REF!", v, f"{ws.title}!{cell.coordinate}")
                        self.assertNotIn("#VALUE!", v, f"{ws.title}!{cell.coordinate}")


class TestRecalculatedValues(unittest.TestCase):
    """LibreOffice recalc roundtrip: formulas must compute the right numbers."""

    @classmethod
    def setUpClass(cls):
        cls.tx = read_csv("sample_transactions.csv")
        cls.inv = read_csv("sample_invoices.csv")
        cls.bud = {r["Category"]: float(r["MonthlyBudget"])
                   for r in read_csv("sample_budgets.csv")}
        cls.wb = recalc_values()

    def total(self, typ):
        return sum(float(r["Amount"]) for r in self.tx if r["Type"] == typ)

    def test_total_income_value(self):
        got = self.wb["Dashboard"]["B3"].value
        self.assertAlmostEqual(got, self.total("Income"), delta=1.0)

    def test_total_expenses_value(self):
        got = self.wb["Dashboard"]["B4"].value
        self.assertAlmostEqual(got, self.total("Expense"), delta=1.0)

    def test_net_profit_value(self):
        got = self.wb["Dashboard"]["B5"].value
        self.assertAlmostEqual(got, self.total("Income") - self.total("Expense"),
                               delta=1.0)

    def test_overdue_amount_value(self):
        expected = sum(float(r["Amount"]) for r in self.inv
                       if r["Status"] == "Unpaid"
                       and date.fromisoformat(r["DueDate"]) < AS_OF)
        got = self.wb["Dashboard"]["B8"].value
        self.assertAlmostEqual(got, expected, delta=1.0)

    def test_outstanding_receivables_value(self):
        expected = sum(float(r["Amount"]) for r in self.inv
                       if r["Status"] == "Unpaid"
                       and date.fromisoformat(r["DueDate"]) >= AS_OF)
        got = self.wb["Dashboard"]["B7"].value
        self.assertAlmostEqual(got, expected, delta=1.0)

    def test_marketing_variance_value(self):
        actual = sum(float(r["Amount"]) for r in self.tx
                     if r["Type"] == "Expense" and r["Category"] == "Marketing")
        expected = self.bud["Marketing"] * 9 - actual
        got = self.wb["Budget vs Actual"]["E4"].value  # Marketing row
        self.assertAlmostEqual(got, expected, delta=1.0)

    def test_invoice_flag_values(self):
        ws = self.wb["Invoices"]
        flags = [ws[f"H{i}"].value for i in range(4, 24)]
        self.assertEqual(flags.count("Paid"), 14)
        self.assertEqual(flags.count("Overdue"), 3)
        self.assertEqual(flags.count("Outstanding"), 3)

    def test_days_outstanding_spot_check(self):
        ws = self.wb["Invoices"]
        # INV-26014: due 2026-09-15 -> 24 days outstanding as of 2026-10-09
        row = next(i for i in range(4, 24) if ws[f"A{i}"].value == "INV-26014")
        self.assertEqual(ws[f"G{row}"].value, 24)

    def test_cumulative_net_final_value(self):
        got = self.wb["ChartData"]["D10"].value
        self.assertAlmostEqual(got, self.total("Income") - self.total("Expense"),
                               delta=1.0)

    def test_tax_prep_office_total(self):
        office_cats = {"Rent", "Utilities", "Office Supplies"}
        expected = sum(float(r["Amount"]) for r in self.tx
                       if r["Type"] == "Expense" and r["Category"] in office_cats)
        ws = self.wb["Tax Prep"]
        row = next(i for i in range(3, 10) if ws[f"A{i}"].value == "Office")
        self.assertAlmostEqual(ws[f"B{row}"].value, expected, delta=1.0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
