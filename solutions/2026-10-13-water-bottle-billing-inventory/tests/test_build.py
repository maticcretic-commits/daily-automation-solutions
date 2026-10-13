#!/usr/bin/env python3
"""Tests for the AquaBill water-bottle billing/inventory workbook.

Covers: determinism, data invariants, formula hygiene (the $H$2:$5000 class),
and LibreOffice headless recalc vs independent Python ground truth.
"""
import csv
import re
import subprocess
import sys
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from generate import SEED  # noqa: E402  (reuse seed constant)

DATADIR = ROOT / "data"
XLSX = ROOT / "water-bottle-billing-inventory.xlsx"
ERROR_TOKENS = ("#DIV/0!", "#N/A", "#NAME?", "#REF!", "#VALUE!", "#NUM!",
                "Err:502", "Err:503", "Err:508", "Err:509", "Err:510", "Err:522")


def r2(x):
    return Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def read_csv(name):
    with open(DATADIR / name, newline="") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def data():
    return {
        "skus": read_csv("skus.csv"),
        "invoices": read_csv("invoices.csv"),
        "lines": read_csv("invoice_lines.csv"),
        "purchases": read_csv("purchases.csv"),
    }


@pytest.fixture(scope="module")
def wb_formulas():
    return load_workbook(XLSX, data_only=False)


@pytest.fixture(scope="module")
def wb_values(tmp_path_factory):
    outdir = tmp_path_factory.mktemp("recalc")
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         "--outdir", str(outdir), str(XLSX)],
        capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, f"soffice failed: {r.stderr[-2000:]}"
    converted = outdir / XLSX.name
    assert converted.exists(), "LO produced no output file"
    return load_workbook(converted, data_only=True)


def ground_truth(data):
    inv = data["invoices"]
    lines = data["lines"]
    purch = data["purchases"]
    skus = data["skus"]
    subtotal, gst, grand, paid = (Decimal("0"),) * 4
    for i in inv:
        subtotal += Decimal(i["subtotal"])
        gst += Decimal(i["gst"])
        grand += Decimal(i["grand"])
        paid += Decimal(i["paid"])
    units = sum(int(l["qty"]) for l in lines)
    bought = {s["sku"]: 0 for s in skus}
    for p in purch:
        bought[p["sku"]] += int(p["qty"])
    sold = {s["sku"]: 0 for s in skus}
    for l in lines:
        sold[l["sku"]] += int(l["qty"])
    closing, status, stockval = {}, {}, Decimal("0")
    for s in skus:
        c = int(s["opening_stock"]) + bought[s["sku"]] - sold[s["sku"]]
        closing[s["sku"]] = c
        status[s["sku"]] = ("OUT OF STOCK" if c <= 0
                            else "REORDER" if c <= int(s["reorder_level"])
                            else "OK")
        stockval += Decimal(c) * Decimal(s["rate"])
    months = {}
    for i in inv:
        m = i["date"][:7]
        months.setdefault(m, [Decimal("0"), Decimal("0")])
        months[m][0] += Decimal(i["subtotal"])
        months[m][1] += Decimal(i["gst"])
    return {
        "revenue": grand, "n_invoices": len(inv), "units": units,
        "avg": grand / len(inv), "gst": gst, "outstanding": grand - paid,
        "reorder_n": sum(1 for v in status.values() if v != "OK"),
        "stockval": stockval, "closing": closing, "status": status,
        "months": months,
    }


# ---------- generator-level tests ----------

def test_determinism(tmp_path):
    import generate as gen_mod
    import random
    outs = []
    for i in range(2):
        d = tmp_path / f"run{i}"
        d.mkdir()
        rng = random.Random(SEED)
        sku_map, invoices, lines, purchases = gen_mod.gen(rng)
        gen_mod.write_csv(str(d / "invoices.csv"), invoices,
                          ["inv_no", "date", "customer", "subtotal", "gst",
                           "grand", "paid"])
        gen_mod.write_csv(str(d / "invoice_lines.csv"), lines,
                          ["inv_no", "date", "line_no", "sku", "qty", "rate"])
        outs.append((d / "invoices.csv").read_bytes() +
                    (d / "invoice_lines.csv").read_bytes())
    assert outs[0] == outs[1]


def test_reasonable_volume(data):
    assert 120 <= len(data["invoices"]) <= 320
    assert len(data["lines"]) >= len(data["invoices"])
    assert len(data["purchases"]) >= 10


def test_invoice_numbers_unique_and_sequential(data):
    nos = [i["inv_no"] for i in data["invoices"]]
    assert len(set(nos)) == len(nos)
    seq = sorted(int(n.split("-")[-1]) for n in nos)
    assert seq == list(range(1, len(nos) + 1))


def test_invoice_math_from_lines(data):
    by_inv = {}
    for l in data["lines"]:
        by_inv.setdefault(l["inv_no"], []).append(l)
    for inv in data["invoices"]:
        sub = r2(sum(Decimal(l["qty"]) * Decimal(l["rate"])
                     for l in by_inv[inv["inv_no"]]))
        g = r2(sub * Decimal("0.18"))
        assert Decimal(inv["subtotal"]) == sub, inv["inv_no"]
        assert Decimal(inv["gst"]) == g, inv["inv_no"]
        assert Decimal(inv["grand"]) == sub + g, inv["inv_no"]
        assert Decimal("0") <= Decimal(inv["paid"]) <= Decimal(inv["grand"])


def test_no_negative_closing_stock(data):
    gt = ground_truth(data)
    assert all(c >= 0 for c in gt["closing"].values())


# ---------- workbook structure / formula hygiene ----------

def test_sheet_names(wb_formulas):
    assert wb_formulas.sheetnames == [
        "Dashboard", "InvoiceLines", "Invoices", "Inventory", "Purchases",
        "GST_Summary", "Daily", "Checks", "Guide"]


def test_no_invalid_range_literals(wb_formulas):
    bad = []
    for ws in wb_formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    if re.search(r"\$[A-Z]{1,3}\$\d+:\$\d+", v):
                        bad.append((ws.title, cell.coordinate, v[:80]))
                    if "[@" in v:
                        bad.append((ws.title, cell.coordinate, "structured ref"))
    assert not bad, f"invalid range literals: {bad[:5]}"


def test_all_formulas_live_not_hardcoded(wb_formulas):
    # Spot-check: computed columns must be formulas, not pasted values.
    ws = wb_formulas["Invoices"]
    for r in (2, 3, 50):
        for col in "DEFHI":
            v = ws[f"{col}{r}"].value
            assert isinstance(v, str) and v.startswith("="), (col, r, v)
    ws = wb_formulas["Inventory"]
    for r in (2, 3):
        for col in "DEFHI":
            v = ws[f"{col}{r}"].value
            assert isinstance(v, str) and v.startswith("="), (col, r, v)


# ---------- LibreOffice recalc vs ground truth ----------

def test_lo_no_formula_errors(wb_values):
    errs = []
    for ws in wb_values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(t in v for t in ERROR_TOKENS):
                    errs.append((ws.title, cell.coordinate, v))
    assert not errs, f"formula errors: {errs[:8]}"


def test_lo_dashboard_kpis(data, wb_values):
    gt = ground_truth(data)
    ws = wb_values["Dashboard"]
    got = {r: ws[f"B{r}"].value for r in range(4, 12)}
    exp = [gt["revenue"], gt["n_invoices"], gt["units"], gt["avg"],
           gt["gst"], gt["outstanding"], gt["reorder_n"], gt["stockval"]]
    for r, e in zip(range(4, 12), exp):
        assert abs(Decimal(str(got[r])) - e) < Decimal("0.05"), (r, got[r], e)


def test_lo_gst_summary(data, wb_values):
    gt = ground_truth(data)
    ws = wb_values["GST_Summary"]
    for i, m in enumerate(["2026-04", "2026-05", "2026-06",
                           "2026-07", "2026-08", "2026-09"]):
        r = i + 2
        tax, g = gt["months"][m]
        cgst = r2(tax * Decimal("0.09"))
        assert abs(Decimal(str(ws[f"B{r}"].value)) - tax) < Decimal("0.05"), m
        assert abs(Decimal(str(ws[f"C{r}"].value)) - cgst) < Decimal("0.05"), m
        assert abs(Decimal(str(ws[f"D{r}"].value)) - cgst) < Decimal("0.05"), m
        # E = C + D by construction; the pair must approximate invoice-level GST
        assert abs(Decimal(str(ws[f"E{r}"].value))
                   - Decimal(str(ws[f"C{r}"].value))
                   - Decimal(str(ws[f"D{r}"].value))) < Decimal("0.005"), m
        assert abs(Decimal(str(ws[f"E{r}"].value)) - (cgst * 2)) \
            < Decimal("0.03"), m
    # whole-period GST tie, loose (per-invoice vs per-month rounding)
    tot_e = sum(Decimal(str(wb_values["GST_Summary"][f"E{r}"].value))
                for r in range(2, 8))
    tot_g = sum(v[1] for v in gt["months"].values())
    assert abs(tot_e - tot_g) < Decimal("1.0")


def test_lo_inventory_closing_and_status(data, wb_values):
    gt = ground_truth(data)
    ws = wb_values["Inventory"]
    for i, s in enumerate(data["skus"]):
        r = i + 2
        assert ws[f"A{r}"].value == s["sku"]
        assert ws[f"F{r}"].value == gt["closing"][s["sku"]], s["sku"]
        assert ws[f"H{r}"].value == gt["status"][s["sku"]], s["sku"]


def test_lo_checks_all_true(wb_values):
    ws = wb_values["Checks"]
    results = [ws[f"B{r}"].value for r in range(2, 10)]
    assert results == [True] * 8, results


def test_lo_daily_revenue_ties(data, wb_values):
    ws = wb_values["Daily"]
    total = sum(c.value for row in ws.iter_rows(min_row=2, max_col=3, max_row=401)
                for c in [row[2]] if isinstance(c.value, (int, float)))
    # Daily revenue is ex-GST: sum of per-line ROUND(qty*rate, 2)
    exp = sum(r2(Decimal(l["qty"]) * Decimal(l["rate"]))
              for l in data["lines"])
    assert abs(Decimal(str(total)) - exp) < Decimal("1.0")


def test_readme_documents_run():
    text = (ROOT / "README.md").read_text()
    for token in ["generate.py", "build_workbook.py", "pytest",
                  "water-bottle-billing-inventory.xlsx", "Checks"]:
        assert token in text, token
