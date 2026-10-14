#!/usr/bin/env python3
"""Tests for StockPulse. Run from the solution dir: python3 -m unittest discover -s tests"""
import csv
import os
import re
import shutil
import subprocess
import tempfile
import unittest
import zipfile

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "stock-portfolio-dashboard.xlsx")
N, W = 30, 33


def read_csv(name):
    with open(os.path.join(BASE, "data", name), newline="") as f:
        return list(csv.DictReader(f))


def ground_truth():
    holdings = {r["ticker"]: (int(r["qty"]), float(r["avg_cost"])) for r in read_csv("holdings.csv")}
    price = {r["ticker"]: [float(r[f"w{i}"]) for i in range(W)] for r in read_csv("prices.csv")}
    inv = sum(q * c for q, c in holdings.values())
    pvals = [sum(holdings[t][0] * price[t][i] for t in holdings) for i in range(W)]
    mv = pvals[-1]
    return holdings, price, inv, mv, pvals


def recalc(path):
    """LibreOffice headless recalc -> openpyxl data_only values."""
    tmp = tempfile.mkdtemp()
    # soffice keeps the INPUT basename for the output file:
    out = os.path.join(tmp, os.path.basename(path))
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, path],
        capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, f"soffice failed: {r.stderr[:400]}"
    wb = openpyxl.load_workbook(out, data_only=True)
    vals = {}
    err_re = re.compile(r"^#(DIV/0!|N/A|NAME\?|NULL!|NUM!|REF!|VALUE!|CALC!|SPILL!|GETTING_DATA|FIELD!)$")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and err_re.match(v):
                    vals[(ws.title, c.coordinate)] = ("ERR", v)
                elif isinstance(v, (int, float, str, bool)):
                    vals[(ws.title, c.coordinate)] = v
    return vals


class TestData(unittest.TestCase):
    def test_csvs_deterministic(self):
        a = {t: v for t, v in ground_truth()[1].items()}
        b = {t: v for t, v in ground_truth()[1].items()}
        self.assertEqual(a, b)

    def test_30_tickers_33_weeks(self):
        _h, price, _i, _m, _p = ground_truth()
        self.assertEqual(len(price), N)
        for t, s in price.items():
            self.assertEqual(len(s), W, t)
            self.assertTrue(all(x > 0 for x in s), t)

    def test_holdings_shape(self):
        h = {r["ticker"]: r for r in read_csv("holdings.csv")}
        self.assertEqual(len(h), N)
        for t, r in h.items():
            self.assertGreater(int(r["qty"]), 0, t)
            self.assertGreater(float(r["avg_cost"]), 0, t)


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(XLSX)

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames,
                         ["Guide", "Holdings", "Prices", "Weekly", "Sector",
                          "Dashboard", "AI_Report", "Checks"])

    def test_xlsx_is_valid_zip(self):
        with zipfile.ZipFile(XLSX) as z:
            self.assertIn("xl/workbook.xml", z.namelist())

    def test_holdings_rows(self):
        ws = self.wb["Holdings"]
        self.assertEqual(ws.max_row, N + 2)  # header + 30 + total
        self.assertEqual(ws["A32"].value, "TOTAL")

    def test_no_short_range_literals(self):
        """AGENTS.md regression: range literals must repeat the column letter
        on both ends (e.g. $H$2:$H$5000, never $H$2:$5000)."""
        bad = re.compile(r"\$[A-Z]+\$\d+:\$\d+")
        hits = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.startswith("=") and bad.search(v):
                        hits.append((ws.title, c.coordinate, v[:60]))
        self.assertEqual(hits, [], f"short range literals: {hits[:5]}")

    def test_no_structured_refs(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.startswith("="):
                        self.assertNotIn("[@", v, f"{ws.title}!{c.coordinate}")
                        self.assertNotIn("tbl", v.lower(), f"{ws.title}!{c.coordinate}")

    def test_holdings_formulas_shape(self):
        ws = self.wb["Holdings"]
        self.assertTrue(str(ws["F2"].value).startswith("=INDEX(Prices!"))
        self.assertEqual(ws["G2"].value, "=D2*E2")
        self.assertEqual(ws["H2"].value, "=D2*F2")
        self.assertEqual(ws["I2"].value, "=H2-G2")
        self.assertIn("IF(G2=0,0", str(ws["J2"].value))

    def test_three_charts(self):
        self.assertEqual(len(self.wb["Dashboard"]._charts), 3)

    def test_ai_report_mentions_real_movers(self):
        _h, price, _i, _m, _p = ground_truth()
        holdings = {r["ticker"]: (int(r["qty"]), float(r["avg_cost"])) for r in read_csv("holdings.csv")}
        pl = {t: price[t][-1] / holdings[t][1] - 1 for t in holdings}
        top = max(pl, key=pl.get)
        ws = self.wb["AI_Report"]
        blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        self.assertIn(top, blob)


class TestRecalc(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.holdings, cls.price, cls.inv, cls.mv, cls.pvals = ground_truth()
        cls.v = recalc(XLSX)

    def _v(self, sheet, coord):
        return self.v[(sheet, coord)]

    def test_no_formula_errors(self):
        errs = [k for k, v in self.v.items() if isinstance(v, tuple)]
        self.assertEqual(errs, [], f"formula errors: {errs[:5]}")

    def test_market_value(self):
        self.assertAlmostEqual(self._v("Dashboard", "B4"), self.mv, delta=0.01)

    def test_invested(self):
        self.assertAlmostEqual(self._v("Dashboard", "B3"), self.inv, delta=0.01)

    def test_pnl_consistency(self):
        b3, b4, b5 = self._v("Dashboard", "B3"), self._v("Dashboard", "B4"), self._v("Dashboard", "B5")
        self.assertAlmostEqual(b5, b4 - b3, delta=0.01)

    def test_allocation_sums_to_100(self):
        s = sum(self._v("Holdings", f"K{i}") for i in range(2, N + 2))
        self.assertAlmostEqual(s, 1.0, delta=0.0005)

    def test_weekly_rollup_ties(self):
        self.assertAlmostEqual(self._v("Weekly", f"C{W + 1}"), self.mv, delta=0.01)

    def test_weekly_return(self):
        exp = self.pvals[-1] / self.pvals[-2] - 1
        self.assertAlmostEqual(self._v("Dashboard", "B7"), exp, delta=1e-9)

    def test_all_checks_pass(self):
        for i in range(2, 10):
            self.assertEqual(self._v("Checks", f"E{i}"), "PASS", f"check row {i}")

    def test_prices_resolve(self):
        n = self._v("Checks", "E5")
        self.assertEqual(n, "PASS")


if __name__ == "__main__":
    unittest.main(verbosity=2)
