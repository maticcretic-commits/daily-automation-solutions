#!/usr/bin/env python3
"""RefreshForge tests: data determinism, pipeline idempotency/dedupe/
quarantine, workbook formula wiring, and LibreOffice headless recalc vs
independent Python ground truth (plus the AGENTS.md regression guards)."""
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
OUT = os.path.join(BASE, "refresh-report.xlsx")
H = 2000


def read_csv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def sandbox_copy():
    """Fresh temp dir with copies of the generated drop files."""
    import tempfile
    tmp = tempfile.mkdtemp(prefix="refreshforge-")
    shutil.copytree(os.path.join(BASE, "drop"), os.path.join(tmp, "drop"))
    shutil.copytree(os.path.join(BASE, "data"), os.path.join(tmp, "data"))
    return tmp


def append_drop_row(tmp, name, row):
    path = os.path.join(tmp, "drop", name)
    rows = read_csv(path)
    with open(path, "a", newline="") as f:
        w = csv.writer(f)
        w.writerow(row)
    return len(rows) + 1


def hist_count(tmp, name):
    path = os.path.join(tmp, "history", name)
    return 0 if not os.path.exists(path) else len(read_csv(path))


class TestData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import generate
        generate.generate()

    def test_deterministic(self):
        import generate
        first = {n: read_csv(os.path.join(BASE, "drop", n))
                 for n in ("sales.csv", "production.csv",
                           "inventory.csv", "issues.csv")}
        generate.generate()
        for n, rows in first.items():
            self.assertEqual(rows, read_csv(os.path.join(BASE, "drop", n)),
                             f"{n} not deterministic")

    def test_shapes(self):
        sales = read_csv(os.path.join(BASE, "drop", "sales.csv"))
        prod = read_csv(os.path.join(BASE, "drop", "production.csv"))
        inv = read_csv(os.path.join(BASE, "drop", "inventory.csv"))
        issues = read_csv(os.path.join(BASE, "drop", "issues.csv"))
        self.assertGreater(len(sales), 100)
        self.assertGreater(len(prod), 30)
        self.assertGreater(len(inv), 30)
        self.assertGreater(len(issues), 10)
        self.assertEqual(len({r["order_id"] for r in sales}), len(sales))
        # breakdown day 14-Oct: Line-B absent, issues spike
        self.assertFalse([r for r in prod
                          if r["date"] == "2026-10-14"
                          and r["line"] == "Line-B"])
        self.assertGreaterEqual(
            sum(1 for r in issues if r["date"] == "2026-10-14"), 4)


class TestPipeline(unittest.TestCase):
    def test_first_refresh_appends_everything(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        entry = refresh.run_refresh(tmp)
        for name in ("sales", "production", "inventory", "issues"):
            s = entry["sources"][name]
            self.assertEqual(s["new"], s["seen"], name)
            self.assertEqual(s["quarantined"], 0, name)
            self.assertEqual(s["duplicates"], 0, name)
            self.assertEqual(hist_count(tmp, s["file"]), s["seen"], name)
        self.assertEqual(entry["status"], "OK")

    def test_second_refresh_is_idempotent(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        refresh.run_refresh(tmp)
        entry = refresh.run_refresh(tmp)
        for name in ("sales", "production", "inventory", "issues"):
            s = entry["sources"][name]
            self.assertEqual(s["new"], 0, name)
            self.assertEqual(s["duplicates"], s["seen"], name)

    def test_delta_append(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        refresh.run_refresh(tmp)
        before = {n: hist_count(tmp, f"{n}.csv")
                  for n in ("sales", "production", "inventory", "issues")}
        append_drop_row(tmp, "sales.csv",
                        ["S9001", "2026-10-24", "Verma Distributors",
                         "OIL-1L", "25", "195.00"])
        append_drop_row(tmp, "sales.csv",
                        ["S9002", "2026-10-24", "Gupta Traders",
                         "DAL-1KG", "40", "148.00"])
        append_drop_row(tmp, "production.csv",
                        ["B9001", "2026-10-24", "Line-A", "OIL-1L",
                         "500", "6"])
        entry = refresh.run_refresh(tmp)
        self.assertEqual(entry["sources"]["sales"]["new"], 2)
        self.assertEqual(entry["sources"]["production"]["new"], 1)
        self.assertEqual(entry["sources"]["inventory"]["new"], 0)
        self.assertEqual(entry["sources"]["issues"]["new"], 0)
        self.assertEqual(hist_count(tmp, "sales.csv"), before["sales"] + 2)
        self.assertEqual(hist_count(tmp, "production.csv"),
                         before["production"] + 1)

    def test_duplicate_rows_are_deduped(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        refresh.run_refresh(tmp)
        # duplicate an existing sales row inside the drop file
        src = os.path.join(tmp, "drop", "sales.csv")
        rows = read_csv(src)
        with open(src, "a", newline="") as f:
            w = csv.writer(f)
            w.writerow([rows[0][c] for c in
                        ["order_id", "date", "customer", "sku", "qty",
                         "unit_price"]])
        entry = refresh.run_refresh(tmp)
        self.assertEqual(entry["sources"]["sales"]["new"], 0)
        self.assertGreaterEqual(entry["sources"]["sales"]["duplicates"], 1)

    def test_bad_rows_are_quarantined(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        refresh.run_refresh(tmp)
        n0 = hist_count(tmp, "sales.csv")
        append_drop_row(tmp, "sales.csv",
                        ["S9003", "2026-13-99", "Verma Distributors",
                         "OIL-1L", "25", "195.00"])      # bad date
        append_drop_row(tmp, "sales.csv",
                        ["S9004", "2026-10-24", "Gupta Traders",
                         "DAL-1KG", "-5", "148.00"])     # negative qty
        append_drop_row(tmp, "inventory.csv",
                        ["T9001", "2026-10-24", "OIL-1L", "SIDEWAYS",
                         "10"])                          # bad movement
        entry = refresh.run_refresh(tmp)
        self.assertEqual(entry["sources"]["sales"]["quarantined"], 2)
        self.assertEqual(entry["sources"]["inventory"]["quarantined"], 1)
        self.assertEqual(hist_count(tmp, "sales.csv"), n0)  # none appended
        self.assertEqual(entry["status"], "OK_WITH_QUARANTINE")

    def test_log_audit_trail(self):
        import refresh
        tmp = sandbox_copy()
        self.addCleanup(shutil.rmtree, tmp, True)
        e1 = refresh.run_refresh(tmp)
        e2 = refresh.run_refresh(tmp)
        self.assertEqual(e1["run_id"], "R0001")
        self.assertEqual(e2["run_id"], "R0002")
        with open(os.path.join(tmp, "refresh_log.json")) as f:
            log = json.load(f)
        self.assertEqual(len(log["runs"]), 2)
        for run in log["runs"]:
            for name, s in run["sources"].items():
                self.assertEqual(len(s["sha1"]), 40, name)


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import generate, refresh, build_workbook
        generate.generate()
        refresh.run_refresh(BASE)   # idempotent — safe on re-runs
        build_workbook.build()
        assert os.path.exists(OUT), "workbook was not built"

    def test_sheets(self):
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        expect = {"Guide", "Refresh_Log", "Raw_Sales", "Raw_Production",
                  "Raw_Inventory", "Raw_Issues", "Daily", "Dashboard",
                  "Checks"}
        self.assertEqual(set(wb.sheetnames), expect)

    def test_formula_columns_present(self):
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        self.assertTrue(str(wb["Raw_Sales"]["G2"].value).startswith("="))
        self.assertTrue(str(wb["Raw_Production"]["G2"].value).startswith("="))
        self.assertTrue(str(wb["Raw_Inventory"]["F2"].value).startswith("="))
        self.assertTrue(str(wb["Daily"]["B2"].value).startswith("="))
        self.assertTrue(str(wb["Dashboard"]["B3"].value).startswith("="))
        self.assertTrue(str(wb["Checks"]["B3"].value).startswith("="))

    def test_no_range_literal_without_column_letter(self):
        """AGENTS.md regression: $H$2:$5000 (missing end column) is invalid."""
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        bad = re.compile(r"\$[A-Z]{1,3}\$\d+:\$\d+")
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for m in bad.finditer(c.value):
                            hits.append(f"{ws.title}!{c.coordinate}: "
                                        f"{m.group(0)}")
        self.assertEqual(hits, [], f"bad range literals: {hits[:5]}")

    def test_no_blank_guard_in_numeric_columns(self):
        """AGENTS.md regression: '' in a numeric formula column -> #VALUE!."""
        from openpyxl import load_workbook
        wb = load_workbook(OUT)
        numeric = {("Raw_Sales", 7), ("Raw_Production", 7),
                   ("Raw_Inventory", 6), ("Daily", 2), ("Daily", 3),
                   ("Daily", 4), ("Daily", 5), ("Daily", 6), ("Daily", 7),
                   ("Dashboard", 2)}
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if ((ws.title, c.column) in numeric
                            and isinstance(c.value, str)
                            and c.value.startswith("=")
                            and '""' in c.value):
                        hits.append(f"{ws.title}!{c.coordinate}")
        self.assertEqual(hits, [], f'"" in numeric columns: {hits[:5]}')

    def test_powerbi_measures_doc(self):
        path = os.path.join(BASE, "powerbi", "MEASURES.md")
        self.assertTrue(os.path.exists(path))
        with open(path) as f:
            text = f.read()
        for m in ("MTD Sales Revenue", "MTD Production Qty", "Average Yield",
                  "Open Issues"):
            self.assertIn(m, text)


def ground_truth():
    """Independent KPI computation straight from history CSVs."""
    sales = read_csv(os.path.join(BASE, "history", "sales.csv"))
    prod = read_csv(os.path.join(BASE, "history", "production.csv"))
    inv = read_csv(os.path.join(BASE, "history", "inventory.csv"))
    issues = read_csv(os.path.join(BASE, "history", "issues.csv"))
    skus = read_csv(os.path.join(BASE, "data", "skus.csv"))
    opening = sum(int(s["opening_qty"]) for s in skus)
    revenue = sum(int(r["qty"]) * float(r["unit_price"]) for r in sales)
    pq = sum(int(r["qty_produced"]) for r in prod)
    rj = sum(int(r["qty_rejected"]) for r in prod)
    yld = 1.0 if pq + rj == 0 else pq / (pq + rj)
    open_issues = sum(1 for r in issues if r["status"] == "Open")
    signed = sum(int(r["qty"]) if r["movement"] == "IN" else -int(r["qty"])
                 for r in inv)
    return {"revenue": revenue, "production": pq, "yield": yld,
            "open_issues": open_issues,
            "on_hand": opening + signed,
            "rows": len(sales) + len(prod) + len(inv) + len(issues)}


def recalc_copy():
    """Headless LibreOffice recalc of a COPY in a separate outdir."""
    tmp = os.path.join(BASE, "tests", "_recalc")
    if os.path.exists(tmp):
        shutil.rmtree(tmp)
    os.makedirs(tmp)
    src = os.path.join(tmp, "src.xlsx")
    shutil.copy(OUT, src)
    outdir = os.path.join(tmp, "out")
    os.makedirs(outdir)
    p = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                        "--outdir", outdir, src],
                       capture_output=True, text=True, timeout=300)
    outs = [f for f in os.listdir(outdir) if f.endswith(".xlsx")]
    assert outs, f"LO convert produced nothing: {p.stderr[-500:]}"
    return os.path.join(outdir, outs[0])


class TestRecalc(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import generate, refresh, build_workbook
        generate.generate()
        refresh.run_refresh(BASE)
        build_workbook.build()
        from openpyxl import load_workbook
        cls.wb = load_workbook(recalc_copy(), data_only=True)
        cls.gt = ground_truth()

    def test_no_formula_errors(self):
        errs = {"#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#N/A", "#NUM!",
                "#NULL!"}
        hits = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v in errs:
                        hits.append(f"{ws.title}!{c.coordinate}={v}")
        self.assertEqual(hits, [], f"formula errors: {hits[:8]}")

    def test_dashboard_kpis(self):
        dash = self.wb["Dashboard"]
        gt = self.gt
        self.assertAlmostEqual(dash["B3"].value, gt["revenue"], places=2)
        self.assertEqual(dash["B4"].value, gt["production"])
        self.assertAlmostEqual(dash["B5"].value, gt["yield"], places=4)
        self.assertEqual(dash["B6"].value, gt["open_issues"])
        self.assertEqual(dash["B7"].value, gt["on_hand"])
        self.assertEqual(dash["B8"].value, gt["rows"])
        self.assertIn(dash["B10"].value, ("OK", "OK_WITH_QUARANTINE"))

    def test_all_checks_true(self):
        ch = self.wb["Checks"]
        for r in range(3, 14):
            self.assertTrue(ch.cell(row=r, column=2).value,
                            f"check row {r} ({ch.cell(row=r, column=1).value}) "
                            "is not TRUE")

    def test_daily_ties(self):
        dy = self.wb["Daily"]
        self.assertAlmostEqual(
            sum(dy.cell(row=r, column=2).value or 0 for r in range(2, 33)),
            self.gt["revenue"], places=2)
        self.assertEqual(
            sum(dy.cell(row=r, column=3).value or 0 for r in range(2, 33)),
            self.gt["production"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
